"""python-pptx 绘制原语：文字、矩形 / 圆角矩形 / 椭圆。所有 box 都是画布逻辑像素 (x, y, w, h)。"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

from .errors import PptError
from .geometry import Box, px_to_emu

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "rounded": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
          "pentagon": MSO_SHAPE.PENTAGON}  # pentagon = 指向右侧的五边形箭头（homePlate）


def rgb(style: dict, token: str) -> RGBColor:
    if token not in style["colors"]:
        raise PptError(f"未知颜色 token: {token}")
    return RGBColor.from_string(style["colors"][token].lstrip("#"))


def emu_box(box: Box) -> tuple[Emu, Emu, Emu, Emu]:
    return tuple(Emu(px_to_emu(v)) for v in box)


def _text_style(style: dict, token: str) -> dict:
    if token not in style["text_styles"]:
        raise PptError(f"未知文字样式 token: {token}")
    return style["text_styles"][token]


def add_run(p, style: dict, text: str, text_style: str, color=None, bold=None):
    """追加一个 run；显式写 a:latin / a:ea / a:cs。"""
    ts = _text_style(style, text_style)
    fonts = style["fonts"]
    run = p.add_run()
    run.text = text
    font = run.font
    font.size = Pt(ts["size"])
    font.bold = ts.get("bold", False) if bold is None else bold
    font.color.rgb = rgb(style, color or ts["color"])
    font.name = fonts["latin"]  # 只写 a:latin，a:ea / a:cs 紧随其后手动补上
    latin = run._r.rPr.find(qn("a:latin"))
    ea = parse_xml(f'<a:ea {nsdecls("a")} typeface="{fonts["cjk"]}"/>')
    cs = parse_xml(f'<a:cs {nsdecls("a")} typeface="{fonts["cjk"]}"/>')
    latin.addnext(ea)
    ea.addnext(cs)


def _setup_frame(tf, valign: str, inset):
    """inset 是四边统一的内边距（px），或 (左右, 上下)。"""
    ix, iy = inset if isinstance(inset, tuple) else (inset, inset)
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = VALIGN[valign]
    tf.margin_left = tf.margin_right = Emu(px_to_emu(ix))
    tf.margin_top = tf.margin_bottom = Emu(px_to_emu(iy))


def set_text(tf, style: dict, text, text_style: str, *, color=None, bold=None, align="left", valign="top", inset=0):
    """写入文字并设置文本框属性；按换行拆成段落，每段一个 run。"""
    _setup_frame(tf, valign, inset)
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        add_run(p, style, line, text_style, color, bold)


def set_paragraphs(tf, style: dict, paras, *, align="left", valign="top", inset=0, gap=0):
    """写入多段文字，每段一个 run、各用各的文字样式；paras = [(text, text_style, color, bold), ...]，
    gap 是第 2 段起每段与上一段的间距（px）。"""
    _setup_frame(tf, valign, inset)
    for i, (text, text_style, color, bold) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if i and gap:
            p.space_before = Emu(px_to_emu(gap))
        add_run(p, style, text, text_style, color, bold)


def add_text(slide, style: dict, box: Box, text, text_style: str, **opts):
    """纯文字框：四边内边距为 0（可用 inset 覆盖）。"""
    tb = slide.shapes.add_textbox(*emu_box(box))
    set_text(tb.text_frame, style, text, text_style, **opts)
    return tb


def add_runs(slide, style: dict, box: Box, runs, *, align="left", valign="top"):
    """单段落、多 run 的文字框（如 24 号标题 + 18 号后半句，或「加粗名称：」+ 正文），
    runs = [(text, text_style) 或 (text, text_style, color), ...]，run 之间自然连排、一起折行。"""
    tb = slide.shapes.add_textbox(*emu_box(box))
    _setup_frame(tb.text_frame, valign, 0)
    p = tb.text_frame.paragraphs[0]
    p.alignment = ALIGN[align]
    for text, text_style, *color in runs:
        add_run(p, style, text, text_style, *color)
    return tb


BULLET = "•"


def add_bullets(slide, style: dict, box: Box, items, text_style: str, *, color=None):
    """要点列表：每项一段，带项目符号；项是字符串，或 {text, level}（level 从 0 起，每级多缩进一个字宽）。"""
    tb = slide.shapes.add_textbox(*emu_box(box))
    _setup_frame(tb.text_frame, "top", 0)
    em = px_to_emu(_text_style(style, text_style)["size"] * 96 / 72)
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else (item["text"], item.get("level", 0))
        p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
        p_pr = p._p.get_or_add_pPr()
        p_pr.set("marL", str(em * (level + 1)))
        p_pr.set("indent", str(-em))
        if i:  # a:pPr 子元素顺序：spcBef 必须在 buFont / buChar 之前
            p_pr.append(parse_xml(f'<a:spcBef {nsdecls("a")}><a:spcPts val="400"/></a:spcBef>'))
        p_pr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="{style["fonts"]["latin"]}"/>'))
        p_pr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="{BULLET}"/>'))
        add_run(p, style, text, text_style, color)
    return tb


def _outer_shadow(style: dict):
    """风格 shadow 小节 → a:outerShdw；distance / blur 是逻辑像素，angle 是度（90 = 向下投影）。"""
    shadow = style.get("shadow")
    if not shadow:
        raise PptError("风格没有定义 shadow，shape 不能用 shadow: true")
    return parse_xml(
        f'<a:outerShdw {nsdecls("a")} blurRad="{px_to_emu(shadow["blur"])}" dist="{px_to_emu(shadow["distance"])}" '
        f'dir="{round(shadow["angle"] * 60000)}" algn="ctr" rotWithShape="0">'
        f'<a:srgbClr val="{rgb(style, shadow["color"])}"><a:alpha val="{round(shadow["alpha"] * 100000)}"/></a:srgbClr>'
        f'</a:outerShdw>')


def add_shape(slide, style: dict, box: Box, fill: str, *, line=None, line_width=0.75, kind="rect", shadow=False, text=None,
              **text_opts):
    """实心形状（kind: rect / rounded / oval）；line 给定才画 line_width pt 的描边。
    默认无阴影；shadow=True 按风格的 shadow 小节画投影。text 直接写在形状里，便于居中。"""
    shape = slide.shapes.add_shape(SHAPES[kind], *emu_box(box))
    shape.shadow.inherit = False  # PowerPoint 认空的 a:effectLst
    shape._element.find(qn("p:style")).find(qn("a:effectRef")).set("idx", "0")  # LibreOffice 仍会按主题画阴影
    if shadow:
        shape._element.spPr.find(qn("a:effectLst")).append(_outer_shadow(style))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(style, fill)
    if line:
        shape.line.color.rgb = rgb(style, line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    if text is not None:
        set_text(shape.text_frame, style, text, **text_opts)
    return shape


def add_image(slide, box: Box, path, *, fit="contain"):
    """把图片放进 box：contain = 等比缩放到整张图都在 box 内、居中；cover = 等比缩放到铺满 box、裁掉超出的部分。"""
    pic = slide.shapes.add_picture(str(path), 0, 0)
    iw, ih = pic.image.size
    x, y, w, h = box
    scale = (max if fit == "cover" else min)(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    if fit == "cover":  # 显示尺寸 = box，按超出的比例裁掉左右或上下
        pic.crop_left = pic.crop_right = (1 - w / dw) / 2
        pic.crop_top = pic.crop_bottom = (1 - h / dh) / 2
        dw, dh = w, h
    pic.left, pic.top, pic.width, pic.height = (Emu(px_to_emu(v)) for v in (x + (w - dw) / 2, y + (h - dh) / 2, dw, dh))
    return pic


def add_arrow(slide, style: dict, begin, end, *, color="muted", width=1):
    """从 begin 到 end（逻辑像素 (x, y)）的直线，end 一侧带三角箭头；无阴影。"""
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(Emu(px_to_emu(v)) for v in (*begin, *end)))
    conn._element.find(qn("p:style")).find(qn("a:effectRef")).set("idx", "0")  # 同 add_shape：不按主题画阴影
    conn.line.color.rgb = rgb(style, color)
    conn.line.width = Pt(width)
    conn.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>'))
    return conn
