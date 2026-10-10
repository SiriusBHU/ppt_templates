"""ppt import：把骨架 pptx 的若干页导成 agent 可读的草稿 shapes.yaml 和原图 original.png。

坐标统一换算成 1280×720 画布的逻辑像素（按页宽缩放），颜色和文字样式对应到默认风格的最近 token。
"""

import colorsys
import shutil
import tempfile
from pathlib import Path

import yaml
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .errors import PptError
from .geometry import CANVAS_H, CANVAS_W, REF_H, REF_W
from .loader import load_style, load_yaml
from .preview import render_pngs
from .validate import TEXT_COLORS

STYLE = "default"
PLACEHOLDER_LABELS = {"图片", "图表"}
# 字号 → 文字样式 token；14、12 和 18 按是否加粗二选一
SIZE_STYLES = {24: "title", 18: ("subtitle", "title_suffix"), 14: ("lead", "emphasis"), 12: ("body", "body_strong"),
               10: "detail", 9: "caption", 8: "footnote"}
PRESET_KINDS = {"rect": "rect", "roundRect": "rounded", "ellipse": "oval", "line": "line", "straightConnector1": "line"}
SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
PRESET_COLORS = {"black": "000000", "white": "FFFFFF"}
DEFAULT_SIZE = 18  # 没写字号时 PowerPoint 的默认值
HEADER_SIZE = 18   # 标题区文字（title / title_suffix / takeaway）的最小字号
EDGE_TOLERANCE = 8  # 判断是否超出 body 区时的容差（px）
NEUTRAL_SPREAD = 24  # 文字颜色三个通道相差不超过这个值就算灰色
LIGHT_TEXT = 0xE0    # 灰色文字三个通道都不低于这个值，当作深色底上的白字


class _Flow(dict):
    """YAML 里写成一行的小映射，如 {hex: '#C00000', token: accent}。"""


class _Dumper(yaml.SafeDumper):
    def ignore_aliases(self, data):  # 同一个颜色对象出现多次时也要写全，不要 &id001 / *id001
        return True


_Dumper.add_representer(_Flow, lambda d, data: d.represent_mapping("tag:yaml.org,2002:map", data, flow_style=True))
_Dumper.add_representer(list, lambda d, data: d.represent_sequence(
    "tag:yaml.org,2002:seq", data, flow_style=all(not isinstance(x, (dict, list)) for x in data)))


def _num(value: float):
    value = round(value, 1)
    return int(value) if value == int(value) else value


class _Context:
    def __init__(self, scale: float, theme: dict, colors: dict, body: list):
        self.scale = scale
        self.theme = theme
        self.colors = colors
        self.body = body
        self.leaves: list[dict] = []


# ---- 颜色 ----

def _theme_colors(master) -> dict[str, str]:
    """母版主题里的配色表 {dk1: 'RRGGBB', accent1: ...}；取不到返回空表。"""
    try:
        scheme = etree.fromstring(master.part.part_related_by(RT.THEME).blob).find(".//" + qn("a:clrScheme"))
    except (KeyError, etree.XMLSyntaxError):
        return {}
    colors = {}
    for child in scheme if scheme is not None else []:
        inner = child[0]
        colors[etree.QName(child).localname] = inner.get("lastClr") if etree.QName(inner).localname == "sysClr" else inner.get("val")
    return colors


def _hex_of(clr, theme: dict) -> str | None:
    """颜色元素（srgbClr / schemeClr / sysClr / prstClr）→ 'RRGGBB'，应用 lumMod / lumOff；解析不了返回 None。"""
    tag = etree.QName(clr).localname
    if tag == "srgbClr":
        base = clr.get("val")
    elif tag == "sysClr":
        base = clr.get("lastClr")
    elif tag == "schemeClr":
        name = clr.get("val")
        base = theme.get(SCHEME_ALIAS.get(name, name))
    elif tag == "prstClr":
        base = PRESET_COLORS.get(clr.get("val"))
    else:
        base = None
    if not base:
        return None
    h, l, s = colorsys.rgb_to_hls(*(int(base[i:i + 2], 16) / 255 for i in (0, 2, 4)))
    for mod in clr:
        name = etree.QName(mod).localname
        if name == "lumMod":
            l *= int(mod.get("val")) / 100000
        elif name == "lumOff":
            l += int(mod.get("val")) / 100000
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
    return f"{round(r * 255):02X}{round(g * 255):02X}{round(b * 255):02X}"


def _fill_in(container, theme: dict):
    """container（spPr / a:ln / tcPr）里直接写的填充：(是否写了, 颜色)；noFill 算写了但没有颜色。"""
    for child in container:
        tag = etree.QName(child).localname
        if tag == "noFill":
            return True, None
        if tag == "solidFill":
            return True, _hex_of(child[0], theme)
        if tag == "gradFill":
            stop = child.find(".//" + qn("a:gs"))
            return True, _hex_of(stop[0], theme) if stop is not None else None
    return False, None


def _style_ref_hex(el, ref: str, theme: dict) -> str | None:
    """形状 p:style 里的 fillRef / lnRef / fontRef 指向的颜色（idx 为 0 表示不画）。"""
    style = el.find(qn("p:style"))
    node = style.find(qn(f"a:{ref}")) if style is not None else None
    if node is None or node.get("idx") == "0" or len(node) == 0:
        return None
    return _hex_of(node[0], theme)


def _paint(el, kind: str, ctx: _Context) -> str | None:
    """形状的填充（kind='fill'）或轮廓（kind='ln'）颜色：先看自己写的，没写再看 p:style。"""
    sppr = el.find(qn("p:spPr"))
    container = sppr if kind == "fill" or sppr is None else sppr.find(qn("a:ln"))
    if container is not None:
        found, hex_ = _fill_in(container, ctx.theme)
        if found:
            return hex_
    return _style_ref_hex(el, "fillRef" if kind == "fill" else "lnRef", ctx.theme)


def _nearest(hex_: str, palette: dict[str, str]) -> tuple[str, int]:
    """palette（{token: '#RRGGBB'}）里与 hex_ 最近的 token 和 RGB 距离的平方；并列取靠前的。"""
    rgb = [int(hex_[i:i + 2], 16) for i in (0, 2, 4)]

    def dist(token: str) -> int:
        return sum((a - int(palette[token][i:i + 2], 16)) ** 2 for a, i in zip(rgb, (1, 3, 5)))

    best = min(palette, key=dist)
    return best, dist(best)


def _color(hex_: str | None, ctx: _Context, text: bool = False) -> _Flow | None:
    """{hex, token}：token 是默认风格里最近的颜色；不完全一致时加 approx。
    文字颜色只在 TEXT_COLORS 里选，且灰色（各通道相近）不按距离就近取靛蓝：偏浅的归 on_primary，其余归 text。"""
    if hex_ is None:
        return None
    channels = [int(hex_[i:i + 2], 16) for i in (0, 2, 4)]
    if text and max(channels) - min(channels) <= NEUTRAL_SPREAD:
        token = "on_primary" if min(channels) >= LIGHT_TEXT else "text"
        dist = sum((c - int(ctx.colors[token][i:i + 2], 16)) ** 2 for c, i in zip(channels, (1, 3, 5)))
    else:
        token, dist = _nearest(hex_, {t: c for t, c in ctx.colors.items() if not text or t in TEXT_COLORS})
    out = _Flow(hex=f"#{hex_.upper()}", token=token)
    if dist:
        out["approx"] = True
    return out


# ---- 文字 ----

def _style_token(size: float, bold: bool) -> str:
    snapped = min(SIZE_STYLES, key=lambda s: (abs(s - size), -s))
    token = SIZE_STYLES[snapped]
    return token[bold] if isinstance(token, tuple) else token


def _run_info(run, default_hex: str | None, ctx: _Context) -> dict:
    size = run.font.size.pt if run.font.size is not None else None
    bold = bool(run.font.bold)
    rpr = run._r.find(qn("a:rPr"))
    found, hex_ = _fill_in(rpr, ctx.theme) if rpr is not None else (False, None)
    info = {"text": run.text, "size": _num(size or DEFAULT_SIZE), "style": _style_token(size or DEFAULT_SIZE, bold),
            "bold": bold, "color": _color(hex_ if found else default_hex, ctx, text=True)}
    if size is None:
        info["size_inherited"] = True
    return info


def _paragraphs(text_frame, default_hex: str | None, ctx: _Context) -> list[dict]:
    out = []
    for para in text_frame.paragraphs:
        text = para.text.replace("\v", "\n")
        if not text.strip():
            continue
        runs = [_run_info(r, default_hex, ctx) for r in para.runs if r.text]
        first = runs[0] if runs else {"size": DEFAULT_SIZE, "style": _style_token(DEFAULT_SIZE, False), "bold": False,
                                      "color": _color(default_hex, ctx, text=True)}
        item = {"text": text, **{k: first[k] for k in ("size", "style", "bold", "color") if first.get(k) is not None}}
        if first.get("size_inherited"):
            item["size_inherited"] = True
        item["align"] = para.alignment.name.lower() if para.alignment is not None else "left"
        if para.level:
            item["level"] = para.level
        ppr = para._p.find(qn("a:pPr"))
        if ppr is not None and (ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buAutoNum")) is not None):
            item["bullet"] = True
        looks = {(r["size"], r["style"], r["bold"], (r["color"] or {}).get("hex")) for r in runs}
        if len(looks) > 1:  # 同一段里字号 / 粗细 / 颜色不一致：逐 run 列出（对应模板的 runs）
            item["runs"] = [{k: r[k] for k in ("text", "size", "style", "bold", "color") if r.get(k) is not None} for r in runs]
        out.append(item)
    return out


# ---- 形状 ----

def _box(x, y, w, h, tf, ctx: _Context) -> list:
    ax, bx, ay, by = tf
    s = ctx.scale
    return [_num((ax * x + bx) * s), _num((ay * y + by) * s), _num(ax * w * s), _num(ay * h * s)]


def _body_box(box: list, ctx: _Context) -> list:
    """画布 box → 相对 body 区左上角、按 1216×560 参考尺寸缩放的 box，可直接写进 layout（body 区不是参考尺寸时才和画布像素不同）。"""
    bx, by, bw, bh = ctx.body
    sx, sy = REF_W / bw, REF_H / bh
    return [_num((box[0] - bx) * sx), _num((box[1] - by) * sy), _num(box[2] * sx), _num(box[3] * sy)]


def _group_transform(el, tf) -> tuple:
    """组内子形状坐标 → 页面 EMU 的变换 (ax, bx, ay, by)：x_页 = ax * x_子 + bx。"""
    xfrm = el.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    off, ext, choff, chext = (xfrm.find(qn(f"a:{n}")) for n in ("off", "ext", "chOff", "chExt"))
    if None in (off, ext, choff, chext):
        return tf
    sx = int(ext.get("cx")) / int(chext.get("cx")) if int(chext.get("cx")) else 1
    sy = int(ext.get("cy")) / int(chext.get("cy")) if int(chext.get("cy")) else 1
    ax, bx, ay, by = tf
    return (ax * sx, ax * (int(off.get("x")) - int(choff.get("x")) * sx) + bx,
            ay * sy, ay * (int(off.get("y")) - int(choff.get("y")) * sy) + by)


def _line_info(el, ctx: _Context) -> dict:
    ln = el.find(qn("p:spPr")).find(qn("a:ln")) if el.find(qn("p:spPr")) is not None else None
    info = {}
    color = _color(_paint(el, "ln", ctx), ctx)
    if color:
        info["line"] = color
    if ln is not None and ln.get("w"):
        info["line_width"] = _num(int(ln.get("w")) * ctx.scale)
    if ln is not None:
        dash = ln.find(qn("a:prstDash"))
        if dash is not None and dash.get("val") != "solid":
            info["dash"] = dash.get("val")
    return info


def _arrow(el) -> str | None:
    ln = el.find(qn("p:spPr")).find(qn("a:ln")) if el.find(qn("p:spPr")) is not None else None
    if ln is None:
        return None
    has = {n: ln.find(qn(f"a:{n}")) is not None and ln.find(qn(f"a:{n}")).get("type", "none") != "none"
           for n in ("headEnd", "tailEnd")}
    return {(True, True): "both", (True, False): "start", (False, True): "end"}.get((has["headEnd"], has["tailEnd"]))


def _table(shape, tf, ctx: _Context) -> dict:
    table = shape.table
    ax, _, ay, _ = tf
    rows = []
    for r in table.rows:
        row = []
        for cell in r.cells:
            if cell.is_spanned:
                row.append({"merged": True})
                continue
            item = {}
            text = _paragraphs(cell.text_frame, ctx.theme.get("dk1"), ctx)
            if text:
                item["text"] = text
            tcpr = cell._tc.find(qn("a:tcPr"))
            fill = _color(_fill_in(tcpr, ctx.theme)[1], ctx) if tcpr is not None else None
            if fill:
                item["fill"] = fill
            if cell.is_merge_origin:
                item["span"] = [cell.span_height, cell.span_width]
            row.append(item)
        rows.append(row)
    return {"rows": len(table.rows), "cols": len(table.columns), "first_row_header": bool(table.first_row),
            "col_widths": [_num(c.width * ax * ctx.scale) for c in table.columns],
            "row_heights": [_num(r.height * ay * ctx.scale) for r in table.rows], "cells": rows}


def _node(shape, tf, ctx: _Context) -> dict | None:
    el = shape._element
    tag = etree.QName(el).localname
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    node = {"id": shape.shape_id, "name": shape.name}
    box = _box(shape.left, shape.top, shape.width, shape.height, tf, ctx)

    if tag == "grpSp":
        node.update(kind="group", box=box, body_box=_body_box(box, ctx))
        node["children"] = _walk(shape.shapes, _group_transform(el, tf), ctx)
        return node

    if tag == "cxnSp":
        xfrm = el.find(qn("p:spPr")).find(qn("a:xfrm"))
        flip_h, flip_v = xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"
        x, y, w, h = box
        node.update(kind="line", box=box, body_box=_body_box(box, ctx),
                    start=[_num(x + (w if flip_h else 0)), _num(y + (h if flip_v else 0))],
                    end=[_num(x + (0 if flip_h else w)), _num(y + (0 if flip_v else h))], **_line_info(el, ctx))
        if arrow := _arrow(el):
            node["arrow"] = arrow
        ctx.leaves.append(node)
        return node

    if tag == "graphicFrame":
        kind = "table" if shape.has_table else "chart" if shape.has_chart else "graphic"
        node.update(kind=kind, box=box, body_box=_body_box(box, ctx))
        if kind == "table":
            node["table"] = _table(shape, tf, ctx)
        ctx.leaves.append(node)
        return node

    if tag == "pic":
        node.update(kind="picture", suggest="image", box=box, body_box=_body_box(box, ctx))
        ctx.leaves.append(node)
        return node

    prst = el.find(".//" + qn("a:prstGeom"))
    prst = prst.get("prst") if prst is not None else None
    cnv = el.find(qn("p:nvSpPr")).find(qn("p:cNvSpPr"))
    if cnv is not None and cnv.get("txBox") == "1" and prst in (None, "rect"):
        kind = "textbox"
    elif prst is None:
        kind = "freeform" if el.find(".//" + qn("a:custGeom")) is not None else "rect"
    else:
        kind = PRESET_KINDS.get(prst, f"shape:{prst}")
    text = _paragraphs(shape.text_frame, _style_ref_hex(el, "fontRef", ctx.theme) or ctx.theme.get("dk1"), ctx) \
        if shape.has_text_frame else []
    if len(text) == 1 and text[0]["text"].strip("「」【】[] ") in PLACEHOLDER_LABELS:
        node.update(kind="placeholder", label=text[0]["text"].strip("「」【】[] "))
        if node["label"] == "图片":
            node["suggest"] = "image"
        text = []
    else:
        node["kind"] = kind
    node.update(box=box, body_box=_body_box(box, ctx))
    if shape.is_placeholder:
        node["layout_placeholder"] = shape.placeholder_format.type.name.lower()
    if (rotation := getattr(shape, "rotation", 0)):
        node["rotation"] = _num(rotation)
    if kind != "line":
        if fill := _color(_paint(el, "fill", ctx), ctx):
            node["fill"] = fill
        node.update(_line_info(el, ctx))
    sppr = el.find(qn("p:spPr"))
    node["shadow"] = sppr is not None and sppr.find(qn("a:effectLst") + "/" + qn("a:outerShdw")) is not None
    if text:
        anchor = shape.text_frame.vertical_anchor
        node["valign"] = anchor.name.lower() if anchor is not None else "top"
        node["text"] = text
    ctx.leaves.append(node)
    return node


def _walk(shapes, tf, ctx: _Context) -> list[dict]:
    return [n for n in (_node(s, tf, ctx) for s in shapes) if n is not None]


# ---- 页框判断 ----

def _max_size(node: dict) -> float:
    return max((p["size"] for p in node.get("text", [])), default=0)


def _guess_body(leaves: list[dict], style: dict, key: str) -> dict:
    """是否适合 frame: standard 且正文区用 style.frame[key]（body 或 body_compact）：左上角有标题文字，且其余元素都在该正文区内。"""
    bx, by, bw, bh = style["frame"][key]["box"]
    compact = key == "body_compact"
    tx, ty, _, th = style["frame"]["title"]["box"]
    top = ty + th if compact else by  # 紧凑正文只要求在标题行之下（标题与正文的间距实际常小于 16px）
    header = [n for n in leaves if n.get("text") and _max_size(n) >= HEADER_SIZE
              and n["box"][1] + n["box"][3] <= by + EDGE_TOLERANCE]  # 标题、后半句（standard 还有副标题）都在正文区上方
    header_ids = {id(n) for n in header}
    titles = [n for n in header if n["box"][0] <= tx + 48 and n["box"][1] <= ty + 48]
    if not titles:
        return {"frame": "none", "reason": f"左上角（x≤{tx + 48}，y≤{ty + 48}）没有 {HEADER_SIZE}pt 以上的标题文字，按整页自由布局处理"}
    title = min(titles, key=lambda n: (n["box"][1], -_max_size(n)))
    outside = [n["id"] for n in leaves
               if id(n) not in header_ids and n["kind"] != "line" and n["box"][2] * n["box"][3] > 0
               and (n["box"][1] < top - EDGE_TOLERANCE or n["box"][0] < bx - EDGE_TOLERANCE
                    or n["box"][0] + n["box"][2] > bx + bw + EDGE_TOLERANCE or n["box"][1] + n["box"][3] > by + bh + EDGE_TOLERANCE)]
    guess = {"title_shape": title["id"]}
    if outside:
        return {"frame": "none", "reason": f"标题在左上角，但有 {len(outside)} 个元素超出 body 区 [{bx}, {by}, {bw}×{bh}]（含页边距、顶部或底部）："
                f"shape id {', '.join(map(str, outside))}", **guess}
    if compact:
        return {"frame": "standard", "body": "compact",
                "reason": f"左上角有标题文字（{_max_size(title):g}pt），没有副标题，其余元素都在 body_compact 区 [{bx}, {by}, {bw}×{bh}] 内："
                          f"模板不要设 takeaway 槽位，正文区用 body_compact；body_box 已按 1216×560 参考尺寸换算（纵向放大 {bh}/{REF_H}）",
                **guess}
    return {"frame": "standard", "body": "standard",
            "reason": f"左上角有标题文字（{_max_size(title):g}pt），其余元素都在 body 区 [{bx}, {by}, {bw}×{bh}] 内", **guess}


def _guess_frame(leaves: list[dict], style: dict) -> dict:
    """先按标准 body 判断；放不下（内容紧贴标题、没有副标题行）再试 body_compact，都不行就取标准判断的结果。"""
    standard = _guess_body(leaves, style, "body")
    if standard["frame"] == "standard" or "body_compact" not in style["frame"]:
        return standard
    compact = _guess_body(leaves, style, "body_compact")
    return compact if compact["frame"] == "standard" else standard


# ---- 入口 ----

def _manifest_purpose(path: Path, n: int, warnings: list[str]) -> str | None:
    """manifest.yaml 里第 n 页的 purpose。slides 可以是列表（元素带 slide / index / n / page 编号，或按顺序）或以页码为键的映射。"""
    if not path.is_file():
        return None
    try:
        data = load_yaml(path)
    except PptError as e:
        warnings.append(f"manifest.yaml 读取失败，purpose 留空：{e}")
        return None
    slides = data.get("slides", data) if isinstance(data, dict) else data
    entry = None
    if isinstance(slides, dict):
        entry = next((v for k, v in slides.items() if str(k) in (str(n), f"slide-{n}", f"slide{n}", f"slide_{n}")), None)
    elif isinstance(slides, list):
        keys = ("slide", "index", "n", "no", "number", "page", "id")
        entry = next((e for e in slides if isinstance(e, dict) and any(str(e.get(k)) == str(n) for k in keys)), None)
        if entry is None and n <= len(slides) and not any(isinstance(e, dict) and any(k in e for k in keys) for e in slides):
            entry = slides[n - 1]
    purpose = entry.get("purpose") if isinstance(entry, dict) else None
    return str(purpose) if purpose else None


def import_skeleton(root: Path, pptx: Path, out_dir: Path, slides: list[int] | None = None) -> tuple[list[Path], list[str]]:
    """把 pptx 的指定页（默认全部，页码从 1 起）导到 out_dir/slide-N/；返回 (各页目录, 警告)。"""
    if not pptx.is_file():
        raise PptError(f"文件不存在: {pptx}")
    try:
        prs = Presentation(str(pptx))
    except Exception as e:  # python-pptx 对坏文件抛各种异常
        raise PptError(f"无法打开 {pptx}：{e}") from e
    total = len(prs.slides)
    wanted = list(dict.fromkeys(slides or range(1, total + 1)))
    bad = [n for n in wanted if not 1 <= n <= total]
    if bad:
        raise PptError(f"页码 {', '.join(map(str, bad))} 超出范围（共 {total} 页）")

    style = load_style(root, STYLE)
    colors = {token: value.upper() for token, value in style["colors"].items()}
    scale = CANVAS_W / prs.slide_width
    warnings = []
    if abs(prs.slide_width / prs.slide_height - CANVAS_W / CANVAS_H) > 0.01:
        warnings.append(f"页面比例不是 16:9（{prs.slide_width}×{prs.slide_height} EMU），坐标仍按页宽缩放到 1280 宽，纵向可能与 720 不符")

    dirs = []
    for n in wanted:
        slide = prs.slides[n - 1]
        ctx = _Context(scale, _theme_colors(slide.slide_layout.slide_master), colors, style["frame"]["body"]["box"])
        shapes = _walk(slide.shapes, (1, 0, 1, 0), ctx)
        guess = _guess_frame(ctx.leaves, style)
        if guess.get("body") == "compact":
            ctx.body = style["frame"]["body_compact"]["box"]
            shapes = _walk(slide.shapes, (1, 0, 1, 0), ctx)  # body_box 要按紧凑 body 的原点和缩放重算
        doc = {"slide": n, "source": pptx.name}
        if purpose := _manifest_purpose(pptx.parent / "manifest.yaml", n, warnings):
            doc["purpose"] = purpose
        doc.update(canvas=[CANVAS_W, CANVAS_H], body_origin=ctx.body[:2], frame_guess=guess, shapes=shapes)
        slide_dir = out_dir / f"slide-{n}"
        slide_dir.mkdir(parents=True, exist_ok=True)
        (slide_dir / "shapes.yaml").write_text(
            yaml.dump(doc, Dumper=_Dumper, allow_unicode=True, sort_keys=False, width=10**6), encoding="utf-8")
        dirs.append(slide_dir)

    if shutil.which("soffice") is None:
        warnings.append("找不到 soffice，没有生成 original.png（安装 LibreOffice 后重新运行）")
    else:
        try:
            with tempfile.TemporaryDirectory() as tmp:
                pngs = render_pngs(prs, Path(tmp))
                for n, slide_dir in zip(wanted, dirs):
                    shutil.copyfile(pngs[n - 1], slide_dir / "original.png")
        except PptError as e:
            warnings.append(f"original.png 生成失败：{e}")
    return dirs, warnings
