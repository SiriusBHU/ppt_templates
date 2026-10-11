"""deck → Presentation：页框 + 布局元素（shape / text / bullets / grid / image / component）。"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from .components import COMPONENTS
from .draw import add_bullets, add_image, add_runs, add_shape, add_text
from .errors import PptError
from .geometry import CANVAS_H, CANVAS_W, REF_H, REF_W, Transform, body_box, grid_cells, px_to_emu, resolve_box
from .images import ImageSource
from .loader import load_style
from .paths import INDEX_BIND, is_empty, lookup
from .schemas import pointer
from .slots import prepare_pages

BLANK_LAYOUT = 6
FRAME_DEFAULT_STYLE = {"title": "title", "title_suffix": "title_suffix", "takeaway": "subtitle"}
CANVAS_AREA = (0, 0, CANVAS_W, CANVAS_H)


def render_deck(deck: dict, root: Path, images: ImageSource | None = None) -> Presentation:
    """images 决定 image 元素的路径相对哪里解析、能引用哪些文件；省略则相对当前目录、不限制。"""
    images = images or ImageSource(Path.cwd())
    pages = prepare_pages(deck, root, images)
    style = load_style(root, deck.get("style", "default"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(px_to_emu(CANVAS_W)), Emu(px_to_emu(CANVAS_H))
    for template, data in pages:
        slide = prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT])
        if template.get("frame", "standard") == "standard":
            _render_frame(slide, style, data)
            bx, by, bw, bh = body_box(style, data)
            _render_elements(slide, style, template["layout"], (0, 0, REF_W, REF_H), data,
                             Transform(bx, by, bw / REF_W, bh / REF_H), images=images)
        else:
            _render_elements(slide, style, template["layout"], CANVAS_AREA, data, Transform(), images=images)
    return prs


def _render_frame(slide, style, data):
    frame = style["frame"]
    _render_elements(slide, style, frame.get("decorations", []), CANVAS_AREA, {}, Transform())
    regions = {name: data[name] for name in ("title", "takeaway") if not is_empty(data.get(name))}
    for name, text in regions.items():
        region = frame[name]
        box = resolve_box(region["box"], CANVAS_W, CANVAS_H)
        text_style = region.get("style", FRAME_DEFAULT_STYLE[name])
        align, valign = region.get("align", "left"), region.get("valign", "top")
        if name == "title" and data.get("title_suffix"):
            # 标题主体与后半句同一行、共用基线；全角空格隔开
            suffix_style = region.get("suffix_style", FRAME_DEFAULT_STYLE["title_suffix"])
            add_runs(slide, style, box, [(text, text_style), ("　" + data["title_suffix"], suffix_style)],
                     align=align, valign=valign)
        else:
            add_text(slide, style, box, text, text_style, align=align, valign=valign)


def _mapped(mapping, value):
    """color_map / fill_map / line_map 的取色：值转成字符串后按键匹配；没命中或没有值返回 None。"""
    if not mapping or value is None:
        return None
    return next((token for key, token in mapping.items() if str(key) == str(value)), None)


def _text_of(scope, el, index):
    """text 元素或 run 的文字：value 是字面量，bind 按 scope 取值（"#" 是 grid 当前项序号）；取不到返回 None。"""
    if "value" in el:
        return el["value"]
    if el["bind"] == INDEX_BIND:
        return str(index)
    return lookup(scope, el["bind"])


def _render_elements(slide, style, elements, area, scope, tf, index=None, *, images=None):
    """area = 父区域 (x, y, w, h)，在参考坐标系中；tf 把参考坐标映射到画布；index 是 grid 当前项的 1 起序号（bind: "#"）；
    images 解析 image 元素的路径。"""
    ax, ay, aw, ah = area
    for el in elements:
        if "when" in el and is_empty(lookup(scope, el["when"])):
            continue
        x, y, w, h = resolve_box(el["box"], aw, ah)
        ref_box = (ax + x, ay + y, w, h)
        kind = el["type"]
        if kind == "shape":
            key = lookup(scope, el["by"]) if "by" in el else None
            add_shape(slide, style, tf(ref_box), _mapped(el.get("fill_map"), key) or el["fill"],
                      line=_mapped(el.get("line_map"), key) or el.get("line"), kind=el.get("shape", "rect"),
                      shadow=el.get("shadow", False), radius=el.get("radius"), dash=el.get("dash", False))
        elif kind == "text" and "runs" in el:
            # 一个段落里连排多个 run；run 没写 style / color 就用元素的（有 by + color_map 时按值取色）；取不到值的 run 跳过，全部取不到则不画
            color = (_mapped(el.get("color_map"), lookup(scope, el["by"])) if "by" in el else None) or el.get("color")
            runs = [(text, run.get("style", el["style"]), run.get("color", color))
                    for run in el["runs"] if (text := _text_of(scope, run, index)) is not None]
            if runs:
                add_runs(slide, style, tf(ref_box), runs, align=el.get("align", "left"), valign=el.get("valign", "top"))
        elif kind == "text":
            text = _text_of(scope, el, index)
            if text is not None:
                key = lookup(scope, el["by"]) if "by" in el else text
                add_text(slide, style, tf(ref_box), text, el["style"], color=_mapped(el.get("color_map"), key) or el.get("color"),
                         align=el.get("align", "left"), valign=el.get("valign", "top"))
        elif kind == "bullets":
            items = lookup(scope, el["bind"])
            if items:
                color = _mapped(el.get("color_map"), lookup(scope, el["by"])) if "by" in el else None
                add_bullets(slide, style, tf(ref_box), items, el["style"], color=color or el.get("color"), marker=el.get("marker", True),
                            children=el.get("children"), group_style=el.get("group_style"))
        elif kind == "grid":
            items = lookup(scope, el["bind"]) or []
            weights = None
            if "weight" in el:  # 数值按数值、列表按条目数，没填的项按 1 算
                weights = [float(len(w) if isinstance(w, list) else w) if w is not None else 1.0
                           for w in (lookup(item, el["weight"]) for item in items)]
            cells = grid_cells(ref_box, el["columns"], el.get("gap", 0), len(items), el.get("row_height"), weights)
            for i, (item, cell) in enumerate(zip(items, cells), 1):
                _render_elements(slide, style, el["item"], cell, item, tf, i, images=images)
        elif kind == "image":
            path = lookup(scope, el["bind"])
            if is_empty(path):  # 没给图片：画一块带引导文字的占位底（模板预览就是这个）
                add_shape(slide, style, tf(ref_box), "surface", text=el["placeholder"], text_style="detail", color="text",
                          align="center", valign="middle", inset=8)
            else:
                add_image(slide, tf(ref_box), images.resolve(path), fit=el.get("fit", "contain"))
        elif kind == "component":
            if el["component"] not in COMPONENTS:
                raise PptError(f"未知组件: {el['component']}")
            data = lookup(scope, el["bind"])
            if data is not None:
                component, box = COMPONENTS[el["component"]], tf(ref_box)
                # prepare_pages 已按参考尺寸检查过；这里按实际画布尺寸兜底，免得画出损坏的形状
                problems = [f"组件 {el['component']}  {pointer(path)}  {msg}" for path, msg in component.check(data, box, el.get("props", {}))]
                if problems:
                    raise PptError("\n".join(problems))
                component.render(slide, style, box, data, el.get("props", {}))
        else:
            raise PptError(f"不支持的布局元素类型: {kind}")
