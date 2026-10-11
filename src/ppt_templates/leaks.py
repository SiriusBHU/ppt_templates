"""ppt leaks：模板 example.yaml 和 template.yaml 里和样板原文重合的片段。

example 只能写引导文字，模板的名称、structure、描述、固定文字也不能带样板的业务内容。从真实页做模板时最容易出的事，
是把样板内容（哪怕换个说法）写进示例，或者把样板的业务词写死在模板里；
这里按 tools/ppt-skeleton 的 check 同一思路找重合：连续 4 个汉字、连续 8 个字母（不分大小写）、3 位以上的同一个数。
"""

import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .errors import PptError

CJK_RUN, ASCII_RUN, DIGIT_RUN = 4, 8, 3
_CJK = re.compile(rf"[一-鿿]{{{CJK_RUN}}}")
_ASCII = re.compile(rf"[A-Za-z]{{{ASCII_RUN}}}")
_DIGITS = re.compile(rf"\d{{{DIGIT_RUN},}}")


def sample_texts(pptx: Path, slides: list[int] | None = None) -> list[str]:
    """样板选定页（默认全部）的所有文字：文本框、组合里的形状、表格单元格、备注。"""
    if not pptx.is_file():
        raise PptError(f"样板 {pptx} 不存在")
    prs = Presentation(pptx)
    total = len(prs.slides)
    for n in slides or []:
        if not 1 <= n <= total:
            raise PptError(f"样板只有 {total} 页，没有第 {n} 页")
    texts = []
    for n, slide in enumerate(prs.slides, 1):
        if slides and n not in slides:
            continue
        texts += _shape_texts(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            texts.append(slide.notes_slide.notes_text_frame.text)
    return texts


def _shape_texts(shapes) -> list[str]:
    texts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts += _shape_texts(shape.shapes)
        elif shape.has_table:
            texts += [cell.text for row in shape.table.rows for cell in row.cells]
        elif shape.has_text_frame:
            texts.append(shape.text_frame.text)
    return texts


# template.yaml 里这些键的值是标识、token 或路径，不是给人看的文字
TEMPLATE_KEYS_SKIPPED = {"id", "roles", "also_fits", "density", "level", "frame", "type", "box", "bind", "by", "when", "weight",
                         "style", "color", "fill", "line", "shape", "component", "fit", "align", "valign", "required",
                         "pattern", "format", "$ref", "$schema"}


def _strings(data, pointer="", skip=frozenset()):
    if isinstance(data, str):
        yield pointer or "/", data
    elif isinstance(data, dict):
        for key, value in data.items():
            if key not in skip:
                yield from _strings(value, f"{pointer}/{key}", skip)
    elif isinstance(data, list):
        for i, value in enumerate(data):
            yield from _strings(value, f"{pointer}/{i}", skip)


def find_leaks(example, texts: list[str], allow=(), *, prefix="", skip=frozenset(), numbers=True) -> list[str]:
    """example 里每个字符串和原文重合的最长片段，形如 `<prefix>/items/0/name  「片段」与样板原文重合`；skip 里的键不查，
    numbers=False 时不查数字（template.yaml 里的数字是字数上限、坐标这类参数）。
    空白不计；allow 里的词先从原文里去掉，不再算重合。各段原文用换行隔开，片段不会跨两段拼出来。"""
    source = "\n".join(re.sub(r"\s+", "", t) for t in texts)
    for word in allow:
        source = source.replace(re.sub(r"\s+", "", word), "\n")
    lower = source.lower()
    known = set(_DIGITS.findall(source)) if numbers else set()
    leaks = []
    for pointer, value in _strings(example, prefix, skip):
        s = re.sub(r"\s+", "", value)
        hit = [False] * len(s)
        for i in range(len(s)):
            for size, pattern, hay in ((CJK_RUN, _CJK, source), (ASCII_RUN, _ASCII, lower)):
                window = s[i:i + size]
                if pattern.fullmatch(window) and window.lower() in hay:
                    hit[i:i + size] = [True] * size
        for m in _DIGITS.finditer(s):
            if m.group() in known:
                hit[m.start():m.end()] = [True] * len(m.group())
        start = None
        for i, flag in enumerate(hit + [False]):
            if flag and start is None:
                start = i
            elif not flag and start is not None:
                leaks.append(f"{pointer}  「{s[start:i]}」与样板原文重合")
                start = None
    return leaks
