"""表格 / 对比矩阵：原生 PPT 表格。表头深色底，正文行浅色交替，可高亮一列（如推荐方案）。
props 的 look: plain 照原页画成无底色、只有横线的表格；text_style / align 调字号和对齐。"""

from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu

from ..draw import ALIGN, add_run, emu_box, rgb
from ..geometry import px_to_emu
from .tokens import COLOR

NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 去掉内置表格样式，颜色全部由风格 token 决定
CELL_PAD = 12

SCHEMA = {
    "type": "object",
    "required": ["columns", "rows"],
    "properties": {
        "columns": {"type": "array", "minItems": 1, "items": {"type": "string"}},
        "rows": {"type": "array", "minItems": 1, "items": {"type": "array", "items": {"type": "string"}}},
        "highlight_column": {"type": "integer", "minimum": 0},
        "widths": {"type": "array", "items": {"type": "number", "exclusiveMinimum": 0}},   # 各列的相对宽度，跟着内容定；不写就按 props
    },
}

PROPS_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "first_column_width": {"type": "number", "exclusiveMinimum": 0},
        "look": {"enum": ["filled", "plain"]},          # filled（默认）：深色表头 + 浅色交替行；plain：白底，表头上下和每行下方一条横线
        "text_style": {"enum": ["body", "detail", "caption"]},   # 单元格文字，默认 body（12 号）
        "align": {"enum": ["left", "center"]},           # 正文格的对齐，默认第一列左对齐、其余居中
        "rule": COLOR,                                    # plain 的横线颜色，默认 text
    },
}


def check(data, box, props=None) -> list[tuple[tuple, str]]:
    """每行格数、widths 的个数都必须等于表头列数。"""
    n = len(data["columns"])
    problems = [(("rows", i), f"第 {i + 1} 行有 {len(row)} 格，表头有 {n} 列")
                for i, row in enumerate(data["rows"]) if len(row) != n]
    if "widths" in data and len(data["widths"]) != n:
        problems.append((("widths",), f"widths 有 {len(data['widths'])} 项，表头有 {n} 列"))
    return problems


def _cell(cell, style, text, text_style, fill, *, color=None, bold=None, align="left", pad_y=CELL_PAD / 2):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(style, fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Emu(px_to_emu(CELL_PAD))
    cell.margin_top = cell.margin_bottom = Emu(px_to_emu(pad_y))
    cell.text_frame.word_wrap = True
    p = cell.text_frame.paragraphs[0]
    p.alignment = ALIGN[align]
    add_run(p, style, text, text_style, color, bold)


def _rule(cell, style, color: str, side: str) -> None:
    """单元格上边（T）或下边（B）画一条 0.75pt 的实线。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn(f"a:ln{side}")):
        tc_pr.remove(old)
    ln = tc_pr.makeelement(qn(f"a:ln{side}"), {"w": "9525", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    fill = ln.makeelement(qn("a:solidFill"), {})
    fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(rgb(style, color))}))
    ln.append(fill)
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
    tc_pr.insert(0 if side == "T" else len(tc_pr.findall(qn("a:lnL")) + tc_pr.findall(qn("a:lnR")) + tc_pr.findall(qn("a:lnT"))), ln)


def _widths(total: float, n: int, first) -> list[float]:
    if first and n > 1:
        return [first] + [(total - first) / (n - 1)] * (n - 1)
    return [total / n] * n


def render(slide, style, box, data, props):
    columns, rows = data["columns"], data["rows"]
    hot = data.get("highlight_column")
    _, _, w, h = box
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), *emu_box(box))
    table = shape.table
    table.first_row = table.horz_banding = False
    table._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_NO_GRID
    ratios = data.get("widths")
    widths = [w * r / sum(ratios) for r in ratios] if ratios else _widths(w, len(columns), props.get("first_column_width"))
    for col, cw in zip(table.columns, widths):
        col.width = Emu(px_to_emu(cw))
    for row in table.rows:
        row.height = Emu(px_to_emu(h / (len(rows) + 1)))

    if props.get("look") == "plain":
        _render_plain(table, style, columns, rows, props)
        return shape
    for c, name in enumerate(columns):
        _cell(table.cell(0, c), style, name, "body_strong", "highlight_strong" if c == hot else "primary",
              color="on_primary", align="left" if c == 0 and c != hot else "center")
    for r, row in enumerate(rows, 1):
        fill = "surface" if r % 2 == 0 else "bg"
        for c, text in enumerate(row):
            if c == hot:
                _cell(table.cell(r, c), style, text, "body", "highlight", color="highlight_strong", bold=True, align="center")
            elif c == 0:
                _cell(table.cell(r, c), style, text, "body_strong", fill)
            else:
                _cell(table.cell(r, c), style, text, "body", fill, align="center")
    return shape


def _render_plain(table, style, columns, rows, props):
    """白底表格：表头和正文同色同字号（表头加粗），表头上方、表头下方和每行下方各一条横线。"""
    text_style, rule = props.get("text_style", "body"), props.get("rule", "text")
    bold_style = "body_strong" if text_style == "body" else text_style
    for r, row in enumerate([columns, *rows]):
        for c, text in enumerate(row):
            align = props.get("align") or ("left" if c == 0 else "center")
            cell = table.cell(r, c)
            _cell(cell, style, text, bold_style if r == 0 else text_style, "bg", bold=True if r == 0 else None, align=align,
                  pad_y=CELL_PAD / 4)
            if r == 0:
                _rule(cell, style, rule, "T")
            _rule(cell, style, rule, "B")
