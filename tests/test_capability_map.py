import pytest
import yaml
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from factories import deck_of
from ppt_templates.components.capability_map import _columns, _rows
from ppt_templates.errors import PptError
from ppt_templates.loader import load_example
from ppt_templates.render import render_deck
from ppt_templates.slots import prepare_pages
from ppt_templates.validate import validate_repo

CAP = "product.roadmap-capmap"
REAL = "product.roadmap-capability-map"
PX = 9525


def capmap_template():
    """product.roadmap-capmap：能力地图（capability-map 组件）；3 个分组共 7 行：
    第 1 组 3 个条目 + 1 个 wide 条目，第 2 组 1 个分组框（2 × 2）+ 3 个条目，第 3 组 4 个条目（默认 2 × 2）；两种图例。"""
    template = {
        "id": CAP,
        "name": "能力地图",
        "roles": ["roadmap"],
        "when_to_use": "测试用：能力地图",
        "frame": "standard",
        "slots": {
            "type": "object",
            "required": ["title", "map"],
            "properties": {"title": {"type": "string", "maxLength": 14}, "map": {"type": "object"}},  # 放宽槽位，错误只来自组件
        },
        "layout": [{"type": "component", "component": "capability-map", "box": [0, 0, 576, 496], "bind": "map"}],
    }
    example = {
        "title": "能力地图",
        "map": {
            "title": "总标题",
            "groups": [
                {"name": "甲组", "items": [{"name": "能力一"}, {"name": "能力二", "tag": "近期"}, {"name": "能力三"}, {"name": "通栏条目", "wide": True}]},
                {"name": "乙组", "items": [
                    {"name": "分组框", "items": [{"name": "子一"}, {"name": "子二", "tag": "远期"}, {"name": "子三"}, {"name": "子四"}]},
                    {"name": "能力四"}, {"name": "能力五", "tag": "远期"}, {"name": "能力六"}]},
                {"name": "丙组", "items": [{"name": "能力七"}, {"name": "能力八"}, {"name": "能力九"}, {"name": "能力十", "tag": "远期"}]},
            ],
            "legend": [{"tag": "近期", "label": "近期落地"}, {"tag": "远期", "label": "远期预埋"}],
        },
    }
    return template, example


def _renderer(make_repo):
    """返回 render(example=None) -> 第 1 页的全部形状；同一个测试里 make_repo 只能调用一次。"""
    template, default = capmap_template()
    repo = make_repo({CAP: (template, default)})
    return lambda example=None: list(render_deck(deck_of(CAP, example or default), repo).slides[0].shapes)


def _shape(shapes, text):
    return next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == text)


def _box(shapes, text):
    sh = _shape(shapes, text)
    return tuple(round(v / PX, 1) for v in (sh.left, sh.top, sh.width, sh.height))


def _rgb(shapes, text):
    return str(_shape(shapes, text).fill.fore_color.rgb)


def test_example_is_valid(make_repo):
    template, example = capmap_template()
    assert validate_repo(make_repo({CAP: (template, example)})).errors == []


def test_the_real_template_example_passes_with_slot_defaults(root):
    """槽位里 wide 有 default: false：补全默认值后，分组框条目上也会带一个 wide: false，不能因此报错。"""
    assert [e for e in validate_repo(root).errors if REAL in e] == []
    assert prepare_pages(deck_of(REAL, load_example(root, REAL)), root)


def test_row_heights_share_the_spare_height_and_fill_down_to_the_legend(make_repo):
    shapes = _renderer(make_repo)()
    # box 原点 = body 左上角 (32, 128)；总标题 24 + 8，图例 24 + 8；中间 432 px 分给 3 个分组：
    # 固定高度 204（组名 / 内边距 / 行距），7 行各 ≥ 28，剩下 32 / 7 平分 → 每行 32.6
    assert _box(shapes, "总标题") == (32, 128, 576, 24)
    item = _box(shapes, "能力一")
    assert item == (40, 192, 181.3, 32.6)                       # 面板 y = 160，组名 32 之下
    heights = {_box(shapes, name)[3] for name in ("能力一", "通栏条目", "子一", "子三", "能力四", "能力七", "能力九")}
    assert heights == {32.6}
    legend = _box(shapes, "近期落地")
    assert legend[1] + legend[3] == 128 + 496                    # 图例在 box 的底边
    last_row = _box(shapes, "能力九")
    assert last_row[1] + last_row[3] + 8 + 8 + 24 == pytest.approx(128 + 496, abs=0.5)   # 末行下沿 + 面板内边距 + 间距 + 图例


def test_rows_flow_by_columns_and_wide_or_group_items_take_a_full_row(make_repo):
    shapes = _renderer(make_repo)()
    one, two, three = _box(shapes, "能力一"), _box(shapes, "能力二"), _box(shapes, "能力三")
    assert one[1] == two[1] == three[1] and one[2] == two[2] == three[2] == 181.3     # 3 个一行，平分 560 - 2 × 8
    assert two[0] == pytest.approx(one[0] + one[2] + 8, abs=0.1)
    wide = _box(shapes, "通栏条目")
    assert wide[0] == 40 and wide[2] == 560 and wide[1] == pytest.approx(one[1] + one[3] + 8, abs=0.1)
    # 默认列数：第 3 组 4 个条目排 2 × 2
    seven, eight, nine = _box(shapes, "能力七"), _box(shapes, "能力八"), _box(shapes, "能力九")
    assert seven[2] == eight[2] == 276 and seven[1] == eight[1] and nine[1] > seven[1] and nine[0] == seven[0]


def test_a_group_box_has_its_own_title_and_inner_grid(make_repo):
    shapes = _renderer(make_repo)()
    title, first, second = _box(shapes, "分组框"), _box(shapes, "子一"), _box(shapes, "子二")
    assert title[2] == 544 and title[0] == 48                                      # 框内再缩进 8
    assert first[0] == 48 and first[2] == second[2] == 268 and second[0] == pytest.approx(48 + 268 + 8, abs=0.1)   # 4 个排 2 × 2
    assert first[1] == pytest.approx(title[1] - 4 + 28, abs=0.1)                    # 标题 4 + 20 + 4 之下
    follow = _box(shapes, "能力四")
    assert follow[0] == 40 and follow[1] > _box(shapes, "子三")[1] + _box(shapes, "子三")[3]   # 分组框之后另起一行，回到面板的宽度


def test_tags_pick_the_fill_and_untagged_items_use_the_first_one(make_repo):
    render = _renderer(make_repo)
    shapes = render()
    assert _rgb(shapes, "能力一") == "FFFFFF" and _rgb(shapes, "能力二") == "FFFFFF"      # 没写 tag / 第 1 种
    assert _rgb(shapes, "子二") == "FFF1E6" and _rgb(shapes, "能力十") == "FFF1E6"        # 第 2 种
    hot = _shape(shapes, "能力十")
    assert str(hot.line.color.rgb) == "D1D5DB" and hot.line.width == 9525
    run = hot.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and run.font.size.pt == 12 and str(run.font.color.rgb) == "000000"   # 文字仍是黑色、加粗
    assert hot.text_frame.paragraphs[0].alignment == 2                                        # PP_ALIGN.CENTER


def test_legend_swatches_follow_the_two_fills_and_are_left_aligned(make_repo):
    render = _renderer(make_repo)
    shapes = render()
    first, second = _box(shapes, "近期落地"), _box(shapes, "远期预埋")
    swatches = [sh for sh in shapes if not sh.has_text_frame or not sh.text_frame.text]
    left = sorted((sh for sh in swatches if round(sh.width / PX) == 14), key=lambda sh: sh.left)
    assert [str(sh.fill.fore_color.rgb) for sh in left] == ["FFFFFF", "FFF1E6"]
    assert round(left[0].left / PX) == 32 and first[0] == 32 + 14 + 6 and second[0] > first[0] + first[2]
    _, example = capmap_template()
    del example["map"]["legend"]
    for group in example["map"]["groups"]:
        for item in group["items"]:
            item.pop("tag", None)
            for sub in item.get("items", []):
                sub.pop("tag", None)
    assert {"近期落地", "远期预埋"}.isdisjoint(sh.text_frame.text for sh in render(example) if sh.has_text_frame)


def test_group_panels_titles_and_the_nested_box(make_repo):
    shapes = _renderer(make_repo)()
    title = _shape(shapes, "甲组")
    run = title.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and run.font.size.pt == 14 and str(run.font.color.rgb) == "283593"       # 组名：14 号加粗、靛蓝
    sub = _shape(shapes, "分组框").text_frame.paragraphs[0].runs[0]
    assert sub.font.bold and sub.font.size.pt == 12 and str(sub.font.color.rgb) == "283593"
    top = _shape(shapes, "总标题").text_frame.paragraphs[0].runs[0]
    assert top.font.bold and top.font.size.pt == 14 and str(top.font.color.rgb) == "000000"
    fills = [str(sh.fill.fore_color.rgb) for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not sh.text_frame.text
             and sh.line.fill.type == MSO_FILL.BACKGROUND]
    assert fills.count("F3F5F8") == 3 and fills.count("FFFFFF") == 1                            # 3 个灰底分组面板、1 个白底分组框，都不描边


def test_rows_stop_growing_at_the_cap_and_the_stack_stays_at_the_top(make_repo):
    _, example = capmap_template()
    example["map"]["groups"] = example["map"]["groups"][:1]
    example["map"]["groups"][0]["items"] = [{"name": "唯一条目"}]
    shapes = _renderer(make_repo)(example)
    assert _box(shapes, "唯一条目")[3] == 56
    legend = _box(shapes, "近期落地")
    assert legend[1] + legend[3] == 128 + 496                                                    # 图例仍贴着底边


def test_without_a_title_the_groups_start_at_the_top(make_repo):
    _, example = capmap_template()
    del example["map"]["title"]
    shapes = _renderer(make_repo)(example)
    assert _box(shapes, "甲组")[1] == 128 + 4 and _box(shapes, "能力一")[1] == 128 + 32


def test_render_example_reopens(make_repo, tmp_path):
    template, example = capmap_template()
    path = tmp_path / "cap.pptx"
    render_deck(deck_of(CAP, example), make_repo({CAP: (template, example)})).save(path)
    texts = {sh.text_frame.text for sh in Presentation(path).slides[0].shapes if sh.has_text_frame}
    assert {"总标题", "甲组", "乙组", "丙组", "分组框", "能力一", "子四", "能力十", "近期落地", "远期预埋"} <= texts


def _prepare_error(repo, example) -> str:
    with pytest.raises(PptError) as exc:
        prepare_pages(deck_of(CAP, example), repo)
    return str(exc.value)


def _group_item(item):
    return lambda m: m["groups"][1]["items"][0].update(item)


@pytest.mark.parametrize("mutate, where, message", [
    (lambda m: m["groups"][0]["items"][0].update(tag="未知"), "groups/0/items/0/tag", "tag「未知」不在 legend 里，可选：近期、远期"),
    (lambda m: m["groups"][1]["items"][0]["items"][1].update(tag="未知"),
     "groups/1/items/0/items/1/tag", "tag「未知」不在 legend 里，可选：近期、远期"),
    (lambda m: m["legend"][1].update(tag="近期"), "legend/1/tag", "tag「近期」重复"),
    (_group_item({"tag": "近期"}), "groups/1/items/0/tag", "有 items 的条目是分组框，不能再写 tag"),
    (_group_item({"wide": True}), "groups/1/items/0/wide", "有 items 的条目是分组框，不能再写 wide"),
])
def test_tag_semantics_are_checked_with_page_and_pointer(make_repo, mutate, where, message):
    template, example = capmap_template()
    mutate(example["map"])
    repo = make_repo({CAP: (template, example)})
    assert f"第 1 页 {CAP}  /pages/0/data/map/{where}  {message}" in _prepare_error(repo, example)
    assert any(f"example.yaml  /map/{where}  " in e and message in e and e.endswith("（组件 capability-map）")
               for e in validate_repo(repo).errors)


def test_a_tag_needs_a_legend(make_repo):
    template, example = capmap_template()
    del example["map"]["legend"]
    repo = make_repo({CAP: (template, example)})
    assert "/map/groups/0/items/1/tag  tag「近期」不在 legend 里，请先在 legend 里定义" in _prepare_error(repo, example)


def _very_wide_legend(m):
    m["legend"][0]["label"] = "例" * 40


@pytest.mark.parametrize("mutate, where, message", [
    (lambda m: m["groups"][0]["items"][0].update(name="Prefill-Decode-Overlap"),
     "groups/0/items/0/name", "英文单词「Prefill-Decode-Overlap」比 181 px 宽的框还宽，会被拦腰折断"),
    (lambda m: m["groups"][1]["items"][0]["items"][0].update(name="Prefill-Decode-Overlap-Pipeline"),
     "groups/1/items/0/items/0/name", "英文单词「Prefill-Decode-Overlap-Pipeline」比 268 px 宽的框还宽"),
    (lambda m: m["groups"][0].update(name="宽" * 40), "groups/0/name", "在 560 px 宽的框里估算要 2 行，只放得下 1 行"),
    (lambda m: m["groups"][1]["items"][0].update(name="宽" * 40), "groups/1/items/0/name", "在 544 px 宽的框里估算要 2 行，只放得下 1 行"),
    (lambda m: m.update(title="题" * 31), "title", "在 576 px 宽的框里估算要 2 行，只放得下 1 行"),
    (_very_wide_legend, "legend", "图例估算要 "),
])
def test_text_that_cannot_fit_is_rejected_with_page_and_pointer(make_repo, mutate, where, message):
    template, example = capmap_template()
    mutate(example["map"])
    repo = make_repo({CAP: (template, example)})
    assert f"第 1 页 {CAP}  /pages/0/data/map/{where}  " in _prepare_error(repo, example)
    assert message in _prepare_error(repo, example)
    assert any(f"example.yaml  /map/{where}  " in e and message in e for e in validate_repo(repo).errors)


def test_too_many_rows_for_the_box_height_are_rejected_at_the_overflowing_group(make_repo):
    template, example = capmap_template()
    six = [{"name": f"能力{i}"} for i in range(6)]
    example["map"]["groups"] = [{"name": f"{n}组", "items": [dict(item) for item in six]} for n in "甲乙丙丁"]
    repo = make_repo({CAP: (template, example)})
    message = _prepare_error(repo, example)
    assert "/pages/0/data/map/groups/3  第 4 组「丁组」放不下：各组自上而下排，到它的下沿估算要 " in message
    assert "估算要 472 px 高，只放得下 464 px" in message                                       # 464 = 496 - 图例 32；总标题也算在内
    example["map"]["groups"] = example["map"]["groups"][:3]
    assert prepare_pages(deck_of(CAP, example), repo)                                            # 3 组 6 行放得下


def test_two_line_items_get_a_taller_row_and_still_fit(make_repo):
    template, example = capmap_template()
    example["map"]["groups"] = example["map"]["groups"][:1]
    example["map"]["groups"][0]["items"] = [{"name": "较长的条目名称要折行"}, {"name": "短"}, {"name": "短"}]
    example["map"]["groups"][0]["columns"] = 3
    template["layout"][0]["box"] = [0, 0, 300, 496]                                              # 窄一点：每个条目只有 ~ 5 个字宽
    repo = make_repo({CAP: (template, example)})
    shapes = list(render_deck(deck_of(CAP, example), repo).slides[0].shapes)
    assert _box(shapes, "较长的条目名称要折行")[3] == _box(shapes, "短")[3] == 56                  # 一行放不下，需要 2 × 20 + 6 = 46，富余后封顶 56


@pytest.mark.parametrize("mutate, expected", [
    (lambda m: m.update(legend=[{"tag": str(i), "label": "例"} for i in range(3)]), "/legend  共 3 项，超过上限 2"),
    (lambda m: m["legend"][0].pop("label"), "/legend/0  缺少必填字段 label"),
    (lambda m: m["groups"].clear(), "/groups  共 0 项，低于下限 1"),
    (lambda m: m["groups"][0].update(items=[]), "/groups/0/items  共 0 项，低于下限 1"),
    (lambda m: m["groups"][0].pop("name"), "/groups/0  缺少必填字段 name"),
    (lambda m: m["groups"][0].update(columns=0), "/groups/0/columns  取值 0，低于下限 1"),
    (lambda m: m["groups"][1]["items"][0]["items"][0].update(items=[{"name": "再嵌一层"}]),
     "/groups/1/items/0/items/0  不允许的字段 items"),
    (lambda m: m.pop("groups"), "  缺少必填字段 groups"),
])
def test_component_schema_errors_are_prefixed_with_the_page_pointer(make_repo, mutate, expected):
    template, example = capmap_template()
    mutate(example["map"])
    repo = make_repo({CAP: (template, example)})
    message = _prepare_error(repo, example)
    assert f"第 1 页 {CAP}  /pages/0/data/map{expected}" in message and "（组件 capability-map）" in message


def test_render_rechecks_against_the_actual_body_area(make_repo):
    template, example = capmap_template()
    repo = make_repo({CAP: (template, example)})
    style_path = repo / "styles" / "default" / "style.yaml"
    style = yaml.safe_load(style_path.read_text(encoding="utf-8"))
    style["frame"]["body"]["box"] = [32, 128, 400, 560]    # body 区比参考尺寸窄得多
    style_path.write_text(yaml.safe_dump(style, allow_unicode=True), encoding="utf-8")
    with pytest.raises(PptError, match="组件 capability-map  /groups/"):
        render_deck(deck_of(CAP, example), repo)


def test_default_columns_and_row_flow():
    plain = {"name": "条目"}
    assert [_columns([plain] * n, None) for n in (1, 2, 3, 4, 5, 6)] == [1, 2, 3, 2, 3, 3]
    assert _columns([plain, plain, {"name": "宽", "wide": True}, {"name": "框", "items": [plain]}], None) == 2   # wide 和分组框不计入
    assert _columns([plain] * 4, 4) == 4                                                                      # 写了 columns 就用它
    assert _rows([plain] * 5, 3) == [(0, 1, 2), (3, 4)]                                                       # 末行不满：照常成一行（渲染时拉宽）
    wide, box = {"name": "宽", "wide": True}, {"name": "框", "items": [plain]}
    assert _rows([plain, plain, wide, plain], 3) == [(0, 1), (2,), (3,)]                                      # wide 截断当前行
    assert _rows([plain, box, plain, plain], 2) == [(0,), (1,), (2, 3)]                                       # 分组框也一样
