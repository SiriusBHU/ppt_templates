import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from factories import deck_of, diagram_template
from ppt_templates.errors import PptError
from ppt_templates.render import render_deck
from ppt_templates.validate import validate_repo


def test_nodes_and_arrow_connectors(make_repo, tmp_path):
    template, example = diagram_template()
    repo = make_repo({"system.solution-flow": (template, example)})
    assert validate_repo(repo).errors == []
    path = tmp_path / "d.pptx"
    render_deck(deck_of("system.solution-flow", example), repo).save(path)
    shapes = list(Presentation(path).slides[0].shapes)
    node = {sh.text_frame.text: sh for sh in shapes if sh.has_text_frame}
    conns = [sh for sh in shapes if sh._element.tag == qn("p:cxnSp")]
    assert len(conns) == 3
    gpu, nvme, store, meta = node["训练进程"], node["本地 NVMe"], node["存储集群"], node["元数据服务"]
    first, _, last = conns
    assert abs(first.begin_x - (gpu.left + gpu.width)) <= 1 and abs(first.end_x - nvme.left) <= 1      # 水平：右 → 左
    assert abs(last.begin_y - (store.top + store.height)) <= 1 and abs(last.end_y - meta.top) <= 1     # 垂直：下 → 上
    assert all(c._element.xpath(".//a:tailEnd/@type") == ["triangle"] for c in conns)
    assert "同步写" in node and "异步下刷" in node
    assert str(nvme.fill.fore_color.rgb) == "FFF1E6"


@pytest.mark.parametrize("mutate, message", [
    (lambda f: f["edges"].append({"from": "gpu", "to": "ghost"}), "框图连线引用了不存在的节点: ghost"),
    (lambda f: f["nodes"].append({"id": "x", "name": "重叠", "pos": [0, 0]}), r"框图节点位置重复: \[0, 0\]"),
    (lambda f: f["nodes"].append({"id": "gpu", "name": "重复", "pos": [1, 0]}), "框图节点 id 重复: gpu"),
])
def test_bad_references_are_rejected(make_repo, mutate, message):
    template, example = diagram_template()
    mutate(example["flow"])
    repo = make_repo({"system.solution-flow": (template, example)})
    with pytest.raises(PptError, match=message):
        render_deck(deck_of("system.solution-flow", example), repo)
