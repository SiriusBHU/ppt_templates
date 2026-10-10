import pytest
import yaml
from pptx import Presentation

from factories import arch_template, deck_of
from ppt_templates.components.partitioned_arch import _inset_x, _lines
from ppt_templates.errors import PptError
from ppt_templates.render import render_deck
from ppt_templates.slots import prepare_pages
from ppt_templates.validate import validate_repo

ARCH = "system.solution-arch"
PX = 9525


def _renderer(make_repo):
    """返回 render(example=None) -> 第 1 页的全部形状；同一个测试里 make_repo 只能调用一次。"""
    template, default = arch_template()
    repo = make_repo({ARCH: (template, default)})
    return lambda example=None: list(render_deck(deck_of(ARCH, example or default), repo).slides[0].shapes)


def _box(shapes, text):
    sh = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == text)
    return tuple(round(v / PX, 1) for v in (sh.left, sh.top, sh.width, sh.height))


def test_example_is_valid(make_repo):
    template, example = arch_template()
    assert validate_repo(make_repo({ARCH: (template, example)})).errors == []


def test_partition_width_follows_module_count(make_repo):
    shapes = _renderer(make_repo)()
    cpu, npu = _box(shapes, "Tool Use (via CPU)"), _box(shapes, "Agent Reasoning (via NPU)")
    # 固定开销（内边距 16 + 层名 56 + 间距 6 = 78）之外的宽度按每层最多的模块数 2 : 3 分；box 原点 = body 左上角 (32, 128)
    assert cpu == (32, 128, 275.6, 28)
    assert npu == (317.6, 128, 378.4, 28)
    assert 275.6 + 10 + 378.4 == 664


def test_layers_modules_and_shared_rows_are_placed_inside_the_partitions(make_repo):
    shapes = _renderer(make_repo)()
    assert _box(shapes, "工具层") == (40, 168, 56, 176)                      # 层名：分区主体左上角 + 内边距 8
    assert _box(shapes, "代码沙箱") == (102, 168, 94.8, 176)                 # 模块区平分剩余宽度，间距 8
    assert _box(shapes, "搜索服务") == (204.8, 168, 94.8, 176)
    assert _box(shapes, "内存池")[0] == 40                                  # 共享层横跨全宽，层名与分区的对齐
    hbm, cxl = _box(shapes, "HBM 池"), _box(shapes, "CXL 池")
    assert hbm[1] == cxl[1] > _box(shapes, "工具层")[1] + 176                # 共享层在分区之下
    assert hbm[0] + hbm[2] + 8 == cxl[0] and cxl[0] + cxl[2] == 32 + 664 - 8 and hbm[1] == 368


def test_highlight_style_and_legend(make_repo):
    render = _renderer(make_repo)
    shapes = render()
    hot = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == "CXL 池")
    assert str(hot.fill.fore_color.rgb) == "FFF1E6" and str(hot.line.color.rgb) == "C00000" and hot.line.width == 19050
    assert "重点模块" in {sh.text_frame.text for sh in shapes if sh.has_text_frame}
    _, example = arch_template()
    for part in example["arch"]["partitions"]:
        for layer in part["layers"]:
            for module in layer["modules"]:
                module["highlight"] = False
    example["arch"]["shared"][0]["modules"][1]["highlight"] = False
    assert "重点模块" not in {sh.text_frame.text for sh in render(example) if sh.has_text_frame}


def test_header_tab_is_a_bordered_box_with_centered_bold_text(make_repo):
    shapes = _renderer(make_repo)()
    tab = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == "Tool Use (via CPU)")
    run = tab.text_frame.paragraphs[0].runs[0]
    assert str(tab.fill.fore_color.rgb) == "FFFFFF" and str(tab.line.color.rgb) == "283593"
    assert run.font.bold and run.font.size.pt == 12 and tab.text_frame.paragraphs[0].alignment == 2  # PP_ALIGN.CENTER


def test_fewer_layers_stretch_to_the_same_partition_height(make_repo):
    _, example = arch_template()
    example["arch"]["partitions"][0]["layers"].append({"name": "编排层", "modules": [{"name": "任务编排"}]})
    shapes = _renderer(make_repo)(example)
    # 左分区 2 层、右分区 1 层：右分区这一层占满主体，左分区两层各占一半
    assert _box(shapes, "推理层")[3] > _box(shapes, "工具层")[3] > 0
    top, bottom = _box(shapes, "工具层"), _box(shapes, "编排层")
    assert bottom[1] == pytest.approx(top[1] + top[3] + 8, abs=0.2) and bottom[3] == top[3]


def _prepare_error(repo, example) -> str:
    with pytest.raises(PptError) as exc:
        prepare_pages(deck_of(ARCH, example), repo)
    return str(exc.value)


@pytest.mark.parametrize("mutate, where, message", [
    (lambda a: a["partitions"][0]["layers"][0]["modules"][0].update(name="长" * 60),
     "partitions/0/layers/0/modules/0/name", "「" + "长" * 60 + "」在 95×176 px 的框里估算要 12 行，只放得下 8 行"),
    (lambda a: a["partitions"][0].update(name="宽" * 20), "partitions/0/name", "估算要 2 行，只放得下 1 行"),
    (lambda a: a["partitions"][0]["layers"][0]["modules"][1].update(name="Prefill-Decode"),
     "partitions/0/layers/0/modules/1/name", "英文单词「Prefill-Decode」比 95 px 宽的框还宽，会被拦腰折断"),
])
def test_text_that_cannot_fit_is_rejected_with_page_and_pointer(make_repo, mutate, where, message):
    template, example = arch_template()
    template["slots"]["properties"]["arch"] = {"type": "object"}  # 放宽槽位的长度上限，让错误只来自组件的尺寸检查
    mutate(example["arch"])
    repo = make_repo({ARCH: (template, example)})
    assert f"第 1 页 {ARCH}  /pages/0/data/arch/{where}  " in _prepare_error(repo, example)
    assert message in _prepare_error(repo, example)
    assert any(f"example.yaml  /arch/{where}  " in e and message in e for e in validate_repo(repo).errors)


def test_rows_get_shorter_as_layers_are_added(make_repo):
    template, example = arch_template()
    template["slots"]["properties"]["arch"] = {"type": "object"}
    example["arch"]["partitions"][1]["name"] = "NPU"  # 只有 1 个模块的分区很窄，放不下长页签
    example["arch"]["partitions"][1]["layers"] = [{"name": f"层{i}", "modules": [{"name": "模块"}]} for i in range(4)]
    example["arch"]["shared"] = [{"name": "共享一", "modules": [{"name": "甲"}]}, {"name": "共享二", "modules": [{"name": "乙"}]}]
    repo = make_repo({ARCH: (template, example)})
    assert prepare_pages(deck_of(ARCH, example), repo)  # 6 行、每行约 58px：放得下 2 行文字
    example["arch"]["partitions"][1]["layers"][0]["name"] = "层" * 9  # 层名标签每行 3 字 → 3 行，放不下
    assert "/arch/partitions/1/layers/0/name  「层层层层层层层层层」在 56×58 px 的框里估算要 3 行，只放得下 2 行" in _prepare_error(repo, example)


@pytest.mark.parametrize("mutate, expected", [
    (lambda a: a["partitions"][0].update(layers=[]), "/partitions/0/layers  共 0 项，低于下限 1"),
    (lambda a: a["partitions"][0]["layers"][0].update(modules=[]), "/partitions/0/layers/0/modules  共 0 项，低于下限 1"),
    (lambda a: a["partitions"][1].pop("name"), "/partitions/1  缺少必填字段 name"),
])
def test_component_schema_errors_are_prefixed_with_the_page_pointer(make_repo, mutate, expected):
    template, example = arch_template()
    template["slots"]["properties"]["arch"] = {"type": "object"}  # 放宽槽位，让错误只来自组件的 schema
    mutate(example["arch"])
    repo = make_repo({ARCH: (template, example)})
    assert f"第 1 页 {ARCH}  /pages/0/data/arch{expected}（组件 partitioned-arch）" in _prepare_error(repo, example)


def test_render_rechecks_text_against_the_actual_body_area(make_repo):
    template, example = arch_template()
    repo = make_repo({ARCH: (template, example)})
    style_path = repo / "styles" / "default" / "style.yaml"
    style = yaml.safe_load(style_path.read_text(encoding="utf-8"))
    style["frame"]["body"]["box"] = [32, 128, 400, 560]    # body 区比参考尺寸窄得多
    style_path.write_text(yaml.safe_dump(style, allow_unicode=True), encoding="utf-8")
    with pytest.raises(PptError, match="组件 partitioned-arch  /partitions/"):
        render_deck(deck_of(ARCH, example), repo)


def test_line_estimate_and_wrap_inset():
    assert _lines("代码沙箱", 69) == 1 and _lines("稀疏注意力", 69) == 2 and _lines("长" * 60, 87) == 12
    assert _lines("Agent Reasoning (via NPU)", 268) == 1
    # 5 个字在每行放 4 字的框里折成 4+1：把内边距加大到每行放 3 字，折成 3+2
    assert _inset_x("module", "稀疏注意力", 77) == (77 - 3 * 16 - 4) / 2
    assert _inset_x("module", "代码沙箱", 77) == 4            # 一行放得下：默认内边距
    assert _inset_x("module", "Agent 框架", 40) == 4          # 含英文：不干预
