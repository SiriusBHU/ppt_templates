from pptx import Presentation
from pptx.oxml.ns import qn

from factories import bullets_template, deck_of
from ppt_templates.render import render_deck
from ppt_templates.validate import validate_repo


def test_bullets_render_with_bullet_char_indent_and_spacing(make_repo, tmp_path):
    template, example = bullets_template()
    repo = make_repo({"common.summary-points": (template, example)})
    assert validate_repo(repo).errors == []
    path = tmp_path / "b.pptx"
    render_deck(deck_of("common.summary-points", example), repo).save(path)
    box = next(sh for sh in Presentation(path).slides[0].shapes
               if sh.has_text_frame and sh.text_frame.text.startswith("检查点写入"))
    paras = box.text_frame.paragraphs
    assert [p.text for p in paras] == ["检查点写入不再阻塞训练", "先写本地 NVMe，再异步下刷", "恢复时间缩短70%"]
    em = 16 * 9525  # body 12pt = 16px
    assert [int(p._p.pPr.get("marL")) for p in paras] == [em, 2 * em, em]
    assert all(int(p._p.pPr.get("indent")) == -em for p in paras)
    assert all(p._p.pPr.find(qn("a:buChar")).get("char") == "•" for p in paras)
    assert paras[0]._p.pPr.find(qn("a:spcBef")) is None
    assert paras[1]._p.pPr.find(qn("a:spcBef")) is not None
    assert box._element.xpath(".//a:r/a:rPr/a:ea/@typeface") == ["微软雅黑"] * 3
