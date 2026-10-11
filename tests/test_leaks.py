import pytest
from pptx import Presentation
from pptx.util import Emu

from factories import mini_template
from ppt_templates.cli import main
from ppt_templates.errors import PptError
from ppt_templates.leaks import find_leaks, sample_texts

BOX = (Emu(914400), Emu(914400), Emu(3657600), Emu(457200))


def make_sample(path, slides):
    """每页给 box / group / cell / notes 几处文字，覆盖原文可能藏身的地方。"""
    prs = Presentation()
    for page in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if "box" in page:
            slide.shapes.add_textbox(*BOX).text_frame.text = page["box"]
        if "group" in page:
            slide.shapes.add_group_shape().shapes.add_textbox(*BOX).text_frame.text = page["group"]
        if "cell" in page:
            slide.shapes.add_table(1, 1, *BOX).table.cell(0, 0).text = page["cell"]
        if "notes" in page:
            slide.notes_slide.notes_text_frame.text = page["notes"]
    prs.save(path)
    return path


@pytest.fixture
def sample(tmp_path):
    return make_sample(tmp_path / "sample.pptx", [
        {"box": "面向示例业务的整体策略", "group": "冷热分层 建模平台", "cell": "Throughput 提升 200%", "notes": "备注里的口径说明"},
        {"box": "第二页只讲时延"},
    ])


def test_sample_texts_reads_boxes_groups_tables_and_notes(sample):
    texts = sample_texts(sample)
    for text in ("面向示例业务的整体策略", "冷热分层 建模平台", "Throughput 提升 200%", "备注里的口径说明", "第二页只讲时延"):
        assert text in texts


def test_sample_texts_checks_the_file_and_slide_numbers(sample, tmp_path):
    with pytest.raises(PptError, match="不存在"):
        sample_texts(tmp_path / "none.pptx")
    with pytest.raises(PptError, match="样板只有 2 页，没有第 3 页"):
        sample_texts(sample, [3])


def test_four_chinese_chars_in_a_row_are_a_leak(sample):
    texts = sample_texts(sample)
    assert find_leaks({"title": "写清冷热分层的做法"}, texts) == ["/title  「冷热分层」与样板原文重合"]
    assert find_leaks({"title": "写清冷热分开的做法"}, texts) == []  # 「冷热分」只有 3 个字


def test_overlapping_windows_merge_and_whitespace_is_ignored(sample):
    # 原文「冷热分层 建模平台」中间有空格，仍按「冷热分层建模」整段报出
    assert find_leaks({"note": "说明冷热分层建模怎么做"}, sample_texts(sample)) == ["/note  「冷热分层建模」与样板原文重合"]


def test_ascii_runs_ignore_case_and_numbers_must_match_exactly(sample):
    texts = sample_texts(sample)
    assert find_leaks({"note": "throughput 翻倍"}, texts) == ["/note  「throughput」与样板原文重合"]
    assert find_leaks({"note": "目标值 200 以上"}, texts) == ["/note  「200」与样板原文重合"]
    assert find_leaks({"note": "目标值 2000 以上，through 一词"}, texts) == []


def test_pointers_follow_nested_lists_and_dicts(sample):
    example = {"groups": [{"name": "分组名称"}, {"items": ["条目名称", "写清整体策略的目标"]}]}
    assert find_leaks(example, sample_texts(sample)) == ["/groups/1/items/1  「整体策略」与样板原文重合"]


def test_allowed_words_and_unselected_slides_are_skipped(sample):
    example = {"title": "写清冷热分层的做法", "note": "只讲时延的那一页"}
    assert find_leaks(example, sample_texts(sample, [1]), allow=["冷热分层"]) == []
    assert find_leaks(example, sample_texts(sample, [2])) == ["/note  「只讲时延」与样板原文重合"]


def test_cli_leaks(make_repo, sample, monkeypatch, capsys):
    template, _ = mini_template()
    repo = make_repo({"common.summary-note": (template, {"title": "总结", "note": "写清面向示例业务的结论"})})
    monkeypatch.chdir(repo)
    assert main(["leaks", "common.summary-note", str(sample), "--slide", "1"]) == 1
    err = capsys.readouterr().err
    assert "/note  「面向示例业务的」与样板原文重合" in err and "1 处" in err
    assert main(["leaks", "common.summary-note", str(sample), "--allow", "面向示例业务"]) == 0
    assert "没有与样板原文重合的片段" in capsys.readouterr().out
    assert main(["leaks", "common.none", str(sample)]) == 1
    assert "缺少 example.yaml" in capsys.readouterr().err


def test_cli_leaks_also_checks_template_text_but_not_boxes_or_numbers(make_repo, sample, monkeypatch, capsys):
    template, _ = mini_template(name="整体策略模板", when_to_use="上限 200 字，用来写冷热分层建模的结论")
    template["layout"][0]["box"] = [0, 0, 1216, 200]  # 坐标、字数上限里的数字不算重合
    repo = make_repo({"common.summary-note": (template, {"title": "总结", "note": "一句话说清楚结论"})})
    monkeypatch.chdir(repo)
    assert main(["leaks", "common.summary-note", str(sample)]) == 1
    err = capsys.readouterr().err
    assert "template.yaml/name  「整体策略」与样板原文重合" in err
    assert "template.yaml/when_to_use  「冷热分层建模」与样板原文重合" in err
    assert "200" not in err.replace("2 处", "") and "example.yaml" not in err
