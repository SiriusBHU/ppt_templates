"""grid 的 weight（按权重分行高 / 列宽）、shape 的 radius（固定圆角）、bullets 的 marker: false（不带项目符号的并列段落）、模板的 structure。"""

import copy

import pytest
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from factories import deck_of
from ppt_templates.geometry import grid_cells
from ppt_templates.render import render_deck
from ppt_templates.validate import validate_repo

TID = "common.overview-weighted"


def weighted_template():
    """frame: none，坐标就是画布像素：外层单列 grid 按 size 分行高，内层每行最多 3 张圆角卡片按 w 分列宽。"""
    card = {"type": "object", "required": ["name"], "properties": {
        "name": {"type": "string", "maxLength": 6},
        "w": {"type": "number", "minimum": 1, "maximum": 2},
        "points": {"type": "array", "maxItems": 3, "items": {"type": "string", "maxLength": 20}},
    }}
    row = {"type": "object", "required": ["cards"], "properties": {
        "size": {"type": "number", "minimum": 1, "maximum": 2},
        "cards": {"type": "array", "minItems": 1, "maxItems": 3, "items": card},
    }}
    template = {
        "id": TID, "name": "按权重分配", "roles": ["overview"], "when_to_use": "测试用", "frame": "none",
        "structure": ["每行一组卡片，行高和卡宽按内容多少分配"],
        "slots": {"type": "object", "required": ["rows"],
                  "properties": {"rows": {"type": "array", "minItems": 1, "maxItems": 3, "items": row}}},
        "layout": [{"type": "grid", "box": [0, 0, 1000, 620], "bind": "rows", "columns": 1, "gap": 20, "weight": "size", "item": [
            {"type": "grid", "box": [0, 0, "100%", "100%"], "bind": "cards", "columns": 3, "gap": 10, "weight": "w", "item": [
                {"type": "shape", "box": [0, 0, "100%", "100%"], "shape": "rounded", "radius": 12, "fill": "bg", "line": "line"},
                {"type": "text", "box": [4, 4, "100%-8", 20], "bind": "name", "style": "body"},
                {"type": "bullets", "box": [4, 30, "100%-8", "100%-34"], "bind": "points", "style": "body", "marker": False},
            ]},
        ]}],
    }
    example = {"rows": [
        {"cards": [{"name": "甲", "w": 2}, {"name": "乙"}]},
        {"size": 2, "cards": [{"name": "丙", "points": ["第一段", "第二段"]}]},
    ]}
    return template, example


def _px(shape):
    return tuple(round(v / 9525, 1) for v in (shape.left, shape.top, shape.width, shape.height))


def _cards(prs):
    return [sh for sh in prs.slides[0].shapes if sh.shape_type == 1 and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]


def test_grid_cells_weights_split_rows_in_one_column():
    cells = grid_cells((0, 0, 100, 330), columns=1, gap=10, count=3, weights=[1, 2, 1])
    assert cells == [(0, 0, 100, 77.5), (0, 87.5, 100, 155), (0, 252.5, 100, 77.5)]


def test_grid_cells_weights_split_each_row_across_its_items():
    cells = grid_cells((0, 0, 310, 210), columns=2, gap=10, count=3, weights=[2, 1, 3])
    assert cells[:2] == [(0, 0, 200, 100), (210, 0, 100, 100)]
    assert cells[2] == (0, 110, 310, 100)  # 最后一行只有一项，按权重占满整行


def test_weighted_rows_and_cards_render_with_fixed_radius(make_repo):
    template, example = weighted_template()
    repo = make_repo({TID: (template, example)})
    assert validate_repo(repo).errors == []
    cards = _cards(render_deck(deck_of(TID, example), repo))
    assert [_px(c) for c in cards] == [(0, 0, 660, 200), (670, 0, 330, 200), (0, 220, 1000, 400)]
    assert [round(c.adjustments[0], 4) for c in cards] == [0.06, round(12 / 200, 4), 0.03]  # 半径 12px ÷ 短边


def test_rounded_without_radius_keeps_the_powerpoint_default(make_repo):
    template, example = weighted_template()
    del template["layout"][0]["item"][0]["item"][0]["radius"]
    repo = make_repo({TID: (template, example)})
    assert all(c._element.find(".//" + qn("a:avLst")).find(qn("a:gd")) is None for c in _cards(render_deck(deck_of(TID, example), repo)))


def test_bullets_without_marker_are_plain_spaced_paragraphs(make_repo):
    template, example = weighted_template()
    repo = make_repo({TID: (template, example)})
    box = next(sh for sh in render_deck(deck_of(TID, example), repo).slides[0].shapes
               if sh.has_text_frame and sh.text_frame.text.startswith("第一段"))
    paras = box.text_frame.paragraphs
    assert [p.text for p in paras] == ["第一段", "第二段"]
    assert all(p._p.pPr.find(qn("a:buChar")) is None and p._p.pPr.get("marL") is None for p in paras)
    assert paras[1]._p.pPr.find(qn("a:spcBef")) is not None


def _errors(make_repo, edit):
    template, example = weighted_template()
    template = copy.deepcopy(template)
    edit(template)
    return validate_repo(make_repo({TID: (template, example)})).errors


def _set_weight_slot(t, **schema):
    t["slots"]["properties"]["rows"]["items"]["properties"]["size"] = schema


@pytest.mark.parametrize("edit, expected", [
    (lambda t: t["layout"][0].update(weight="nope"), "/layout/0/weight  槽位 nope 不存在"),
    (lambda t: _set_weight_slot(t, type="number"), "/layout/0/weight  weight 的槽位必须是写了 minimum（> 0）和 maximum 的 number"),
    (lambda t: _set_weight_slot(t, type="string", maxLength=3), "weight 的槽位必须是写了 minimum"),
    (lambda t: _set_weight_slot(t, type="number", minimum=0, maximum=2), "weight 的槽位必须是写了 minimum"),
    (lambda t: t["layout"][0].update(row_height=100), "/layout/0/weight  单列 grid 的 weight 分配行高，不能再写 row_height"),
    (lambda t: t["layout"][0]["item"][0]["item"][0].update(shape="rect"), "/layout/0/item/0/item/0/radius  radius 只用于 shape: rounded"),
    (lambda t: t.update(structure=["说明", 3]), "/structure/1"),
])
def test_layout_option_errors(make_repo, edit, expected):
    assert any(expected in e for e in _errors(make_repo, edit))


def test_geometry_uses_the_least_favourable_weights(make_repo):
    # 等分时每行 (620-40)/3≈193 高放得下；按权重最不利（1 : 2 : 2）只有 116，内层 bullets 需要 34 + 100
    def edit(t):
        t["layout"][0]["item"][0]["item"][2]["box"] = [4, 30, "100%-8", 100]
    assert any("/layout/0/item/0/item/2/box" in e and "116" in e for e in _errors(make_repo, edit))


def _variant(**edits):
    template, example = weighted_template()
    template = copy.deepcopy(template)
    rows = template["slots"]["properties"]["rows"]["items"]["properties"]
    card = rows["cards"]["items"]["properties"]
    card["points"] = {"type": "array", "maxItems": 3, "items": {"type": "object", "required": ["text"], "properties": {
        "text": {"type": "string"}, "name": {"type": "string"}, "highlight": {"type": "string"}}}}
    card["tone"] = {"type": "string", "enum": ["红", "蓝"]}
    for key, value in edits.items():
        template[key] = value
    return template, example


def test_bullets_name_highlight_and_color_map(make_repo):
    template, _ = _variant()
    points = template["layout"][0]["item"][0]["item"][2]
    points.update(by="tone", color_map={"红": "accent", "蓝": "primary_text"})
    example = {"rows": [{"cards": [{"name": "甲", "tone": "红", "points": [{"name": "名称", "text": "说明，", "highlight": "强调"}]}]}]}
    repo = make_repo({TID: (template, example)})
    assert validate_repo(repo).errors == []
    box = next(sh for sh in render_deck(deck_of(TID, example), repo).slides[0].shapes
               if sh.has_text_frame and sh.text_frame.text.startswith("名称"))
    runs = box.text_frame.paragraphs[0].runs
    assert [(r.text, bool(r.font.bold), str(r.font.color.rgb)) for r in runs] == [
        ("名称：", True, "C00000"), ("说明，", False, "C00000"), ("强调", True, "0000FF")]


def test_bullets_color_map_needs_by(make_repo):
    template, example = _variant()
    template["layout"][0]["item"][0]["item"][2]["color_map"] = {"红": "accent"}
    errors = validate_repo(make_repo({TID: (template, example)})).errors
    assert any("bullets 的 color_map 必须配合 by" in e for e in errors)


def test_dashed_outline(make_repo):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    template, example = weighted_template()
    template["layout"][0]["item"][0]["item"][0]["dash"] = True
    repo = make_repo({TID: (template, example)})
    assert validate_repo(repo).errors == []
    assert all(c.line.dash_style == MSO_LINE_DASH_STYLE.DASH for c in _cards(render_deck(deck_of(TID, example), repo)))


def test_dash_needs_line(make_repo):
    def edit(t):
        card = t["layout"][0]["item"][0]["item"][0]
        card["dash"] = True
        del card["line"]
    assert any("/layout/0/item/0/item/0/dash  dash 是描边的虚线样式" in e for e in _errors(make_repo, edit))


def test_weight_can_be_the_length_of_a_list(make_repo):
    template, example = weighted_template()
    template["layout"][0]["weight"] = "cards"  # 卡片多的行更高
    example = {"rows": [{"cards": [{"name": "甲"}]}, {"cards": [{"name": "乙"}, {"name": "丙"}, {"name": "丁"}]}]}
    repo = make_repo({TID: (template, example)})
    assert validate_repo(repo).errors == []
    cards = _cards(render_deck(deck_of(TID, example), repo))
    assert _px(cards[0])[3] == 150 and _px(cards[1])[3] == 450   # (620-20) 按 1 : 3 分
