import pytest
import yaml
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

from factories import cards_template, deck_of, mini_template, rows_template
from ppt_templates.cli import main
from ppt_templates.errors import PptError
from ppt_templates.loader import load_style
from ppt_templates.render import render_deck
from ppt_templates.slots import fill_defaults, prepare_pages


def _arch_layers(data):
    return [layer for part in data["arch"]["partitions"] for layer in part["layers"]] + data["arch"].get("shared", [])


@pytest.fixture
def pptx_path(deck, root, tmp_path):
    path = tmp_path / "out.pptx"
    render_deck(deck, root).save(path)
    return path


def test_render_example_reopens_with_one_slide_of_16_9(pptx_path):
    prs = Presentation(pptx_path)
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    assert len(prs.slides) == 1


def test_render_contains_title_and_every_module_name(deck, pptx_path):
    data = deck["pages"][0]["data"]
    all_texts = [sh.text_frame.text for sh in Presentation(pptx_path).slides[0].shapes if sh.has_text_frame]
    texts = set(all_texts)
    assert f"{data['title']}　{data['title_suffix']}" in texts
    assert data["takeaway"] in texts
    for layer in _arch_layers(data):
        assert layer["name"] in texts
        for module in layer["modules"]:
            assert module["name"] in texts
    for part in data["arch"]["partitions"]:
        assert part["name"] in texts
    for item in data["challenges"] + data["principles"]:
        assert f"{item['name']}：{item['detail']}" in texts  # runs：名称、冒号、说明连排在同一个文本框里
    assert {str(i) for i in range(1, len(data["key_techs"]) + 1)} <= texts  # bind: "#" 渲染出的序号
    assert all_texts.count("1") == 1  # 只有关键技术的序号，不画页码


def test_every_run_has_cjk_and_cs_fonts(pptx_path):
    slide = Presentation(pptx_path).slides[0]
    runs = slide._element.xpath(".//a:r")
    assert runs
    for tag in ("latin", "ea", "cs"):
        assert slide._element.xpath(f".//a:r/a:rPr/a:{tag}/@typeface") == ["微软雅黑"] * len(runs)


def _shape_with_text(pptx_path, text):
    return next(sh for sh in Presentation(pptx_path).slides[0].shapes if sh.has_text_frame and sh.text_frame.text == text)


def _title_text(deck):
    data = deck["pages"][0]["data"]
    return f"{data['title']}　{data['title_suffix']}"


def test_default_style_draws_nothing_before_the_title(pptx_path):
    first = Presentation(pptx_path).slides[0].shapes[0]  # 页框不画标题下的分隔线：第一个形状就是标题
    assert first.has_text_frame and first.top == 32 * 9525


@pytest.mark.parametrize("decorations", [None, []])
def test_style_without_or_with_empty_decorations_is_valid(make_repo, decorations):
    template, example = mini_template()
    repo = make_repo({template["id"]: (template, example)})
    path = repo / "styles" / "default" / "style.yaml"
    style = yaml.safe_load(path.read_text(encoding="utf-8"))
    style["frame"].pop("decorations", None)
    if decorations is not None:
        style["frame"]["decorations"] = decorations
    path.write_text(yaml.safe_dump(style, allow_unicode=True), encoding="utf-8")
    slide = render_deck(deck_of(template["id"], example), repo).slides[0]
    assert [sh.text_frame.text for sh in slide.shapes] == [example["title"], example["note"]]  # 只有页框和内容，没有装饰


def test_frame_decorations_have_no_outline_and_no_shadow(make_repo):
    template, example = mini_template()
    repo = make_repo({template["id"]: (template, example)})
    path = repo / "styles" / "default" / "style.yaml"
    style = yaml.safe_load(path.read_text(encoding="utf-8"))
    style["frame"]["decorations"] = [{"type": "shape", "box": [32, 116, 1216, 1], "fill": "line"}]
    path.write_text(yaml.safe_dump(style, allow_unicode=True), encoding="utf-8")
    decoration = render_deck(deck_of(template["id"], example), repo).slides[0].shapes[0]  # 装饰先于其他内容绘制
    assert (decoration.left, decoration.top, decoration.width, decoration.height) == (32 * 9525, 116 * 9525, 1216 * 9525, 9525)
    assert decoration._element.xpath("./p:spPr/a:ln/a:noFill")
    assert decoration._element.xpath("./p:spPr/a:effectLst[not(a:outerShdw)]")


def test_positions_follow_frame_regions_and_body_origin(deck, pptx_path):
    title = _shape_with_text(pptx_path, _title_text(deck))
    assert (title.left, title.top, title.width, title.height) == (32 * 9525, 32 * 9525, 1216 * 9525, 40 * 9525)
    label = _shape_with_text(pptx_path, _arch_layers(deck["pages"][0]["data"])[0]["name"])  # 第 1 个分区第 1 层的层名；
    # partitioned-arch 的 box 原点 = body 左上角 (32, 128 + 108)；页签 28 + 间距 4 + 内边距 8
    assert (label.left, label.top, label.width) == ((32 + 8) * 9525, (128 + 108 + 28 + 4 + 8) * 9525, 56 * 9525)


def test_margins_are_equal_on_all_sides(root):
    from ppt_templates.loader import load_style
    frame = load_style(root, "default")["frame"]
    x, y, w, h = frame["body"]["box"]
    assert frame["title"]["box"][:2] == [32, 32]
    assert (x, 1280 - x - w, 720 - y - h) == (32, 32, 32)


def test_title_and_suffix_share_one_line_with_24_and_18_pt(deck, pptx_path):
    runs = _shape_with_text(pptx_path, _title_text(deck)).text_frame.paragraphs[0].runs
    assert [(r.font.size.pt, bool(r.font.bold)) for r in runs] == [(24, True), (18, True)]


def test_title_without_suffix_is_a_single_run(deck, root, tmp_path):
    del deck["pages"][0]["data"]["title_suffix"]
    path = tmp_path / "t.pptx"
    render_deck(deck, root).save(path)
    runs = _shape_with_text(path, deck["pages"][0]["data"]["title"]).text_frame.paragraphs[0].runs
    assert len(runs) == 1 and runs[0].font.size.pt == 24


def test_every_font_size_is_in_the_allowed_set(pptx_path):
    sizes = {r.font.size.pt for sh in Presentation(pptx_path).slides[0].shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs}
    assert sizes <= {24, 18, 14, 12, 10, 9, 8}


def test_highlighted_module_style(deck, pptx_path):
    modules = deck["pages"][0]["data"]["arch"]["shared"][0]["modules"]
    name_of = lambda hot: next(m["name"] for m in modules if m.get("highlight", False) == hot)  # noqa: E731
    hot, plain = _shape_with_text(pptx_path, name_of(True)), _shape_with_text(pptx_path, name_of(False))
    assert hot.line.width == 19050 and plain.line.width == 9525
    assert str(hot.line.color.rgb) == "C00000" and str(hot.fill.fore_color.rgb) == "FFF1E6"
    run = hot.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == "C00000"
    assert plain.text_frame.paragraphs[0].runs[0].font.bold


def test_legend_only_when_a_module_is_highlighted(deck, root, tmp_path):
    def texts(d):
        path = tmp_path / "t.pptx"
        render_deck(d, root).save(path)
        return {sh.text_frame.text for sh in Presentation(path).slides[0].shapes if sh.has_text_frame}

    assert "本次新增 / 改动" in texts(deck)
    for layer in _arch_layers(deck["pages"][0]["data"]):
        for module in layer["modules"]:
            module["highlight"] = False
    assert "本次新增 / 改动" not in texts(deck)


def test_title_over_limit_is_rejected_with_json_pointer(deck, root):
    deck["pages"][0]["data"]["title"] = "标" * 30
    with pytest.raises(PptError) as exc:
        render_deck(deck, root)
    assert "第 1 页 system.solution-arch-overview  /pages/0/data/title  长度 30，超过上限 14" in str(exc.value)


def test_nested_error_pointer_and_page_number(deck, root):
    deck["pages"].append({"template": deck["pages"][0]["template"], "data": {"title": "x", "challenges": []}})
    with pytest.raises(PptError) as exc:
        prepare_pages(deck, root)
    assert "第 2 页 system.solution-arch-overview  /pages/1/data/challenges  共 0 项，低于下限 2" in str(exc.value)
    assert "/pages/1/data  缺少必填字段 takeaway、principles、arch、key_techs" in str(exc.value)


def test_unknown_template_is_reported(deck, root):
    deck["pages"][0]["template"] = "system.nope"
    with pytest.raises(PptError, match="第 1 页 system.nope  /pages/0/template"):
        prepare_pages(deck, root)


def test_fill_defaults_recurses_into_arrays_and_objects(deck, root):
    template, data = prepare_pages(deck, root)[0]
    modules = [m for layer in _arch_layers(data) for m in layer["modules"]]
    assert all("highlight" in m for m in modules)
    assert modules[0]["highlight"] is False
    # 不修改调用方传入的数据
    assert "highlight" not in deck["pages"][0]["data"]["arch"]["partitions"][0]["layers"][0]["modules"][0]
    assert fill_defaults({"type": "object", "properties": {"a": {"default": 1}}}, {}) == {"a": 1}


def test_cli_render_writes_pptx(deck, tmp_path, monkeypatch, root):
    import yaml

    deck_path = tmp_path / "deck.yaml"
    deck_path.write_text(yaml.safe_dump(deck, allow_unicode=True), encoding="utf-8")
    monkeypatch.chdir(root)
    assert main(["render", str(deck_path), "-o", str(tmp_path / "o" / "out.pptx")]) == 0
    assert len(Presentation(tmp_path / "o" / "out.pptx").slides) == 1


def test_cli_exits_nonzero_and_prints_pointer_on_invalid_data(deck, tmp_path, monkeypatch, root, capsys):
    import yaml

    deck["pages"][0]["data"]["title"] = "标" * 30
    deck_path = tmp_path / "deck.yaml"
    deck_path.write_text(yaml.safe_dump(deck, allow_unicode=True), encoding="utf-8")
    monkeypatch.chdir(root)
    assert main(["render", str(deck_path), "-o", str(tmp_path / "out.pptx")]) == 1
    assert "/pages/0/data/title" in capsys.readouterr().err
    assert not (tmp_path / "out.pptx").exists()


def test_deck_structure_errors_are_reported(root):
    with pytest.raises(PptError, match="deck 结构错误"):
        render_deck({"pages": []}, root)


def _renderer(make_repo, template):
    """返回 render(data) -> 第 1 页的全部形状；同一个测试里 make_repo 只能调用一次，所以一个模板对应一个 renderer。"""
    repo = make_repo({template["id"]: (template, None)})
    return lambda data: list(render_deck(deck_of(template["id"], data), repo).slides[0].shapes)


def _shapes(make_repo, template, data):
    return _renderer(make_repo, template)(data)


def _px(shape):
    return tuple(v / 9525 for v in (shape.left, shape.top, shape.width, shape.height))


def _text_px(shapes, text):
    return _px(next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == text))


def test_grid_columns_is_a_maximum_so_few_items_fill_the_width(make_repo):
    template, _ = cards_template()  # columns: 4, gap: 20, box 1200×400，body 区原点 (32, 128)
    shapes = _shapes(make_repo, template, {"title": "t", "cards": ["甲", "乙"]})
    assert _text_px(shapes, "甲") == (32, 128, 590, 400)
    assert _text_px(shapes, "乙") == (32 + 610, 128, 590, 400)


def test_grid_uses_all_columns_once_items_reach_the_maximum(make_repo):
    template, _ = cards_template()
    shapes = _shapes(make_repo, template, {"title": "t", "cards": ["甲", "乙", "丙", "丁", "戊"]})
    assert _text_px(shapes, "丁") == (32 + 3 * 305, 128, 285, 190)
    assert _text_px(shapes, "戊") == (32, 128 + 210, 285, 190)


def test_grid_row_height_fixes_cell_height_and_stacks_from_top(make_repo):
    template, _ = cards_template(columns=2, row_height=60)
    shapes = _shapes(make_repo, template, {"title": "t", "cards": ["甲", "乙", "丙"]})
    assert _text_px(shapes, "甲") == (32, 128, 590, 60)
    assert _text_px(shapes, "丙") == (32, 128 + 80, 590, 60)


def test_grid_index_bind_renders_one_based_numbers_in_item_order(make_repo):
    template, _ = cards_template()
    template["layout"][0]["item"].append({"type": "text", "box": [0, 0, 20, 20], "bind": "#", "style": "body"})
    shapes = _shapes(make_repo, template, {"title": "t", "cards": ["甲", "乙", "丙"]})
    numbers = {sh.text_frame.text: _px(sh) for sh in shapes if sh.has_text_frame and sh.text_frame.text.isdigit()}
    assert list(numbers) == ["1", "2", "3"]
    cell_w = (1200 - 2 * 20) / 3  # 3 项 → 3 列
    assert [numbers[n][0] for n in numbers] == pytest.approx([32 + i * (cell_w + 20) for i in range(3)])


def _when_template(layout, **props):
    slots = {"type": "object", "required": ["title"],
             "properties": {"title": {"type": "string"}, "note": {}, **props}}
    return mini_template(slots=slots, layout=layout)[0]


FLAG = {"type": "shape", "box": [0, 100, 50, 50], "fill": "accent", "when": "note"}


@pytest.mark.parametrize("note, shown", [
    (None, False), ("", False), ([], False), ({}, False),
    ("x", True), (["a"], True), ({"k": 1}, True), (0, True), (False, True),
])
def test_when_hides_element_for_empty_values(make_repo, note, shown):
    data = {"title": "t", "note": note}
    shapes = _shapes(make_repo, _when_template([FLAG]), data)
    assert ((32, 228, 50, 50) in [_px(sh) for sh in shapes]) is shown


@pytest.mark.parametrize("kind, expected", [
    (None, MSO_SHAPE.RECTANGLE), ("rect", MSO_SHAPE.RECTANGLE), ("rounded", MSO_SHAPE.ROUNDED_RECTANGLE), ("oval", MSO_SHAPE.OVAL)])
def test_shape_kind_maps_to_autoshape_type(make_repo, kind, expected):
    shape = {"type": "shape", "box": [0, 100, 50, 50], "fill": "accent", **({"shape": kind} if kind else {})}
    shapes = _shapes(make_repo, _when_template([shape]), {"title": "t"})
    assert next(sh for sh in shapes if _px(sh) == (32, 228, 50, 50)).auto_shape_type == expected


def test_when_is_skipped_for_a_missing_slot(make_repo):
    shapes = _shapes(make_repo, _when_template([FLAG]), {"title": "t"})
    assert (32, 228, 50, 50) not in [_px(sh) for sh in shapes]


def test_when_resolves_in_the_grid_item_scope(make_repo):
    item = {"type": "object", "properties": {"name": {"type": "string"}, "badge": {"type": "string"}}}
    grid = {"type": "grid", "box": [0, 0, 400, 100], "bind": "cards", "columns": 2, "gap": 0, "item": [
        {"type": "text", "box": [0, 0, 100, 20], "bind": "name", "style": "body"},
        {"type": "shape", "box": [0, 30, 10, 10], "fill": "accent", "when": "badge"}]}
    template = _when_template([grid], cards={"type": "array", "maxItems": 4, "items": item})
    shapes = _shapes(make_repo, template, {"title": "t", "cards": [{"name": "甲", "badge": "新"}, {"name": "乙"}]})
    assert [_px(sh) for sh in shapes if sh.width == 10 * 9525] == [(32, 158, 10, 10)]


def test_when_false_skips_component_data_checks(make_repo):
    component = {"type": "component", "component": "layered-arch", "box": [0, 0, 600, 300], "bind": "layers", "when": "layers"}
    template = _when_template([component], layers={"type": "array"})
    assert _shapes(make_repo, template, {"title": "t", "layers": []})  # [] 不满足 layered-arch 的 minItems，但元素不渲染


def _color(root, token):
    return load_style(root, "default")["colors"][token].lstrip("#")


def _text_template(**text):
    layout = [{"type": "text", "box": [0, 0, 200, 20], "bind": "name", "style": "body", **text}]
    props = {"name": {"type": "string"}, "level": {"enum": ["高", "中", "低"]}, "n": {"type": "integer"}}
    return _when_template(layout, **props)


def _run_color(shapes, text):
    shape = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == text)
    return str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb)


def test_text_color_map_matches_bound_value_and_overrides_color(make_repo, root):
    render = _renderer(make_repo, _text_template(bind="level", color="muted", color_map={"高": "highlight_strong", "低": "accent"}))
    for level, token in (("高", "highlight_strong"), ("低", "accent"), ("中", "muted")):
        assert _run_color(render({"title": "t", "level": level}), level) == _color(root, token)


def test_text_color_map_without_match_or_color_uses_style_color(make_repo, root):
    template = _text_template(bind="level", color_map={"高": "highlight_strong"})
    shapes = _shapes(make_repo, template, {"title": "t", "level": "低"})
    assert _run_color(shapes, "低") == _color(root, "text")


def test_text_color_map_by_other_path_and_stringified_values(make_repo, root):
    render = _renderer(make_repo, _text_template(by="n", color_map={"3": "accent"}))
    assert _run_color(render({"title": "t", "name": "甲", "n": 3}), "甲") == _color(root, "accent")
    assert _run_color(render({"title": "t", "name": "甲", "n": 4}), "甲") == _color(root, "text")


def test_text_color_map_ignores_a_missing_by_value(make_repo, root):
    template = _text_template(by="n", color_map={"None": "accent"})
    shapes = _shapes(make_repo, template, {"title": "t", "name": "甲"})
    assert _run_color(shapes, "甲") == _color(root, "text")


def _swatch(render, level):
    return next(sh for sh in render({"title": "t", "level": level}) if _px(sh) == (32, 228, 50, 50))


def test_shape_fill_map_and_line_map_override_fill_and_line(make_repo, root):
    shape = {"type": "shape", "box": [0, 100, 50, 50], "fill": "surface", "line": "line", "by": "level",
             "fill_map": {"高": "highlight"}, "line_map": {"高": "highlight_strong"}}
    render = _renderer(make_repo, _when_template([shape], level={"enum": ["高", "中"]}))
    hot = _swatch(render, "高")
    assert str(hot.fill.fore_color.rgb) == _color(root, "highlight")
    assert str(hot.line.color.rgb) == _color(root, "highlight_strong")
    plain = _swatch(render, "中")
    assert str(plain.fill.fore_color.rgb) == _color(root, "surface")
    assert str(plain.line.color.rgb) == _color(root, "line")


def test_shape_line_map_adds_a_line_only_on_match(make_repo, root):
    shape = {"type": "shape", "box": [0, 100, 50, 50], "fill": "surface", "by": "level", "line_map": {"高": "highlight_strong"}}
    render = _renderer(make_repo, _when_template([shape], level={"enum": ["高", "中"]}))
    hot, plain = _swatch(render, "高"), _swatch(render, "中")
    assert str(hot.line.color.rgb) == _color(root, "highlight_strong")
    assert plain.line.fill.type == MSO_FILL.BACKGROUND  # 无描边


def _run_info(shape):
    paragraphs = shape.text_frame.paragraphs
    assert len(paragraphs) == 1  # runs 连排在同一个段落里，自然折行
    return [(r.text, bool(r.font.bold), str(r.font.color.rgb)) for r in paragraphs[0].runs]


def _runs_template(runs, **text):
    layout = [{"type": "text", "box": [0, 0, 400, 40], "style": "body", "runs": runs, **text}]
    return _when_template(layout, name={"type": "string"}, detail={"type": "string"})


def test_runs_are_consecutive_runs_in_one_paragraph_with_own_style_and_color(make_repo, root):
    runs = [{"bind": "name", "style": "body_strong"}, {"value": "："}, {"bind": "detail", "color": "accent"}]
    shapes = _shapes(make_repo, _runs_template(runs), {"title": "t", "name": "名称", "detail": "说明"})
    shape = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == "名称：说明")
    assert _run_info(shape) == [("名称", True, _color(root, "text")), ("：", False, _color(root, "text")),
                                ("说明", False, _color(root, "accent"))]
    assert [r.font.size.pt for r in shape.text_frame.paragraphs[0].runs] == [12, 12, 12]


def test_runs_fall_back_to_the_element_color_and_keep_alignment(make_repo, root):
    shapes = _shapes(make_repo, _runs_template([{"bind": "name"}], color="primary", align="center", valign="middle"),
                     {"title": "t", "name": "名称"})
    shape = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == "名称")
    assert _run_info(shape) == [("名称", False, _color(root, "primary"))]
    assert shape.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


def test_runs_skip_missing_values(make_repo):
    shapes = _shapes(make_repo, _runs_template([{"bind": "name"}, {"value": "："}, {"bind": "detail"}]), {"title": "t", "name": "名称"})
    assert "名称：" in [sh.text_frame.text for sh in shapes if sh.has_text_frame]


def test_runs_draw_nothing_when_every_run_is_missing(make_repo):
    shapes = _shapes(make_repo, _runs_template([{"bind": "name"}, {"bind": "detail"}]), {"title": "t"})
    assert [sh.text_frame.text for sh in shapes if sh.has_text_frame] == ["t"]  # 只有页框的标题


def test_runs_resolve_in_the_grid_item_scope_and_support_the_index_bind(make_repo):
    item = [{"type": "text", "box": [0, 0, "100%", 20], "style": "body",
             "runs": [{"bind": "#", "style": "body_strong"}, {"value": ". "}, {"bind": "."}]}]
    grid = {"type": "grid", "box": [0, 0, 400, 100], "bind": "cards", "columns": 1, "item": item}
    template = _when_template([grid], cards={"type": "array", "maxItems": 4, "items": {"type": "string"}})
    shapes = _shapes(make_repo, template, {"title": "t", "cards": ["甲", "乙"]})
    assert [sh.text_frame.text for sh in shapes if sh.has_text_frame][-2:] == ["1. 甲", "2. 乙"]


def test_shape_shadow_follows_the_style_and_is_off_by_default(make_repo, root):
    card = {"type": "shape", "box": [0, 100, 50, 50], "fill": "bg", "shadow": True}
    plain = {"type": "shape", "box": [0, 200, 50, 50], "fill": "bg"}
    shapes = _shapes(make_repo, _when_template([card, plain]), {"title": "t"})
    on, off = (next(sh for sh in shapes if _px(sh) == (32, 128 + y, 50, 50)) for y in (100, 200))
    spec = load_style(root, "default")["shadow"]
    (shdw,) = on._element.xpath("./p:spPr/a:effectLst/a:outerShdw")
    assert (shdw.get("blurRad"), shdw.get("dist"), shdw.get("dir")) == (
        str(round(spec["blur"] * 9525)), str(round(spec["distance"] * 9525)), str(round(spec["angle"] * 60000)))
    assert on._element.xpath("./p:spPr/a:effectLst/a:outerShdw/a:srgbClr/@val") == [_color(root, spec["color"])]
    assert on._element.xpath("./p:spPr/a:effectLst/a:outerShdw/a:srgbClr/a:alpha/@val") == [str(round(spec["alpha"] * 100000))]
    assert on._element.xpath("./p:style/a:effectRef/@idx") == ["0"]
    assert off._element.xpath("./p:spPr/a:effectLst") and not off._element.xpath("./p:spPr/a:effectLst/*")


def test_shadow_without_a_style_shadow_section_is_a_clear_error(make_repo):
    card = {"type": "shape", "box": [0, 100, 50, 50], "fill": "bg", "shadow": True}
    template, example = mini_template(layout=[card])
    repo = make_repo({template["id"]: (template, example)})
    path = repo / "styles" / "default" / "style.yaml"
    style = yaml.safe_load(path.read_text(encoding="utf-8"))
    del style["shadow"]
    path.write_text(yaml.safe_dump(style, allow_unicode=True), encoding="utf-8")
    with pytest.raises(PptError, match="风格没有定义 shadow"):
        render_deck(deck_of(template["id"], example), repo)


# ---- 嵌套 grid：外层每行一个领域，内层是该行的卡片 ----

def _rows_shapes(make_repo, rows):
    template, _ = rows_template()
    return _shapes(make_repo, template, {"title": "t", "rows": rows})


def test_nested_grid_resolves_binds_per_level_and_columns_per_row(make_repo):
    shapes = _rows_shapes(make_repo, [
        {"label": "行一", "cards": [{"name": "甲"}]},
        {"label": "行二", "cards": [{"name": "乙"}, {"name": "丙"}, {"name": "丁"}]},
    ])
    # 外层：1 列、row_height 100、gap 10，body 原点 (32, 128)，grid box 原点 (0, 0)；内层 box 从行内 x=110 起、宽 1000-110
    assert _text_px(shapes, "行一") == (32, 128, 100, 100)
    assert _text_px(shapes, "行二") == (32, 128 + 110, 100, 100)
    assert _text_px(shapes, "甲") == (32 + 110 + 4, 128 + 4, 890 - 8, 20)            # 1 张卡片：占满内层整行（列数 = min(4, 1)）
    cw = (890 - 2 * 8) / 3                                                           # 3 张卡片 → 3 列
    for k, name in enumerate("乙丙丁"):
        assert _text_px(shapes, name) == pytest.approx((32 + 110 + k * (cw + 8) + 4, 128 + 110 + 4, cw - 8, 20))


def test_nested_grid_index_bind_counts_the_inner_grid_from_one_in_every_row(make_repo):
    shapes = _rows_shapes(make_repo, [
        {"label": "行一", "cards": [{"name": "甲"}, {"name": "乙"}]},
        {"label": "行二", "cards": [{"name": "丙"}, {"name": "丁"}, {"name": "戊"}]},
    ])
    numbers = [sh.text_frame.text for sh in shapes if sh.has_text_frame and sh.text_frame.text.isdigit()]
    assert numbers == ["1", "2", "1", "2", "3"]


def test_nested_grid_draws_every_inner_cell_with_its_own_elements(make_repo):
    shapes = _rows_shapes(make_repo, [{"label": "行一", "cards": [{"name": "甲"}, {"name": "乙"}]}])
    cards = [_px(sh) for sh in shapes if sh.width == 441 * 9525]  # 内层 2 列：(890 - 8) / 2 宽的 surface 卡片底
    assert cards == [(32 + 110, 128, 441, 100), (32 + 110 + 449, 128, 441, 100)]
    assert [sh.text_frame.text for sh in shapes if sh.has_text_frame and sh.text_frame.text == "配图"] == ["配图", "配图"]
