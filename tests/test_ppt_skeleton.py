"""tools/ppt-skeleton 的测试：脚本不是包，按文件路径导入。"""
import importlib.util
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt

SCRIPT = Path(__file__).resolve().parents[1] / "tools" / "ppt-skeleton" / "scripts" / "skeleton.py"
spec = importlib.util.spec_from_file_location("ppt_skeleton", SCRIPT)
sk = importlib.util.module_from_spec(spec)
_dont_write = sys.dont_write_bytecode
sys.dont_write_bytecode = True  # 不要在 tools/ppt-skeleton/ 里留下 __pycache__
spec.loader.exec_module(sk)
sys.dont_write_bytecode = _dont_write

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
CT_XML = "application/xml"
SENSITIVE = "星云\nProject-Nebula\n华兴银行\n"


def png_bytes():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (20, 20), "red").save(buf, "PNG")
    buf.seek(0)
    return buf


def add_part(source, uri, ct, blob, reltype):
    package = source if hasattr(source, "iter_parts") else source.package
    part = Part(PackURI(uri), ct, package, blob)
    return source.relate_to(part, reltype)


def build_deck(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_only, blank = prs.slide_layouts[5], prs.slide_layouts[6]

    # ---- 第 1 页：文本框、圆角矩形、组合、表格、图片、图表、备注、超链接
    s1 = prs.slides.add_slide(title_only)
    s1.shapes.title.text = "星云平台系统方案总览"
    tb = s1.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5), Inches(1))
    para = tb.text_frame.paragraphs[0]
    r1 = para.add_run()
    r1.text = "100M长序列："
    r1.font.bold = True
    r1.font.size = Pt(16)
    r1.font.name = "微软雅黑"
    r2 = para.add_run()
    r2.text = "KV Cache 随序列线性增长"
    r2.font.size = Pt(16)
    p2 = tb.text_frame.add_paragraph()
    r = p2.add_run()
    r.text = "详见内部链接"
    r.hyperlink.address = "http://intra.example.com/nebula"
    tb.name = "星云文本框"

    rr = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(1.5), Inches(3), Inches(1.2))
    rr.fill.solid()
    rr.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    rr.line.color.rgb = RGBColor(0xFF, 0xC0, 0x00)
    rr.text_frame.text = "核心技术：混合精度训练框架"
    run = rr.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(14)
    run.font.name = "微软雅黑"
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    rr.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    rr.click_action.hyperlink.address = "http://intra.example.com/click"
    rr._element.nvSpPr.cNvPr.set("descr", "星云核心技术示意")

    grp = s1.shapes.add_group_shape()
    g1 = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(3), Inches(3), Inches(1))
    g1.text_frame.text = "子模块一：数据预处理流水线"
    g2 = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.8), Inches(3), Inches(3), Inches(1))
    g2.text_frame.text = "子模块二：分布式存储引擎"

    tbl = s1.shapes.add_table(2, 2, Inches(7), Inches(3), Inches(5), Inches(1.2)).table
    tbl.cell(0, 0).text = "指标名称"
    tbl.cell(0, 1).text = "当前水平"
    tbl.cell(1, 0).text = "推理吞吐量"
    tbl.cell(1, 1).text = "20480 TPS"

    pic = s1.shapes.add_picture(png_bytes(), Inches(0.5), Inches(5), Inches(2), Inches(1.5))
    pic._element.nvPicPr.cNvPr.set("descr", "星云Logo")
    pic._element.nvPicPr.cNvPr.set("name", "星云Logo.png")
    cd = CategoryChartData()
    cd.categories = ["星云指标甲", "星云指标乙"]
    cd.add_series("星云系列", (1234, 5678))
    s1.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3), Inches(5), Inches(4), Inches(2), cd)

    s1.notes_slide.notes_text_frame.text = "备注：给华兴银行的报价为 888 万"

    # ---- 第 2 页（不选）：含敏感词和数字
    s2 = prs.slides.add_slide(blank)
    t2 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    t2.text_frame.text = "机密：华兴银行专项 Project-Nebula 预算 20480 万"

    # ---- 第 3 页：被隐藏、含数字、点击跳转到第 2 页
    s3 = prs.slides.add_slide(blank)
    t3 = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    t3.text_frame.text = "吞吐提升 3.2 倍，延迟下降 45%"
    t4 = s3.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1))
    t4.text_frame.text = "客户：华兴银行联合验证"
    t4.click_action.target_slide = s2
    s3.shapes.add_picture(png_bytes(), Inches(9), Inches(1), Inches(2), Inches(2))
    s3._element.set("show", "0")
    s3.notes_slide.notes_text_frame.text = "第三页备注"

    # ---- 母版：文字 + 图片
    master = prs.slide_master
    master_tree = master._element.find("p:cSld/p:spTree", NS)
    img_part, rid = master.part.get_or_add_image_part(png_bytes())
    master_tree.append(etree.fromstring(
        '<p:sp xmlns:a="%(a)s" xmlns:p="%(p)s"><p:nvSpPr><p:cNvPr id="90" name="页脚公司名"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
        '</p:nvSpPr><p:spPr><a:xfrm><a:off x="100000" y="6400000"/><a:ext cx="3000000" cy="300000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>'
        "<a:p><a:r><a:rPr lang=\"zh-CN\"/><a:t>星云科技内部资料严禁外传</a:t></a:r></a:p></p:txBody></p:sp>" % NS))
    master_tree.append(etree.fromstring(
        '<p:pic xmlns:a="%(a)s" xmlns:p="%(p)s" xmlns:r="%(r)s"><p:nvPicPr><p:cNvPr id="91" name="公司Logo"/>'
        '<p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="' % NS + rid + '"/><a:stretch><a:fillRect/>'
        '</a:stretch></p:blipFill><p:spPr><a:xfrm><a:off x="100000" y="100000"/><a:ext cx="800000" cy="400000"/>'
        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'))

    # ---- 文档属性、各种杂项部件
    cp = prs.core_properties
    cp.author, cp.last_modified_by, cp.title = "张三", "李四", "星云项目汇报"
    cp.subject, cp.keywords, cp.comments, cp.category = "华兴银行", "星云", "内部", "机密"
    pkg = prs.part.package
    for rel in pkg._rels.values():
        if rel.reltype == RT.EXTENDED_PROPERTIES:
            rel.target_part._blob = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
                'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
                "<Company>星云科技有限公司</Company><Manager>王总</Manager><TitlesOfParts><vt:vector size=\"1\" "
                'baseType="lpstr"><vt:lpstr>星云平台系统方案总览</vt:lpstr></vt:vector></TitlesOfParts></Properties>'
            ).encode("utf-8")
    add_part(pkg, "/docProps/custom.xml",
             "application/vnd.openxmlformats-officedocument.custom-properties+xml",
             "<Properties xmlns='http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'/>".encode(),
             RT.CUSTOM_PROPERTIES)
    assert any(r.reltype == RT.THUMBNAIL for r in pkg._rels.values())  # 默认模板自带缩略图
    add_part(prs.part, "/customXml/item1.xml", CT_XML, "<root>星云</root>".encode(), RT.CUSTOM_XML)
    add_part(prs.part, "/ppt/commentAuthors.xml",
             "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml",
             ("<p:cmAuthorLst xmlns:p='%(p)s'><p:cmAuthor id='0' name='张三' initials='ZS' lastIdx='1' clrIdx='0'/>"
              "</p:cmAuthorLst>" % NS).encode(), RT.COMMENT_AUTHORS)
    add_part(s1.part, "/ppt/comments/comment1.xml",
             "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
             ("<p:cmLst xmlns:p='%(p)s'><p:cm authorId='0' dt='2025-01-01T00:00:00.000' idx='1'><p:pos x='1' y='1'/>"
              "<p:text>批注里的星云机密</p:text></p:cm></p:cmLst>" % NS).encode(), RT.COMMENTS)
    tag_rid = add_part(s1.part, "/ppt/tags/tag1.xml",
                        "application/vnd.openxmlformats-officedocument.presentationml.tags+xml",
                        ("<p:tagLst xmlns:p='%(p)s'><p:tag name='PROJ' val='星云'/></p:tagLst>" % NS).encode(),
                        RT.TAGS)
    s1._element.find("p:cSld", NS).append(etree.fromstring(
        '<p:custDataLst xmlns:p="%(p)s" xmlns:r="%(r)s"><p:tags r:id="' % NS + tag_rid + '"/></p:custDataLst>'))

    # 伪造一个 SmartArt 图形框（带 diagram 数据部件）
    dgm_rid = add_part(s1.part, "/ppt/diagrams/data1.xml",
                   "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
                   "<dgm:dataModel xmlns:dgm='http://schemas.openxmlformats.org/drawingml/2006/diagram'/>".encode(),
                   "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData")
    s1.shapes._spTree.append(etree.fromstring(
        '<p:graphicFrame xmlns:a="%(a)s" xmlns:p="%(p)s" xmlns:r="%(r)s"><p:nvGraphicFramePr>'
        '<p:cNvPr id="95" name="星云架构图"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="8000000" y="5000000"/><a:ext cx="3000000" cy="1500000"/></p:xfrm><a:graphic>'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
        '<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" r:dm="' % NS
        + dgm_rid + '" r:lo="' + dgm_rid + '" r:qs="' + dgm_rid + '" r:cs="' + dgm_rid
        + '"/></a:graphicData></a:graphic></p:graphicFrame>'))
    prs.save(path)


def zip_names(path):
    with zipfile.ZipFile(path) as zf:
        return zf.namelist()


def zip_text(path):
    with zipfile.ZipFile(path) as zf:
        return "\n".join(zf.read(n).decode("utf-8", "ignore") for n in zf.namelist()
                         if n.endswith((".xml", ".rels")))


def fill(work, overrides=None):
    """给每个 id 填一个不含原文的引导文字；overrides 按原文子串覆盖。"""
    texts = json.loads((work / "texts.json").read_text(encoding="utf-8"))
    reps = {}
    for i, e in enumerate(texts["items"], 1):
        reps[e["id"]] = "引导文字第%s号位" % "甲乙丙丁戊己庚辛壬癸"[i % 10]
        for sub, val in (overrides or {}).items():
            if sub in e["text"]:
                reps[e["id"]] = val
    (work / "replacements.json").write_text(json.dumps(reps, ensure_ascii=False), encoding="utf-8")
    man = (work / "manifest.yaml").read_text(encoding="utf-8").replace('purpose: ""', 'purpose: "通用示意页"')
    (work / "manifest.yaml").write_text(man, encoding="utf-8")
    return reps


@pytest.fixture
def env(tmp_path):
    orig = tmp_path / "orig.pptx"
    build_deck(orig)
    (tmp_path / "words.txt").write_text(SENSITIVE, encoding="utf-8")
    return {"orig": orig, "work": tmp_path / "work", "out": tmp_path / "out" / "skeleton.pptx",
            "words": tmp_path / "words.txt", "tmp": tmp_path}


def run(*argv):
    return sk.main([str(a) for a in argv])


def make_skeleton(env, slides="1,3", overrides=None):
    assert run("extract", env["orig"], "--slides", slides, "-o", env["work"]) == 0
    fill(env["work"], overrides)
    assert run("apply", env["orig"], env["work"], "-o", env["out"]) == 0


def check(env, *extra):
    return run("check", env["orig"], env["out"], env["work"], "--words", env["words"], *extra)


# --------------------------------------------------------------------------- extract

def test_help_runs_standalone():
    res = subprocess.run([sys.executable, str(SCRIPT), "--help"], capture_output=True, text=True)
    assert res.returncode == 0
    assert "extract" in res.stdout and "check" in res.stdout


def test_extract_ids_stable_and_complete(env):
    run("extract", env["orig"], "--slides", "1,3", "-o", env["work"])
    first = (env["work"] / "texts.json").read_text(encoding="utf-8")
    run("extract", env["orig"], "--slides", "1,3", "-o", env["work"], "--force")
    assert (env["work"] / "texts.json").read_text(encoding="utf-8") == first

    data = json.loads(first)
    assert data["slides"] == [1, 3]
    items = data["items"]
    texts = [e["text"] for e in items]
    assert "星云平台系统方案总览" in texts
    assert "100M长序列：" in texts and "KV Cache 随序列线性增长" in texts  # 同段里格式不同（加粗/常规）拆成两条
    assert "子模块二：分布式存储引擎" in texts  # 组合内
    assert "20480 TPS" in texts  # 表格单元格
    assert "吞吐提升 3.2 倍，延迟下降 45%" in texts
    assert not any("华兴银行专项" in t or "备注" in t or "内部资料" in t for t in texts)  # 未选页、备注、母版不在内
    assert [e["id"] for e in items] == sorted(e["id"] for e in items)
    assert len({e["id"] for e in items}) == len(items)
    cell = next(e for e in items if e["text"] == "20480 TPS")
    assert cell["kind"] == "表格单元格" and "行2列2" in cell["location"]
    lead = next(e for e in items if e["text"] == "100M长序列：")
    para = next(e for e in items if e["text"].startswith("KV Cache"))
    assert lead["location"].endswith("片段1/2") and para["location"].endswith("片段2/2")
    assert para["font_pt"] == 16 and para["chars"] == len("KVCache随序列线性增长")
    assert para["budget"][0] < para["chars"] < para["budget"][1]
    title = next(e for e in items if e["kind"] == "占位符:title")
    assert title["font_pt"] == 44 and title["box_pct"] is not None  # 字号、位置从版式/母版继承
    assert cell["box_pct"][2] < 50  # 表格单元格给的是单元格的框，不是整张表
    assert json.loads((env["work"] / "replacements.json").read_text(encoding="utf-8")) == {e["id"]: "" for e in items}
    assert (env["work"] / "manifest.yaml").read_text(encoding="utf-8").count("purpose:") == 2


def test_extract_does_not_overwrite_filled_work_dir(env):
    run("extract", env["orig"], "--slides", "1", "-o", env["work"])
    assert run("extract", env["orig"], "--slides", "1", "-o", env["work"]) == 1


def test_extract_rejects_bad_slides(env):
    assert run("extract", env["orig"], "--slides", "9", "-o", env["work"]) == 1
    assert run("extract", env["orig"], "--slides", "1,x", "-o", env["work"]) == 1


# --------------------------------------------------------------------------- apply

def test_apply_requires_every_replacement(env, capsys):
    run("extract", env["orig"], "--slides", "1", "-o", env["work"])
    assert run("apply", env["orig"], env["work"], "-o", env["out"]) == 1
    assert "没有填写" in capsys.readouterr().err
    assert not env["out"].exists()


def test_apply_refuses_output_inside_work_dir(env):
    run("extract", env["orig"], "--slides", "1", "-o", env["work"])
    fill(env["work"])
    assert run("apply", env["orig"], env["work"], "-o", env["work"] / "x.pptx") == 1


def test_apply_rejects_changed_source(env, tmp_path):
    run("extract", env["orig"], "--slides", "1", "-o", env["work"])
    fill(env["work"])
    other = tmp_path / "other.pptx"
    prs = Presentation(env["orig"])
    prs.slides[0].shapes.title.text = "改过的标题"
    prs.save(other)
    assert run("apply", other, env["work"], "-o", env["out"]) == 1


def test_apply_keeps_layout_and_replaces_everything(env):
    make_skeleton(env)
    src, out = Presentation(env["orig"]), Presentation(env["out"])
    assert len(out.slides) == 2
    s_src, s_out = src.slides[0], out.slides[0]
    assert out.slide_width == src.slide_width and out.slide_height == src.slide_height

    # 顶层形状几何完全一致，顺序不变
    geo = lambda s: [(sh.left, sh.top, sh.width, sh.height) for sh in s.shapes]  # noqa: E731
    assert geo(s_out) == geo(s_src)
    assert [sh.shape_type for sh in s_out.shapes if sh.shape_type == MSO_SHAPE_TYPE.GROUP]

    rr_src = next(sh for sh in s_src.shapes if sh.has_text_frame and "核心技术" in sh.text_frame.text)
    rr_out = list(s_out.shapes)[list(s_src.shapes).index(rr_src)]
    assert rr_out.fill.fore_color.rgb == rr_src.fill.fore_color.rgb
    assert rr_out.line.color.rgb == rr_src.line.color.rgb
    run_src, run_out = rr_src.text_frame.paragraphs[0].runs[0], rr_out.text_frame.paragraphs[0].runs[0]
    assert (run_out.font.size, run_out.font.name, run_out.font.color.rgb) == (
        run_src.font.size, run_src.font.name, run_src.font.color.rgb)
    assert rr_out.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert "核心技术" not in rr_out.text_frame.text

    # 同段里格式不同的两个 run 各自替换，各自保留加粗/字号
    tb_src = next(sh for sh in s_src.shapes if sh.has_text_frame and sh.text_frame.text.startswith("100M"))
    tb_out = list(s_out.shapes)[list(s_src.shapes).index(tb_src)]
    runs = tb_out.text_frame.paragraphs[0].runs
    assert len(runs) == 2 and runs[0].font.bold is True and runs[1].font.bold is None
    assert runs[0].font.size == runs[1].font.size == Pt(16) and runs[0].font.name == "微软雅黑"
    assert "100M" not in tb_out.text_frame.text

    # 组合内的子形状位置不变
    g_src = next(sh for sh in s_src.shapes if sh.shape_type == MSO_SHAPE_TYPE.GROUP)
    g_out = list(s_out.shapes)[list(s_src.shapes).index(g_src)]
    assert [(c.left, c.top, c.width, c.height) for c in g_out.shapes] == [
        (c.left, c.top, c.width, c.height) for c in g_src.shapes]
    assert all("子模块" not in c.text_frame.text for c in g_out.shapes)

    # 表格结构保留，内容替换
    t_src = next(sh for sh in s_src.shapes if sh.has_table).table
    t_out = next(sh for sh in s_out.shapes if sh.has_table).table
    assert [c.width for c in t_out.columns] == [c.width for c in t_src.columns]
    assert [r.height for r in t_out.rows] == [r.height for r in t_src.rows]
    assert "20480" not in t_out.cell(1, 1).text

    # 图片 / 图表 / SmartArt -> 灰色矩形 + 标签
    labels = {sh.text_frame.text for sh in s_out.shapes
              if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.text_frame.text in ("图片", "图表", "图示")}
    assert labels == {"图片", "图表", "图示"}
    pic_src = next(sh for sh in s_src.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    pic_out = list(s_out.shapes)[list(s_src.shapes).index(pic_src)]
    assert pic_out.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and pic_out.text_frame.text == "图片"
    assert pic_out.fill.fore_color.rgb == RGBColor.from_string("E7E6E6")

    # 形状名、替代文字被抹掉
    names = [sh.name for sh in s_out.shapes]
    assert all(n.startswith("Shape ") for n in names)
    assert "descr" not in etree.tostring(s_out._element).decode()

    # 隐藏页取消隐藏；幻灯片部件按新顺序重新编号
    assert out.slides[1]._element.get("show") is None
    assert sorted(n for n in zip_names(env["out"]) if n.startswith("ppt/slides/slide")) == [
        "ppt/slides/slide1.xml", "ppt/slides/slide2.xml"]


def test_apply_strips_hidden_channels(env):
    make_skeleton(env)
    names = zip_names(env["out"])
    for forbidden in ("ppt/media/", "ppt/charts/", "ppt/embeddings/", "ppt/diagrams/", "ppt/notesSlides/",
                      "ppt/notesMasters/", "ppt/comments", "ppt/commentAuthors", "customXml/", "ppt/tags/",
                      "docProps/custom.xml", "docProps/thumbnail"):
        assert not any(n.startswith(forbidden) for n in names), forbidden
    blob = zip_text(env["out"])
    for secret in ("星云", "华兴", "张三", "李四", "王总", "intra.example.com", "Project-Nebula", "hlinkClick",
                   "hyperlink", "20480", "888", "批注"):
        assert secret not in blob, secret

    out = Presentation(env["out"])
    cp = out.core_properties
    assert (cp.author, cp.title, cp.last_modified_by, cp.subject, cp.keywords, cp.comments, cp.category) == (
        "", "", "", "", "", "", "")
    assert not out.slides[0].has_notes_slide
    # 母版：文字被清空、图片被换成矩形
    master_xml = etree.tostring(out.slide_master._element).decode()
    assert "内部资料" not in master_xml and "<p:pic" not in master_xml
    # 未被引用的版式被丢弃（默认模板有 11 个）
    assert len(out.slide_layouts) < 11
    # 输出能再次打开，且 python-pptx 能遍历全部形状
    assert sum(1 for s in out.slides for _ in s.shapes) > 0


def test_apply_respects_slide_selection_and_order(env):
    make_skeleton(env, slides="3,1")
    out = Presentation(env["out"])
    assert len(out.slides) == 2
    first_texts = [sh.text_frame.text for sh in out.slides[0].shapes if sh.has_text_frame]
    assert first_texts and not any("吞吐" in t for t in first_texts)
    # 第 3 页（现在在最前面）只有 2 个文本框 + 1 个图片矩形
    assert len(list(out.slides[0].shapes)) == 3
    assert len(list(out.slides[1].shapes)) > 5
    # 未选的第 2 页在包里不应残留
    assert "机密" not in zip_text(env["out"])

    make_skeleton({**env, "work": env["tmp"] / "work2", "out": env["tmp"] / "o2.pptx"}, slides="2")
    assert len(Presentation(env["tmp"] / "o2.pptx").slides) == 1


@pytest.mark.skipif(sk.find_soffice() is None, reason="需要 soffice")
def test_render_produces_images(env):
    make_skeleton(env)
    assert run("render", env["out"], "-o", env["tmp"] / "png") == 0
    produced = list((env["tmp"] / "png").iterdir())
    assert produced


# --------------------------------------------------------------------------- check

def test_check_passes_on_clean_skeleton(env, capsys):
    make_skeleton(env)
    assert check(env) == 0
    assert "通过" in capsys.readouterr().out


def test_check_flags_original_substring(env, capsys):
    make_skeleton(env, overrides={"混合精度": "核心技术：混合精度训练框架的升级版"})
    assert check(env) == 1
    out = capsys.readouterr().out
    assert "[a]" in out and "混合精度训练框架" in out
    assert "[g]" not in out


def test_check_flags_identical_text(env, capsys):
    make_skeleton(env, overrides={"混合精度": "核心技术：混合精度训练框架"})
    assert check(env) == 1
    out = capsys.readouterr().out
    assert "[g]" in out and "[a]" in out


def test_check_flags_three_digit_number(env, capsys):
    make_skeleton(env, overrides={"20480": "吞吐量 20480"})
    assert check(env) == 1
    out = capsys.readouterr().out
    assert "[b]" in out and "20480" in out


def test_check_allows_short_numbers(env):
    make_skeleton(env, overrides={"3.2": "提升 N 倍，下降 45%"})
    assert check(env) == 0


def test_check_flags_sensitive_word_case_insensitive(env, capsys):
    make_skeleton(env, overrides={"客户": "客户 project-nebula 验证"})
    assert check(env) == 1
    out = capsys.readouterr().out
    assert "[c]" in out and "Project-Nebula" in out


def test_check_flags_sensitive_word_split_across_runs(env, capsys):
    # 同一段里两个格式不同的 run 各写一半：「星」+「云」
    make_skeleton(env, overrides={"100M": "星", "KV Cache": "云"})
    assert check(env) == 1
    assert "[c]" in capsys.readouterr().out


def test_check_flags_sensitive_word_in_manifest(env, capsys):
    make_skeleton(env)
    man = env["work"] / "manifest.yaml"
    man.write_text(man.read_text(encoding="utf-8").replace("通用示意页", "华兴银行专用页", 1), encoding="utf-8")
    assert check(env) == 1
    assert "manifest.yaml" in capsys.readouterr().out


def test_check_flags_manifest_problems(env, capsys):
    make_skeleton(env)
    man = env["work"] / "manifest.yaml"
    man.write_text(man.read_text(encoding="utf-8").replace('"通用示意页"', '""', 1), encoding="utf-8")
    assert check(env) == 1
    assert "[h]" in capsys.readouterr().out


def test_check_allow_word_is_reported_and_masks(env, capsys):
    make_skeleton(env, overrides={"混合精度": "核心技术：混合精度训练框架的升级版"})
    assert check(env, "--allow", "核心技术：混合精度训练框架") == 0  # 被允许的片段整体屏蔽
    out = capsys.readouterr().out
    assert "已允许的词" in out and "混合精度训练框架" in out


def test_check_flags_tampered_package(env, capsys):
    make_skeleton(env)
    bad = env["tmp"] / "bad.pptx"

    def edit(name, data):
        if name == "docProps/core.xml":
            return data.replace(b"</cp:coreProperties>", "<dc:creator>张三</dc:creator></cp:coreProperties>".encode())
        if name == "ppt/slides/_rels/slide1.xml.rels":
            return data.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                                b'officeDocument/2006/relationships/hyperlink" Target="http://x.example/" '
                                b'TargetMode="External"/></Relationships>')
        return data

    with zipfile.ZipFile(env["out"]) as zin, zipfile.ZipFile(bad, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item.filename, edit(item.filename, zin.read(item.filename)))
        zout.writestr("ppt/media/image9.png", b"x")
        zout.writestr("ppt/notesSlides/notesSlide1.xml", b"<x/>")
    env = {**env, "out": bad}
    assert check(env) == 1
    out = capsys.readouterr().out
    for tag in ("[d]", "[e]", "[f]"):
        assert tag in out
    assert "ppt/media/image9.png" in out and "docProps/core.xml" in out


def test_check_does_not_scan_texts_json_but_catches_leak_in_replacements(env, capsys):
    make_skeleton(env)
    reps = json.loads((env["work"] / "replacements.json").read_text(encoding="utf-8"))
    first = next(iter(reps))
    reps[first] = "吞吐提升 3.2 倍，延迟下降"  # 来自原文，且 OUT 没用到 -> 仍应在 replacements 里被发现
    (env["work"] / "replacements.json").write_text(json.dumps(reps, ensure_ascii=False), encoding="utf-8")
    assert check(env) == 1
    assert "replacements.json" in capsys.readouterr().out


# --------------------------------------------------------------------------- 怪结构

def build_tricky(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tree = s.shapes._spTree
    _img, img_rid = s.part.get_or_add_image_part(png_bytes())
    vid_rid = s.part.relate_to("file:///C:/内部/视频.mp4", RT.VIDEO, is_external=True)
    ole = Part(PackURI("/ppt/embeddings/oleObject1.bin"), "application/vnd.openxmlformats-officedocument.oleObject",
               prs.part.package, "secret 星云".encode("utf-8"))
    ole_rid = s.part.relate_to(ole, RT.OLE_OBJECT)

    def sp(id_, name, x, text_xml, extra_sppr="", ph=""):
        return (
            '<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr txBox="1"/><p:nvPr>%s</p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="%d" y="500000"/><a:ext cx="3000000" cy="800000"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>%s</p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>%s'
            "</p:txBody></p:sp>" % (id_, name, ph, x, extra_sppr, text_xml))

    ns = ('xmlns:a="%(a)s" xmlns:p="%(p)s" xmlns:r="%(r)s" '
          'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"' % NS)
    fld = '<a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum"><a:rPr lang="zh-CN"/><a:t>%s</a:t></a:fld>'
    frag = [
        # AlternateContent 包着一个文本形状
        '<mc:AlternateContent><mc:Choice Requires="p14">'
        + sp(10, "替代", 300000, '<a:p><a:r><a:rPr lang="zh-CN"/><a:t>替代内容里的文字内容</a:t></a:r></a:p>')
        + "</mc:Choice><mc:Fallback>"
        + sp(11, "回退", 300000, '<a:p><a:r><a:rPr lang="zh-CN"/><a:t>回退里的文字内容</a:t></a:r></a:p>')
        + "</mc:Fallback></mc:AlternateContent>",
        # 图片填充的形状
        sp(12, "图填充", 4000000, '<a:p><a:r><a:rPr lang="zh-CN"/><a:t>图片底纹上的文字</a:t></a:r></a:p>',
           '<a:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>' % img_rid),
        # 软换行 + 域；第一段两个格式不同的行（中间是换行），第二段同格式
        sp(13, "多段", 7000000,
           '<a:p><a:r><a:rPr lang="zh-CN" b="1"/><a:t>加粗的标题行</a:t></a:r><a:br><a:rPr lang="zh-CN"/></a:br>'
           '<a:r><a:rPr lang="zh-CN" b="0"/><a:t>常规的说明行</a:t></a:r></a:p>'
           '<a:p><a:r><a:rPr lang="zh-CN"/><a:t>第一行文字</a:t></a:r><a:br><a:rPr lang="zh-CN"/></a:br>'
           '<a:r><a:rPr lang="zh-CN" dirty="0"/><a:t>第二行文字</a:t></a:r>' + fld % "7" + "</a:p>"),
        # 页脚、页码占位符
        sp(14, "页脚", 100000, '<a:p><a:r><a:rPr lang="zh-CN"/><a:t>某某公司内部页脚</a:t></a:r></a:p>',
           ph='<p:ph type="ftr" sz="quarter" idx="11"/>'),
        sp(15, "页码", 5000000, "<a:p>" + fld % "‹#›" + "</a:p>", ph='<p:ph type="sldNum" sz="quarter" idx="12"/>'),
        # 带外部链接的视频（椭圆形状）
        '<p:pic><p:nvPicPr><p:cNvPr id="16" name="视频"/><p:cNvPicPr/><p:nvPr><a:videoFile r:link="%s"/></p:nvPr>'
        '</p:nvPicPr><p:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/></a:stretch></p:blipFill><p:spPr>'
        '<a:xfrm><a:off x="500000" y="3000000"/><a:ext cx="2000000" cy="1200000"/></a:xfrm>'
        '<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom></p:spPr></p:pic>' % (vid_rid, img_rid),
        # OLE 对象
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="17" name="OLE"/><p:cNvGraphicFramePr/><p:nvPr/>'
        '</p:nvGraphicFramePr><p:xfrm><a:off x="3000000" y="3000000"/><a:ext cx="2000000" cy="1200000"/></p:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">'
        '<p:oleObj r:id="%s" name="Worksheet"/></a:graphicData></a:graphic></p:graphicFrame>' % ole_rid,
    ]
    for f in frag:
        tree.append(etree.fromstring("<root %s>%s</root>" % (ns, f))[0])
    # 背景图
    s._element.find("p:cSld", NS).insert(0, etree.fromstring(
        '<p:bg xmlns:a="%(a)s" xmlns:p="%(p)s" xmlns:r="%(r)s"><p:bgPr><a:blipFill><a:blip r:embed="' % NS
        + img_rid + '"/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>'))
    # 演示文稿级：节名
    pres = prs.part._element
    sld_id = pres.find("p:sldIdLst/p:sldId", NS).get("id")
    pres.append(etree.fromstring(
        '<p:extLst xmlns:p="%s"><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
        '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"><p14:section '
        'name="星云章节" id="{AAAAAAAA-0000-0000-0000-000000000000}"><p14:sldIdLst><p14:sldId id="%s"/></p14:sldIdLst>'
        "</p14:section></p14:sectionLst></p:ext></p:extLst>" % (NS["p"], sld_id)))
    prs.save(path)


def test_tricky_structures(tmp_path):
    orig = tmp_path / "t.pptx"
    build_tricky(orig)
    work, out = tmp_path / "w", tmp_path / "o.pptx"
    assert run("extract", orig, "--slides", "1", "-o", work) == 0
    items = json.loads((work / "texts.json").read_text(encoding="utf-8"))["items"]
    texts = [e["text"] for e in items]
    # AlternateContent 取 Choice；页脚/页码占位符不提取；格式不同的两行拆成两条，同格式的软换行留在一条里
    assert "替代内容里的文字内容" in texts and "回退里的文字内容" not in texts
    assert "某某公司内部页脚" not in texts and "‹#›" not in texts
    assert "加粗的标题行" in texts and "常规的说明行" in texts
    assert "第一行文字\n第二行文字7" in texts
    assert "图片底纹上的文字" in texts

    fill(work)
    assert run("apply", orig, work, "-o", out) == 0
    blob = zip_text(out)
    for gone in ("AlternateContent", "替代", "回退", "星云", "页脚", "videoFile", "oleObj", "sectionLst",
                 "视频", "blip r:embed", "C:/内部"):
        assert gone not in blob, gone
    assert not [n for n in zip_names(out) if n.startswith(("ppt/media", "ppt/embeddings"))]
    prs = Presentation(out)
    shapes = list(prs.slides[0].shapes)
    # 组与组之间的软换行原样保留；同一组内的换行随替换文字走（没写 \\n 就没有）
    assert etree.tostring(shapes[2]._element).decode().count("<a:br") == 1
    # 视频变灰色矩形（保留椭圆形状），OLE 变「对象」矩形
    autos = [sh for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert any(sh.text_frame.text == "对象" for sh in autos)
    assert any(sh.auto_shape_type == MSO_SHAPE.OVAL for sh in autos)
    # 图片填充和背景图都变成纯色
    assert "blipFill" not in etree.tostring(prs.slides[0]._element).decode()
    assert sk.main(["check", str(orig), str(out), str(work)]) == 0


def add_second_master(prs):
    """复制出第二个母版（带两个版式和主题），母版里放一段文字，用来验证未使用的母版会被丢弃。"""
    from pptx.parts.slide import SlideLayoutPart, SlideMasterPart

    pkg = prs.part.package
    m1 = prs.slide_master.part
    m2 = SlideMasterPart.load(PackURI("/ppt/slideMasters/slideMaster2.xml"), m1.content_type, pkg, m1.blob)
    n = 20
    for rid, rel in sorted(m1.rels.items(), key=lambda kv: int(kv[0][3:])):
        if rel.reltype == RT.SLIDE_LAYOUT:
            n += 1
            lp = SlideLayoutPart.load(PackURI("/ppt/slideLayouts/slideLayout%d.xml" % n),
                                      rel.target_part.content_type, pkg, rel.target_part.blob)
            lp.relate_to(m2, RT.SLIDE_MASTER)
            assert m2.relate_to(lp, RT.SLIDE_LAYOUT) == rid
        elif rel.reltype == RT.THEME:
            th = Part.load(PackURI("/ppt/theme/theme9.xml"), rel.target_part.content_type, pkg, rel.target_part.blob)
            assert m2.relate_to(th, RT.THEME) == rid
    m2._element.find("p:cSld/p:spTree", NS).append(etree.fromstring(
        '<p:sp xmlns:a="%(a)s" xmlns:p="%(p)s"><p:nvSpPr><p:cNvPr id="77" name="第二母版文字"/><p:cNvSpPr txBox="1"/>'
        '<p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="1" y="1"/><a:ext cx="999" cy="999"/></a:xfrm></p:spPr>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>第二母版里的星云机密</a:t></a:r></a:p></p:txBody></p:sp>' % NS))
    rid = prs.part.relate_to(m2, RT.SLIDE_MASTER)
    prs.part._element.find("p:sldMasterIdLst", NS).append(etree.fromstring(
        '<p:sldMasterId xmlns:p="%(p)s" xmlns:r="%(r)s" id="2147484000" r:id="' % NS + rid + '"/>'))


def test_unused_masters_and_layouts_are_dropped(tmp_path):
    orig = tmp_path / "m.pptx"
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "标题文字内容"
    add_second_master(prs)
    prs.save(orig)
    assert len(Presentation(orig).slide_masters) == 2
    assert "第二母版里的星云机密" in zip_text(orig)

    work, out = tmp_path / "w", tmp_path / "o.pptx"
    assert run("extract", orig, "--slides", "1", "-o", work) == 0
    fill(work)
    assert run("apply", orig, work, "-o", out) == 0
    res = Presentation(out)
    assert len(res.slide_masters) == 1 and len(res.slide_layouts) == 1
    names = zip_names(out)
    assert len([n for n in names if n.startswith("ppt/slideMasters/slideMaster")]) == 1
    assert len([n for n in names if n.startswith("ppt/slideLayouts/slideLayout")]) == 1
    assert len([n for n in names if n.startswith("ppt/theme/theme")]) == 1
    assert "第二母版" not in zip_text(out)
    assert sk.main(["check", str(orig), str(out), str(work)]) == 0


def test_check_ignores_own_placeholder_labels(tmp_path):
    orig = tmp_path / "l.pptx"
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "图片"  # 原文里恰好有一段文字就叫「图片」
    s.shapes.add_picture(png_bytes(), Inches(1), Inches(3), Inches(2), Inches(2))
    prs.save(orig)
    work, out = tmp_path / "w", tmp_path / "o.pptx"
    run("extract", orig, "--slides", "1", "-o", work)
    fill(work)
    run("apply", orig, work, "-o", out)
    assert sk.main(["check", str(orig), str(out), str(work)]) == 0
