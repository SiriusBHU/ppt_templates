import pytest
from pptx import Presentation

from factories import deck_of, table_template
from ppt_templates.errors import PptError
from ppt_templates.render import render_deck
from ppt_templates.slots import prepare_pages
from ppt_templates.validate import validate_repo


def test_table_renders_native_table_with_header_and_highlight(make_repo, tmp_path):
    template, example = table_template()
    repo = make_repo({"common.overview-table": (template, example)})
    assert validate_repo(repo).errors == []
    path = tmp_path / "t.pptx"
    render_deck(deck_of("common.overview-table", example), repo).save(path)
    frame = next(sh for sh in Presentation(path).slides[0].shapes if sh.has_table)
    table = frame.table
    assert (len(table.rows), len(table.columns)) == (4, 4)
    assert frame._element.xpath(".//a:tableStyleId/text()") == ["{2D5ABB26-0587-4C30-8999-92F81FD0307C}"]
    assert table.columns[0].width == 200 * 9525
    assert str(table.cell(0, 1).fill.fore_color.rgb) == "283593"   # 表头
    assert str(table.cell(0, 2).fill.fore_color.rgb) == "C00000"   # 高亮列的表头
    hot = table.cell(1, 2)
    assert str(hot.fill.fore_color.rgb) == "FFF1E6"
    assert hot.text_frame.paragraphs[0].runs[0].font.bold
    first = table.cell(1, 0)
    assert first.text == "带宽" and first.text_frame.paragraphs[0].runs[0].font.bold
    runs = frame._element.xpath(".//a:r")
    assert frame._element.xpath(".//a:r/a:rPr/a:ea/@typeface") == ["微软雅黑"] * len(runs)


def _prepare_error(repo, template, data) -> str:
    with pytest.raises(PptError) as exc:
        prepare_pages(deck_of(template["id"], data), repo)
    return str(exc.value)


def test_table_row_length_must_match_columns(make_repo):
    template, example = table_template()
    example["matrix"]["rows"][1] = ["延迟"]
    repo = make_repo({"common.overview-table": (template, example)})
    assert ("第 1 页 common.overview-table  /pages/0/data/matrix/rows/1  第 2 行有 1 格，表头有 4 列（组件 table）"
            in _prepare_error(repo, template, example))
    assert ("templates/common/overview-table/example.yaml  /matrix/rows/1  第 2 行有 1 格，表头有 4 列（组件 table）"
            in validate_repo(repo).errors)


def test_deck_data_is_checked_against_the_component_schema_even_if_the_slot_is_loose(make_repo):
    template, example = table_template()
    template["slots"]["properties"]["matrix"] = {"type": "object"}
    repo = make_repo({"common.overview-table": (template, example)})
    bad = {"title": "表", "matrix": {"columns": ["a"], "rows": "x"}}
    assert ("第 1 页 common.overview-table  /pages/0/data/matrix/rows  类型应为 array（组件 table）"
            in _prepare_error(repo, template, bad))
    assert "缺少必填字段 rows（组件 table）" in _prepare_error(repo, template, {"title": "表", "matrix": {"columns": ["a"]}})


def test_component_inside_grid_reports_item_pointer(make_repo):
    template, example = table_template()
    template["slots"]["properties"]["matrix"] = {"type": "array", "maxItems": 2, "items": {"type": "object"}}
    template["layout"] = [{"type": "grid", "box": [0, 0, 1216, 360], "bind": "matrix", "columns": 2, "gap": 16, "item": [
        {"type": "component", "component": "table", "box": [0, 0, "100%", "100%"], "bind": "."}]}]
    repo = make_repo({"common.overview-table": (template, example)})
    cell = {"columns": ["a", "b"], "rows": [["x", "y"]]}
    data = {"title": "表", "matrix": [cell, {**cell, "rows": [["x"]]}]}
    assert ("/pages/0/data/matrix/1/rows/0  第 1 行有 1 格，表头有 2 列（组件 table）"
            in _prepare_error(repo, template, data))
