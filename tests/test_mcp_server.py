import asyncio
import io
import shutil

import pytest
from mcp import Client
from mcp.server.mcpserver import Image
from PIL import Image as PilImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from conftest import TEMPLATE_ID
from factories import deck_of, image_template, mini_template
from ppt_templates import mcp_server
from ppt_templates.errors import PptError
from ppt_templates.loader import load_example

needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="需要 LibreOffice (soffice)")
TOOLS = {"get_taxonomy", "list_templates", "get_template", "validate_deck", "render_deck", "preview_deck"}


@pytest.fixture
def in_repo(root, monkeypatch):
    monkeypatch.chdir(root)


@pytest.fixture
def mini_repo(make_repo, monkeypatch):
    """含一个 common.summary-note 模板的临时仓库，当前目录切到仓库里（out/ 也写在这里）。"""
    repo = make_repo({"common.summary-note": mini_template()})
    monkeypatch.chdir(repo)
    return repo


def mini_deck(**data):
    return deck_of("common.summary-note", {"title": "总结", "note": "一句话说清楚结论", **data})


def test_get_taxonomy(in_repo):
    tax = mcp_server.get_taxonomy()
    assert {"trend", "system", "common"} <= {t["id"] for t in tax["topics"]}
    assert all(t["name"] for t in tax["topics"])
    problem = next(r for r in tax["roles"] if r["id"] == "problem")
    assert problem["question"] == "哪里不行，代价多大"


def test_list_templates_filters_like_the_cli(in_repo):
    assert [e["id"] for e in mcp_server.list_templates(role="cover")] == ["common.cover-basic"]
    assert TEMPLATE_ID in [e["id"] for e in mcp_server.list_templates(topic="hardware")]  # also_fits
    assert TEMPLATE_ID not in [e["id"] for e in mcp_server.list_templates(topic="hardware", role="cover")]
    assert all(e["density"] == "read" for e in mcp_server.list_templates(density="read"))
    assert [e["id"] for e in mcp_server.list_templates(tag="封面")] == ["common.cover-basic"]
    assert len(mcp_server.list_templates()) > 20


def test_get_template_returns_contract_without_layout(in_repo, root):
    t = mcp_server.get_template(TEMPLATE_ID)
    assert t["id"] == TEMPLATE_ID and t["slots"]["type"] == "object"
    assert t["guidance"] and t["when_to_use"] and t["when_not_to_use"] and t["roles"] == ["solution", "feature"]
    assert t["example"] == load_example(root, TEMPLATE_ID)
    assert "layout" not in t


@pytest.mark.parametrize("bad", ["system.nope", "../etc.passwd", "system.solution-arch-overview/../../x", "nodot", ""])
def test_get_template_rejects_unknown_or_malformed_id(in_repo, bad):
    with pytest.raises(PptError, match=r"模板 .* 不存在|模板 id"):
        mcp_server.get_template(bad)


def test_validate_deck_ok(in_repo, deck):
    assert mcp_server.validate_deck(deck) == {"ok": True, "errors": []}


def test_validate_deck_reports_chinese_errors(mini_repo):
    result = mcp_server.validate_deck(mini_deck(note="长" * 61))
    assert result["ok"] is False
    assert "第 1 页 common.summary-note  /pages/0/data/note  长度 61，超过上限 60" in result["errors"]


def test_validate_deck_reports_structure_errors(mini_repo):
    result = mcp_server.validate_deck({"pages": []})
    assert result["ok"] is False and "/pages  共 0 项，低于下限 1" in result["errors"]


def test_validate_deck_checks_the_style_exists(mini_repo):
    deck = mini_deck()
    deck["style"] = "nope"
    result = mcp_server.validate_deck(deck)
    assert result["ok"] is False and any("文件不存在" in e and "nope" in e for e in result["errors"])


def test_style_cannot_escape_the_styles_directory(mini_repo):
    deck = mini_deck()
    deck["style"] = "../schema"
    result = mcp_server.validate_deck(deck)
    assert result["ok"] is False and any("/style  " in e and "不符合格式" in e for e in result["errors"])
    with pytest.raises(PptError, match="/style"):
        mcp_server.render_deck(deck)


def test_render_deck_writes_into_out_and_returns_absolute_path(mini_repo):
    path = mcp_server.render_deck(mini_deck())
    assert path == str((mini_repo / "out" / "deck.pptx").resolve())
    assert len(Presentation(path).slides) == 1
    custom = mcp_server.render_deck(mini_deck(), filename="第一版.pptx")
    assert custom == str((mini_repo / "out" / "第一版.pptx").resolve())
    assert sorted(p.name for p in (mini_repo / "out").iterdir()) == ["deck.pptx", "第一版.pptx"]


@pytest.mark.parametrize("bad", ["../x.pptx", "a/b.pptx", "/tmp/x.pptx", "x.txt", "x", ".pptx", "..\\x.pptx", "", "x.pptx.txt"])
def test_render_deck_rejects_unsafe_filenames(mini_repo, bad):
    with pytest.raises(PptError, match="文件名"):
        mcp_server.render_deck(mini_deck(), filename=bad)
    assert not (mini_repo / "out").exists()
    assert not (mini_repo.parent / "x.pptx").exists()


def test_render_deck_validates_first(mini_repo):
    with pytest.raises(PptError, match="长度 61，超过上限 60"):
        mcp_server.render_deck(mini_deck(note="长" * 61))
    assert not (mini_repo / "out" / "deck.pptx").exists()


@needs_soffice
def test_preview_deck_returns_1280_wide_pngs(mini_repo):
    images = mcp_server.preview_deck(mini_deck())
    assert len(images) == 1 and isinstance(images[0], Image)
    with PilImage.open(io.BytesIO(images[0].data)) as img:
        assert img.size == (1280, 720)


def test_preview_deck_without_soffice_is_a_clear_error(mini_repo, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(PptError, match="找不到 soffice"):
        mcp_server.preview_deck(mini_deck())


def test_preview_deck_validates_first(mini_repo):
    with pytest.raises(PptError, match="长度 61，超过上限 60"):
        mcp_server.preview_deck(mini_deck(note="长" * 61))


# ---- 图片路径：相对仓库根目录，且只能引用仓库之内的文件 ----

PHOTO_ID = "common.overview-photo"


@pytest.fixture
def photo_repo(make_repo, monkeypatch, tmp_path):
    """含 image 模板的临时仓库（tmp_path/repo）；仓库内有 assets/in.png，仓库外有 tmp_path/out.png 和指向它的符号链接 link.png。"""
    repo = make_repo({PHOTO_ID: image_template()})
    for path in (repo / "assets" / "in.png", tmp_path / "out.png"):
        path.parent.mkdir(exist_ok=True)
        PilImage.new("RGB", (40, 20), "red").save(path)
    (repo / "link.png").symlink_to(tmp_path / "out.png")
    monkeypatch.chdir(repo)
    return repo


def photo_deck(photo):
    return deck_of(PHOTO_ID, {"title": "图", "photo": photo})


def test_images_inside_the_repo_resolve_against_the_repo_root(photo_repo, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # 工作目录不影响：相对路径从仓库根起算（find_repo_root 靠包所在位置找到仓库）
    monkeypatch.setattr(mcp_server, "find_repo_root", lambda: photo_repo)
    for photo in ("assets/in.png", str(photo_repo / "assets" / "in.png")):
        assert mcp_server.validate_deck(photo_deck(photo)) == {"ok": True, "errors": []}
    prs = Presentation(mcp_server.render_deck(photo_deck("assets/in.png")))
    assert len([sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1


@pytest.mark.parametrize("photo", ["../out.png", "link.png", "OUTSIDE", "/etc/passwd.png"])
def test_images_outside_the_repo_are_rejected(photo_repo, tmp_path, photo):
    photo = str(tmp_path / "out.png") if photo == "OUTSIDE" else photo
    result = mcp_server.validate_deck(photo_deck(photo))
    assert result["ok"] is False
    assert any("/pages/0/data/photo  图片" in e and "不在允许的目录内" in e for e in result["errors"])
    with pytest.raises(PptError, match="不在允许的目录内"):
        mcp_server.render_deck(photo_deck(photo))
    assert not (photo_repo / "out").exists()
    with pytest.raises(PptError, match="不在允许的目录内"):
        mcp_server.preview_deck(photo_deck(photo))


def test_images_in_the_repo_still_need_a_supported_extension_and_must_exist(photo_repo):
    (photo_repo / "assets" / "notes.txt").write_text("x", encoding="utf-8")
    errors = mcp_server.validate_deck(photo_deck("assets/notes.txt"))["errors"]
    assert any("格式不支持" in e for e in errors)
    errors = mcp_server.validate_deck(photo_deck("assets/none.png"))["errors"]
    assert any("图片 assets/none.png 不存在" in e for e in errors)


@needs_soffice
def test_preview_deck_draws_a_placeholder_when_no_image_is_given(photo_repo):
    assert len(mcp_server.preview_deck(deck_of(PHOTO_ID, {"title": "图"}))) == 1


# ---- 经 MCP 协议（内存连接）调用 ----

def call(name, arguments=None):
    async def run():
        async with Client(mcp_server.mcp) as client:
            return await client.call_tool(name, arguments or {})

    return asyncio.run(run())


def test_server_lists_all_tools_with_chinese_descriptions(in_repo):
    async def run():
        async with Client(mcp_server.mcp) as client:
            return (await client.list_tools()).tools

    tools = asyncio.run(run())
    assert {t.name for t in tools} == TOOLS
    assert all(any("一" <= c <= "鿿" for c in t.description) for t in tools)
    render = next(t for t in tools if t.name == "render_deck")
    assert set(render.input_schema["properties"]) == {"deck", "filename"} and render.input_schema["required"] == ["deck"]


def test_call_tool_over_mcp_returns_structured_results(in_repo, deck):
    result = call("validate_deck", {"deck": deck})
    assert not result.is_error and result.structured_content == {"ok": True, "errors": []}
    listed = call("list_templates", {"role": "cover"})
    assert [e["id"] for e in listed.structured_content["result"]] == ["common.cover-basic"]
    got = call("get_template", {"template_id": TEMPLATE_ID})
    assert got.structured_content["id"] == TEMPLATE_ID and "layout" not in got.structured_content


def test_call_tool_over_mcp_surfaces_chinese_errors(mini_repo):
    result = call("render_deck", {"deck": mini_deck(), "filename": "../x.pptx"})
    assert result.is_error and "文件名" in result.content[0].text
    assert not (mini_repo.parent / "x.pptx").exists()
    missing = call("get_template", {"template_id": "system.nope"})
    assert missing.is_error and "模板 system.nope 不存在" in missing.content[0].text


def test_render_over_mcp_writes_the_file(mini_repo):
    result = call("render_deck", {"deck": mini_deck()})
    assert not result.is_error and (mini_repo / "out" / "deck.pptx").is_file()


@needs_soffice
def test_preview_over_mcp_returns_image_content(mini_repo):
    result = call("preview_deck", {"deck": mini_deck()})
    assert not result.is_error and [c.type for c in result.content] == ["image"]
    assert result.content[0].mime_type == "image/png"
