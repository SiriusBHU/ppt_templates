import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from factories import deck_of, image_template, rows_template
from ppt_templates.cli import main
from ppt_templates.errors import PptError
from ppt_templates.images import ImageSource
from ppt_templates.render import render_deck
from ppt_templates.slots import prepare_pages
from ppt_templates.validate import validate_repo

PHOTO = "common.overview-photo"
ROWS = "common.overview-rows"


def make_png(path, size=(200, 100), color="red"):
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)
    return path


def _px(shape):
    return tuple(v / 9525 for v in (shape.left, shape.top, shape.width, shape.height))


def _pictures(prs):
    return [sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]


def _render(make_repo, template, data, images=None, **kw):
    repo = make_repo({template["id"]: (template, None)}, **kw)
    return render_deck(deck_of(template["id"], data), repo, images)


# ---- 占位框 ----

@pytest.mark.parametrize("data", [{"title": "t"}, {"title": "t", "photo": ""}])
def test_missing_or_empty_path_draws_a_placeholder_with_guide_text(make_repo, data):
    template, _ = image_template()
    prs = _render(make_repo, template, data)
    assert _pictures(prs) == []
    (box,) = _placeholders(prs)
    assert _px(box) == (32, 128, 400, 400)  # body 区原点 (32, 128)
    assert str(box.fill.fore_color.rgb) == "F3F5F8"  # surface
    assert not box._element.xpath(".//a:outerShdw") and box.line.fill.type == 5  # 无阴影、无描边（MSO_FILL.BACKGROUND）
    paragraph = box.text_frame.paragraphs[0]
    assert paragraph.alignment == PP_ALIGN.CENTER and box.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
    run = paragraph.runs[0]
    assert run.font.size.pt == 10 and not run.font.bold and str(run.font.color.rgb) == "000000"  # detail + text


def _placeholders(prs):
    return [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text == "放一张示意图"]


@pytest.mark.parametrize("when, shown", [("title", True), ("photo", False)])
def test_when_applies_to_the_placeholder_too(make_repo, when, shown):
    template, _ = image_template(when=when)
    assert bool(_placeholders(_render(make_repo, template, {"title": "t"}))) is shown


# ---- 插入图片 ----

def test_contain_keeps_the_aspect_ratio_and_centers(make_repo, tmp_path):
    png = make_png(tmp_path / "wide.png", (200, 100))  # 2:1 放进 400×400 → 400×200，上下各空 100
    template, _ = image_template()
    prs = _render(make_repo, template, {"title": "t", "photo": str(png)})
    (pic,) = _pictures(prs)
    assert _px(pic) == (32, 128 + 100, 400, 200)
    assert (pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom) == (0, 0, 0, 0)
    assert _placeholders(prs) == []


def test_contain_with_a_tall_image_leaves_space_left_and_right(make_repo, tmp_path):
    png = make_png(tmp_path / "tall.png", (100, 200))
    template, _ = image_template()
    (pic,) = _pictures(_render(make_repo, template, {"title": "t", "photo": str(png)}))
    assert _px(pic) == (32 + 100, 128, 200, 400)


def test_cover_fills_the_box_and_crops_the_overflow(make_repo, tmp_path):
    png = make_png(tmp_path / "wide.png", (200, 100))  # 铺满 400×400 要放大到 800×400，左右各裁掉 1/4
    template, _ = image_template(fit="cover")
    (pic,) = _pictures(_render(make_repo, template, {"title": "t", "photo": str(png)}))
    assert _px(pic) == (32, 128, 400, 400)
    assert pic.crop_left == pytest.approx(0.25) and pic.crop_right == pytest.approx(0.25)
    assert pic.crop_top == 0 and pic.crop_bottom == 0


def test_cover_with_a_tall_image_crops_top_and_bottom(make_repo, tmp_path):
    png = make_png(tmp_path / "tall.png", (100, 200))
    template, _ = image_template(fit="cover")
    (pic,) = _pictures(_render(make_repo, template, {"title": "t", "photo": str(png)}))
    assert pic.crop_top == pytest.approx(0.25) and pic.crop_bottom == pytest.approx(0.25)
    assert pic.crop_left == 0 and pic.crop_right == 0


def test_picture_is_scaled_with_the_body_area(make_repo, tmp_path):
    png = make_png(tmp_path / "square.png", (100, 100))
    template, _ = image_template(box=[0, 0, 600, 560])  # compact body：纵向放大 600/560
    prs = _render(make_repo, template, {"title": "t", "photo": str(png)}, compact_body=True)
    (pic,) = _pictures(prs)
    assert _px(pic) == pytest.approx((32, 88, 600, 600))  # 600×560 的参考框 → 600×600，方图撑满


# ---- 路径解析 ----

def test_relative_path_resolves_against_the_given_base_not_the_cwd(make_repo, tmp_path, monkeypatch):
    make_png(tmp_path / "deck-dir" / "img" / "a.png")
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    monkeypatch.chdir(elsewhere)
    template, _ = image_template()
    prs = _render(make_repo, template, {"title": "t", "photo": "img/a.png"}, ImageSource(tmp_path / "deck-dir"))
    assert len(_pictures(prs)) == 1


def test_default_base_is_the_current_directory(make_repo, tmp_path, monkeypatch):
    make_png(tmp_path / "work" / "a.png")
    monkeypatch.chdir(tmp_path / "work")
    template, _ = image_template()
    assert len(_pictures(_render(make_repo, template, {"title": "t", "photo": "a.png"}))) == 1


def test_absolute_paths_are_allowed_without_a_confine(make_repo, tmp_path):
    png = make_png(tmp_path / "abs.png")
    template, _ = image_template()
    prs = _render(make_repo, template, {"title": "t", "photo": str(png)}, ImageSource(tmp_path / "other"))
    assert len(_pictures(prs)) == 1


@pytest.mark.parametrize("name", ["a.png", "a.PNG", "a.jpg", "a.jpeg", "a.gif", "a.bmp", "a.tiff", "a.tif"])
def test_supported_extensions(tmp_path, name):
    Image.new("RGB", (8, 8)).save(tmp_path / name)  # PIL 按扩展名选格式
    assert ImageSource(tmp_path).resolve(name) == (tmp_path / name).resolve()


@pytest.mark.parametrize("name, message", [
    ("missing.png", "图片 missing.png 不存在"),
    ("a.webp", "图片 a.webp 的格式不支持，只能用 png / jpg / jpeg / gif / bmp / tif / tiff"),
    ("a.svg", "格式不支持"),
    ("noext", "格式不支持"),
    ("fake.png", "图片 fake.png 无法读取"),
    ("a\0b.png", "无法读取"),
])
def test_bad_image_paths_are_chinese_errors(tmp_path, name, message):
    (tmp_path / "fake.png").write_text("不是图片", encoding="utf-8")
    with pytest.raises(PptError, match=message):
        ImageSource(tmp_path).resolve(name)


def test_a_directory_is_not_an_image(tmp_path):
    (tmp_path / "dir.png").mkdir()
    with pytest.raises(PptError, match="不存在"):
        ImageSource(tmp_path).resolve("dir.png")


def test_symlink_to_a_non_image_is_rejected(tmp_path):
    (tmp_path / "secret.txt").write_text("x", encoding="utf-8")
    (tmp_path / "link.png").symlink_to(tmp_path / "secret.txt")
    with pytest.raises(PptError, match="实际指向 secret.txt，格式不支持"):
        ImageSource(tmp_path).resolve("link.png")


# ---- confine ----

@pytest.fixture
def confined(tmp_path):
    inside = tmp_path / "repo"
    make_png(inside / "assets" / "ok.png")
    make_png(tmp_path / "outside.png")
    (inside / "escape.png").symlink_to(tmp_path / "outside.png")
    return ImageSource(inside, confine=inside), tmp_path


def test_confine_allows_files_inside(confined):
    images, tmp_path = confined
    assert images.resolve("assets/ok.png") == (tmp_path / "repo" / "assets" / "ok.png").resolve()
    assert images.resolve(str(tmp_path / "repo" / "assets" / "ok.png"))  # 绝对路径只要在里面也行


@pytest.mark.parametrize("name", ["../outside.png", "assets/../../outside.png", "escape.png"])
def test_confine_rejects_files_outside_even_via_dotdot_or_symlinks(confined, name):
    images, _ = confined
    with pytest.raises(PptError, match="不在允许的目录内"):
        images.resolve(name)


def test_confine_rejects_absolute_paths_outside(confined):
    images, tmp_path = confined
    with pytest.raises(PptError, match="不在允许的目录内"):
        images.resolve(str(tmp_path / "outside.png"))


def test_confine_checks_before_existence(confined):
    images, _ = confined
    with pytest.raises(PptError, match="不在允许的目录内"):  # 不泄露仓库之外的文件是否存在
        images.resolve("../nope.png")


# ---- 校验 ----

def _prepare(repo, template, data, images=None):
    return prepare_pages(deck_of(template["id"], data), repo, images)


def test_validate_reports_missing_and_bad_images_with_pointers(make_repo, tmp_path):
    template, _ = image_template()
    repo = make_repo({PHOTO: (template, None)})
    with pytest.raises(PptError) as exc:
        _prepare(repo, template, {"title": "t", "photo": "nope.png"}, ImageSource(tmp_path))
    assert f"/pages/0/data/photo  图片 nope.png 不存在（相对路径从 {tmp_path} 起算）" in str(exc.value)
    with pytest.raises(PptError, match="/pages/0/data/photo  图片 a.webp 的格式不支持"):
        _prepare(repo, template, {"title": "t", "photo": "a.webp"}, ImageSource(tmp_path))


def test_validate_accepts_existing_images_and_empty_values(make_repo, tmp_path):
    make_png(tmp_path / "a.png")
    template, _ = image_template()
    repo = make_repo({PHOTO: (template, None)})
    for photo in ("a.png", "", None):
        data = {"title": "t", **({} if photo is None else {"photo": photo})}
        assert _prepare(repo, template, data, ImageSource(tmp_path))


def test_hidden_images_are_not_checked(make_repo, tmp_path):
    template, _ = image_template(when="flag")
    template["slots"]["properties"]["flag"] = {"type": "string"}
    template["layout"].append({"type": "text", "box": [0, 410, 100, 20], "bind": "flag", "style": "body"})
    repo = make_repo({PHOTO: (template, None)})
    assert _prepare(repo, template, {"title": "t", "photo": "nope.png"}, ImageSource(tmp_path))  # flag 为空 → 不渲染也不查


def test_render_rejects_a_bad_path_before_drawing(make_repo, tmp_path):
    template, _ = image_template()
    with pytest.raises(PptError, match="/pages/0/data/photo  图片 nope.png 不存在"):
        _render(make_repo, template, {"title": "t", "photo": "nope.png"}, ImageSource(tmp_path))


def test_repo_validation_checks_example_images_relative_to_the_template_dir(make_repo):
    template, _ = image_template()
    repo = make_repo({PHOTO: (template, {"title": "t", "photo": "pic.png"})})
    errors = validate_repo(repo).errors
    assert any("example.yaml  /photo  图片 pic.png 不存在" in e for e in errors)
    make_png(repo / "templates" / "common" / "overview-photo" / "pic.png")
    assert [e for e in validate_repo(repo).errors if "example.yaml" in e] == []


# ---- grid 里的 image ----

def test_images_inside_nested_grid_items_use_their_own_scope_and_report_deep_pointers(make_repo, tmp_path):
    make_png(tmp_path / "a.png")
    template, _ = rows_template()
    repo = make_repo({ROWS: (template, None)})
    data = {"title": "t", "rows": [
        {"label": "行一", "cards": [{"name": "甲", "pic": "a.png"}]},
        {"label": "行二", "cards": [{"name": "乙"}, {"name": "丙", "pic": "gone.png"}]},
    ]}
    with pytest.raises(PptError) as exc:
        _prepare(repo, template, data, ImageSource(tmp_path))
    assert "/pages/0/data/rows/1/cards/1/pic  图片 gone.png 不存在" in str(exc.value)
    data["rows"][1]["cards"][1]["pic"] = "a.png"
    prs = render_deck(deck_of(ROWS, data), repo, ImageSource(tmp_path))
    assert len(_pictures(prs)) == 2  # 甲、丙有图
    placeholders = [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text == "配图"]
    assert len(placeholders) == 1  # 乙没有图 → 占位框


# ---- CLI ----

def _write_deck(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(deck_of(PHOTO, data), allow_unicode=True), encoding="utf-8")
    return path


def test_cli_render_resolves_images_against_the_deck_directory(make_repo, tmp_path, monkeypatch):
    repo = make_repo({PHOTO: image_template()})
    make_png(tmp_path / "work" / "deck-dir" / "pic.png")
    deck = _write_deck(tmp_path / "work" / "deck-dir" / "deck.yaml", {"title": "t", "photo": "pic.png"})
    monkeypatch.chdir(repo)
    out = tmp_path / "out.pptx"
    assert main(["render", str(deck), "-o", str(out)]) == 0
    assert len(_pictures(Presentation(out))) == 1


def test_cli_validate_deck_reports_missing_images(make_repo, tmp_path, monkeypatch, capsys):
    repo = make_repo({PHOTO: image_template()})
    deck = _write_deck(tmp_path / "work" / "deck.yaml", {"title": "t", "photo": "pic.png"})
    monkeypatch.chdir(repo)
    assert main(["validate", str(deck)]) == 1
    assert "/pages/0/data/photo  图片 pic.png 不存在" in capsys.readouterr().err
    make_png(tmp_path / "work" / "pic.png")
    assert main(["validate", str(deck)]) == 0


def test_cli_preview_of_a_template_id_resolves_example_images_against_the_template_dir(make_repo, monkeypatch, tmp_path):
    template, _ = image_template()
    repo = make_repo({PHOTO: (template, {"title": "t", "photo": "pic.png"})})
    make_png(repo / "templates" / "common" / "overview-photo" / "pic.png")
    monkeypatch.chdir(repo)
    seen = []

    def fake_render_pngs(prs, out_dir, width=1920):
        seen.append(prs)
        out_dir.mkdir(parents=True, exist_ok=True)
        return [out_dir / "slide-01.png"]

    monkeypatch.setattr("ppt_templates.cli.render_pngs", fake_render_pngs)
    assert main(["preview", PHOTO, "-o", str(tmp_path / "pv")]) == 0
    assert len(_pictures(seen[0])) == 1
