import pytest
import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

from factories import deck_of, roadmap_template
from ppt_templates.components.lane_roadmap import _lines
from ppt_templates.errors import PptError
from ppt_templates.render import render_deck
from ppt_templates.slots import prepare_pages
from ppt_templates.validate import validate_repo

ROAD = "product.roadmap-lanes"
PX = 9525


def _renderer(make_repo):
    """返回 render(example=None) -> 第 1 页的全部形状；同一个测试里 make_repo 只能调用一次。"""
    template, default = roadmap_template()
    repo = make_repo({ROAD: (template, default)})
    return lambda example=None: list(render_deck(deck_of(ROAD, example or default), repo).slides[0].shapes)


def _shape(shapes, text):
    return next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.split("\n")[0] == text)


def _box(shapes, text):
    sh = _shape(shapes, text)
    return tuple(round(v / PX, 1) for v in (sh.left, sh.top, sh.width, sh.height))


def test_example_is_valid(make_repo):
    template, example = roadmap_template()
    assert validate_repo(make_repo({ROAD: (template, example)})).errors == []


def test_columns_are_equal_and_items_span_their_periods(make_repo):
    shapes = _renderer(make_repo)()
    # box 原点 = body 左上角 (32, 128)；标签区 80 + 间距 8，剩下 1128 px 分给 4 列，每列 282；条目和时间轴标签两侧各留 3 px
    assert [_box(shapes, p)[:3] for p in ("2026", "2027", "2028", "2029")] == [(123, 128, 276), (405, 128, 276), (687, 128, 276), (969, 128, 276)]
    assert _box(shapes, "芯片甲")[0::2] == (123, 276)                    # 一列
    assert _box(shapes, "芯片乙")[0::2] == (405, 558)                    # 2027~2028 两列
    assert _box(shapes, "控制器")[0::2] == (123, 1122)                   # 2026~2029 整条时间轴：4 × 282 - 6
    assert _box(shapes, "旧时代")[0::2] == (123, 558) and _box(shapes, "新时代")[0::2] == (687, 558)


def test_lane_height_follows_the_content_and_fills_the_body(make_repo):
    _, example = roadmap_template()
    plain = [{"name": f"泳道{n}", "items": [{"name": "条目", "from": "2026"}]} for n in "甲乙"]
    example["roadmap"]["lanes"] += plain
    shapes = _renderer(make_repo)(example)
    label_a, label_b, label_c = _box(shapes, "智算"), _box(shapes, "存储"), _box(shapes, "泳道甲")
    assert label_a[0] == 32 and label_a[2] == 80 and label_a[1] == 128 + 26 + 4        # 时间轴一行 26 + 间距 4
    assert label_a[3] > label_c[3] >= 44 and label_c[3] == label_b[3]                  # 有备注的泳道更高，都没到 140 的上限
    assert label_b[1] == pytest.approx(label_a[1] + label_a[3] + 6, abs=0.2)          # 泳道间距 6
    era = _box(shapes, "旧时代")
    assert era[1] == pytest.approx(_box(shapes, "泳道乙")[1] + _box(shapes, "泳道乙")[3] + 10, abs=0.2)   # 时代行紧跟最后一条泳道
    assert era[1] + era[3] == pytest.approx(128 + 560 - 28, abs=0.5)                  # 富余的高度分完后，刚好留出图例的位置


def test_a_lone_lane_stops_growing_at_the_height_cap(make_repo):
    _, example = roadmap_template()
    example["roadmap"]["lanes"], example["roadmap"]["eras"] = example["roadmap"]["lanes"][1:], []
    assert _box(_renderer(make_repo)(example), "存储")[3] == 140


def test_item_box_sits_above_the_lane_arrow(make_repo):
    shapes = _renderer(make_repo)()
    label, item = _box(shapes, "智算"), _box(shapes, "芯片甲")
    arrows = [sh for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(arrows) == 2                                                           # 每条泳道一条
    first = arrows[0]
    assert round(first.begin_x / PX, 1) == 120 and round(first.end_x / PX, 1) == 120 + 1128   # 横贯整条时间轴
    assert first.begin_y == first.end_y and first.begin_y / PX > item[1] + item[3]    # 在条目之下
    assert first.begin_y / PX < label[1] + label[3]                                   # 还在泳道之内
    assert first._element.xpath(".//a:ln/a:tailEnd/@type") == ["triangle"]
    assert first._element.xpath(".//p:style/a:effectRef/@idx") == ["0"]               # 无阴影


def test_highlight_style_notes_and_legend(make_repo):
    render = _renderer(make_repo)
    shapes = render()
    hot, plain = _shape(shapes, "芯片乙"), _shape(shapes, "芯片甲")
    assert str(hot.fill.fore_color.rgb) == "FFF1E6" and str(hot.line.color.rgb) == "C00000" and hot.line.width == 19050
    assert str(plain.fill.fore_color.rgb) == "FFFFFF" and str(plain.line.color.rgb) == "D1D5DB" and plain.line.width == 9525
    name, *notes = (p.runs[0] for p in hot.text_frame.paragraphs)
    assert [p.runs[0].text for p in hot.text_frame.paragraphs] == ["芯片乙", "IP已确定", "SoC待评审"]
    assert name.font.bold and name.font.size.pt == 12 and str(name.font.color.rgb) == "C00000"
    assert [(r.font.bold, r.font.size.pt, str(r.font.color.rgb)) for r in notes] == [(False, 10, "C00000")] * 2   # 备注不加粗
    run = plain.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == "000000"                       # 非重点：只靠底色和边框淡化，文字仍是黑色
    assert "本次重点" in {sh.text_frame.text for sh in shapes if sh.has_text_frame}
    _, example = roadmap_template()
    example["roadmap"]["lanes"][0]["items"][1]["highlight"] = False
    example["roadmap"]["eras"][1]["highlight"] = False
    assert "本次重点" not in {sh.text_frame.text for sh in render(example) if sh.has_text_frame}


def test_eras_are_arrow_shapes_and_faded_ones_are_plain(make_repo):
    shapes = _renderer(make_repo)()
    old, new = _shape(shapes, "旧时代"), _shape(shapes, "新时代")
    assert old.auto_shape_type == new.auto_shape_type == MSO_SHAPE.PENTAGON
    assert str(old.fill.fore_color.rgb) == "F3F5F8" and str(old.line.color.rgb) == "D1D5DB"
    assert str(new.fill.fore_color.rgb) == "FFF1E6" and str(new.line.color.rgb) == "C00000"
    assert [p.runs[0].text for p in new.text_frame.paragraphs] == ["新时代", "多智能体"]
    assert [p.runs[0].font.bold for p in new.text_frame.paragraphs] == [True, False]   # 名称加粗，说明不加粗
    assert [p.runs[0].font.size.pt for p in new.text_frame.paragraphs] == [12, 10]
    assert str(old.text_frame.paragraphs[0].runs[0].font.color.rgb) == "000000"


def test_axis_and_lane_labels(make_repo):
    shapes = _renderer(make_repo)()
    chip, label = _shape(shapes, "2026"), _shape(shapes, "智算")
    assert str(chip.fill.fore_color.rgb) == "F3F5F8" and str(chip.text_frame.paragraphs[0].runs[0].font.color.rgb) == "283593"
    assert str(label.fill.fore_color.rgb) == "283593" and str(label.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert label.text_frame.paragraphs[0].runs[0].font.bold and label.text_frame.paragraphs[0].alignment == 2   # PP_ALIGN.CENTER


def test_without_eras_there_is_no_era_row(make_repo):
    _, example = roadmap_template()
    del example["roadmap"]["eras"]
    shapes = _renderer(make_repo)(example)
    assert not [sh for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.PENTAGON]


def test_render_example_reopens(make_repo, tmp_path):
    template, example = roadmap_template()
    path = tmp_path / "road.pptx"
    render_deck(deck_of(ROAD, example), make_repo({ROAD: (template, example)})).save(path)
    texts = {sh.text_frame.text.split("\n")[0] for sh in Presentation(path).slides[0].shapes if sh.has_text_frame}
    assert {"产品路线图", "2026", "2029", "智算", "存储", "芯片甲", "芯片乙", "控制器", "旧时代", "新时代"} <= texts


def _prepare_error(repo, example) -> str:
    with pytest.raises(PptError) as exc:
        prepare_pages(deck_of(ROAD, example), repo)
    return str(exc.value)


def _relaxed():
    """放宽槽位对 roadmap 的限制，让错误只来自组件的 schema 和语义检查。"""
    template, example = roadmap_template()
    template["slots"]["properties"]["roadmap"] = {"type": "object"}
    return template, example


@pytest.mark.parametrize("mutate, where, message", [
    (lambda r: r["lanes"][0]["items"][1].update(to="2030"),
     "lanes/0/items/1/to", "「2030」不在 periods 里，可选：2026、2027、2028、2029"),
    (lambda r: r["lanes"][1]["items"][0].update({"from": "去年"}),
     "lanes/1/items/0/from", "「去年」不在 periods 里"),
    (lambda r: r["eras"][0].update({"from": "2025"}), "eras/0/from", "「2025」不在 periods 里"),
    (lambda r: r["lanes"][0]["items"][1].update({"from": "2028", "to": "2027"}),
     "lanes/0/items/1/to", "to「2027」早于 from「2028」"),
    (lambda r: r["eras"][1].update({"from": "2029", "to": "2028"}), "eras/1/to", "to「2028」早于 from「2029」"),
    (lambda r: r["lanes"][0]["items"][0].update(to="2027"),
     "lanes/0/items/1", "与「芯片甲」占用了同一时间段，同一行里的条目不能重叠"),
    (lambda r: r["eras"][1].update({"from": "2027"}), "eras/1", "与「旧时代」占用了同一时间段"),
])
def test_period_semantics_are_checked_with_page_and_pointer(make_repo, mutate, where, message):
    template, example = _relaxed()
    mutate(example["roadmap"])
    repo = make_repo({ROAD: (template, example)})
    assert f"第 1 页 {ROAD}  /pages/0/data/roadmap/{where}  {message}" in _prepare_error(repo, example)
    assert any(f"example.yaml  /roadmap/{where}  " in e and message in e and e.endswith("（组件 lane-roadmap）")
               for e in validate_repo(repo).errors)


def test_items_in_different_lanes_may_share_periods(make_repo):
    template, example = _relaxed()
    assert prepare_pages(deck_of(ROAD, example), make_repo({ROAD: (template, example)}))   # 两条泳道都占 2026，不算重叠


def _rename_period(r, old, new):
    r["periods"][r["periods"].index(old)] = new
    for entry in [i for lane in r["lanes"] for i in lane["items"]] + r["eras"]:
        for key in ("from", "to"):
            if entry.get(key) == old:
                entry[key] = new


def _long_era_name(r):
    r["eras"][0].pop("to")
    r["eras"][0]["name"] = "Physical-AI-Embodied-Era-Long"


@pytest.mark.parametrize("mutate, where, message", [
    (lambda r: _rename_period(r, "2026", "二〇二六年上半年到二〇二七年下半年"), "periods/0",
     "「二〇二六年上半年到二〇二七年下半年」在 276×26 px 的框里估算要 2 行，只放得下 1 行"),
    (lambda r: r["lanes"][0]["items"][0].update(name="Prefill-Decode-Overlap-Pipeline-X"),
     "lanes/0/items/0/name", "英文单词「Prefill-Decode-Overlap-Pipeline-X」比 276 px 宽的框还宽，会被拦腰折断"),
    (lambda r: r["lanes"][0]["items"][0].update(notes=["Heterogeneous-Interconnect-Fabric-Plane"]),
     "lanes/0/items/0/notes/0", "英文单词「Heterogeneous-Interconnect-Fabric-Plane」比 276 px 宽的框还宽"),
    (_long_era_name, "eras/0/name", "英文单词「Physical-AI-Embodied-Era-Long」比 276 px 宽的框还宽"),
])
def test_text_that_cannot_fit_is_rejected_with_page_and_pointer(make_repo, mutate, where, message):
    template, example = _relaxed()
    mutate(example["roadmap"])
    repo = make_repo({ROAD: (template, example)})
    assert f"第 1 页 {ROAD}  /pages/0/data/roadmap/{where}  " in _prepare_error(repo, example)
    assert message in _prepare_error(repo, example)
    assert any(f"example.yaml  /roadmap/{where}  " in e and message in e for e in validate_repo(repo).errors)


def test_too_many_lines_for_the_body_height_are_rejected(make_repo):
    template, example = _relaxed()
    lane = {"name": "泳道", "items": [{"name": "条目", "from": "2026", "notes": ["备注甲", "备注乙", "备注丙"] * 2}]}
    example["roadmap"]["lanes"] = [dict(lane, name=f"泳道{i}") for i in range(4)]
    repo = make_repo({ROAD: (template, example)})
    message = _prepare_error(repo, example)
    assert "/pages/0/data/roadmap/lanes  4 条泳道和时代行估算要 " in message and "px 高，只放得下 " in message
    example["roadmap"]["lanes"] = example["roadmap"]["lanes"][:2]
    assert prepare_pages(deck_of(ROAD, example), repo)                                 # 2 条泳道放得下


@pytest.mark.parametrize("mutate, expected", [
    (lambda r: r.update(periods=["2026", "2026"]), "/periods  数组中有重复项"),
    (lambda r: r["lanes"][0].update(items=[]), "/lanes/0/items  共 0 项，低于下限 1"),
    (lambda r: r["lanes"][0]["items"][0].pop("from"), "/lanes/0/items/0  缺少必填字段 from"),
    (lambda r: r["eras"][0].pop("name"), "/eras/0  缺少必填字段 name"),
    (lambda r: r.pop("lanes"), "  缺少必填字段 lanes"),
])
def test_component_schema_errors_are_prefixed_with_the_page_pointer(make_repo, mutate, expected):
    template, example = _relaxed()
    mutate(example["roadmap"])
    repo = make_repo({ROAD: (template, example)})
    message = _prepare_error(repo, example)
    assert f"第 1 页 {ROAD}  /pages/0/data/roadmap{expected}" in message and "（组件 lane-roadmap）" in message


def test_render_rechecks_against_the_actual_body_area(make_repo):
    template, example = roadmap_template()
    repo = make_repo({ROAD: (template, example)})
    style_path = repo / "styles" / "default" / "style.yaml"
    style = yaml.safe_load(style_path.read_text(encoding="utf-8"))
    style["frame"]["body"]["box"] = [32, 128, 1216, 200]    # body 区比参考尺寸矮得多
    style_path.write_text(yaml.safe_dump(style, allow_unicode=True), encoding="utf-8")
    with pytest.raises(PptError, match="组件 lane-roadmap  /lanes  "):
        render_deck(deck_of(ROAD, example), repo)


def test_line_estimate():
    assert _lines("二〇二六年上半年", 100, 16) == 2 and _lines("2026", 100, 16) == 1
    assert _lines("备" * 12, 160, 40 / 3) == 1 and _lines("备" * 13, 160, 40 / 3) == 2
    assert _lines("Agentic AI时代", 146, 16) == 1
