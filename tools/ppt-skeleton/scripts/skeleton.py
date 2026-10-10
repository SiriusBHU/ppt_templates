#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ppt-skeleton：把 PPT 页脱敏成只保留版式的“骨架” pptx。

只依赖 Python 标准库 + python-pptx（python-pptx 自带 lxml）。

    extract IN.pptx --slides 1,3-5 -o WORK_DIR      提取文字，生成待填文件
    apply   IN.pptx WORK_DIR -o OUT.pptx            写入引导文字并剥离敏感内容
    check   ORIG.pptx OUT.pptx WORK_DIR [--words W] 审计 OUT 是否还有泄露
    render  OUT.pptx -o DIR                         用 soffice 转 PNG 供人工查看（可选）
"""
from __future__ import annotations

import argparse
import copy
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml import parse_xml
except ImportError as _exc:  # pragma: no cover
    sys.stderr.write("缺少依赖：请先运行 pip install python-pptx（%s）\n" % _exc)
    raise SystemExit(2)


class SkeletonError(Exception):
    """可预期的用户错误：打印中文提示后以退出码 1 结束。"""


# --------------------------------------------------------------------------- 常量

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A = "{%s}" % NS_A
P = "{%s}" % NS_P
R = "{%s}" % NS_R
MC = "{%s}" % NS_MC

SHAPE_TAGS = {P + t for t in ("sp", "grpSp", "graphicFrame", "pic", "cxnSp", "contentPart")}
NV_TAGS = {P + t for t in ("nvSpPr", "nvPicPr", "nvGraphicFramePr", "nvGrpSpPr", "nvCxnSpPr", "nvContentPartPr")}
FOOTER_PH = {"dt", "ftr", "sldNum"}
HLINK_TAGS = {A + "hlinkClick", A + "hlinkHover", A + "hlinkMouseOver"}
MEDIA_LOCALS = {"videoFile", "audioFile", "wavAudioFile", "quickTimeFile", "audioCd", "media"}

PLACEHOLDER_LABELS = {"图片", "图表", "图示", "对象"}  # 占位矩形上的标签，不算“与原文相同”
GREY_FILL = "E7E6E6"
GREY_LINE = "BFBFBF"
GREY_TEXT = "7F7F7F"
FIXED_DATE = datetime.datetime(2000, 1, 1)

RECT_TPL = (
    '<p:sp xmlns:a="%s" xmlns:p="%s" xmlns:r="%s">'
    '<p:nvSpPr><p:cNvPr id="{id}" name="Shape {id}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
    '<p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
    '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
    '<a:ln w="9525"><a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:ln></p:spPr>'
    '<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr">'
    "<a:noAutofit/></a:bodyPr><a:lstStyle/>{para}</p:txBody></p:sp>"
) % (NS_A, NS_P, NS_R, GREY_FILL, GREY_LINE)
LABEL_PARA = (
    '<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="zh-CN" altLang="en-US" sz="{sz}" dirty="0">'
    '<a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:rPr><a:t>{label}</a:t></a:r></a:p>' % GREY_TEXT
)
EMPTY_PARA = '<a:p><a:endParaRPr lang="zh-CN" altLang="en-US" dirty="0"/></a:p>'

APP_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
    "<Application>Microsoft Office PowerPoint</Application></Properties>"
).encode("utf-8")
PRES_PROPS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<p:presentationPr xmlns:a="%s" xmlns:r="%s" xmlns:p="%s"/>' % (NS_A, NS_R, NS_P)
).encode("utf-8")
VIEW_PROPS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<p:viewPr xmlns:a="%s" xmlns:r="%s" xmlns:p="%s"/>' % (NS_A, NS_R, NS_P)
).encode("utf-8")

PRESENTATION_DROP = (
    "notesMasterIdLst", "handoutMasterIdLst", "embeddedFontLst", "custShowLst",
    "custDataLst", "extLst", "photoAlbum", "modifyVerifier",
)

# check 用
LEAK_THRESHOLD = 40  # 汉字权重 8、其余字符权重 5：5 个汉字 = 8 个 ASCII = 40
TEXT_ATTRS = {"name", "descr", "title", "tooltip", "initials", "matchingName", "author", "styleName"}
APP_TEXT_TAGS = {"Company", "Manager", "HyperlinkBase", "Template", "lpstr", "lpwstr"}
CORE_IGNORE = {"created", "modified", "lastPrinted", "revision", "language"}
GENERIC_NAME = re.compile(r"^(Shape|Layout|Theme|Table Style)( \d+)?$")
ALLOWED_PARTS = re.compile(
    r"^(\[Content_Types\]\.xml|_rels/\.rels|docProps/(core|app)\.xml"
    r"|ppt/(presentation|presProps|viewProps|tableStyles)\.xml|ppt/_rels/presentation\.xml\.rels"
    r"|ppt/(slides|slideLayouts|slideMasters|theme)/(_rels/)?[^/]+\.(xml|rels))$"
)
PART_CATEGORIES = (
    ("ppt/media/", "图片/音视频"), ("ppt/embeddings/", "嵌入对象"), ("ppt/charts/", "图表"),
    ("ppt/diagrams/", "SmartArt"), ("ppt/notesSlides/", "演讲者备注"), ("ppt/notesMasters/", "备注母版"),
    ("ppt/handoutMasters/", "讲义母版"), ("ppt/comment", "批注"), ("ppt/authors", "批注作者"),
    ("ppt/people", "批注作者"), ("customXml/", "自定义 XML"), ("docProps/custom", "自定义属性"),
    ("docProps/thumbnail", "缩略图"), ("ppt/tags/", "标签"), ("ppt/fonts/", "内嵌字体"),
    ("ppt/ink/", "墨迹"), ("ppt/drawings/", "绘图"), ("ppt/printerSettings/", "打印设置"),
)


# --------------------------------------------------------------------------- 通用小工具

def local(tag):
    return etree.QName(tag).localname if isinstance(tag, str) else ""


def _int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def read_json(path):
    with open(path, encoding="utf-8-sig") as f:
        return json.load(f)


def write_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def parse_slides(spec, total):
    """“1,3-5” -> [1,3,4,5]（按给定顺序，去重）；all 表示全部。"""
    spec = spec.strip()
    if spec.lower() == "all":
        return list(range(1, total + 1))
    out = []
    for part in re.split(r"[,\s，、]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:[-~–](\d+))?", part)
        if not m:
            raise SkeletonError("无法解析页码“%s”，格式示例：1,3-5" % part)
        lo = int(m.group(1))
        hi = int(m.group(2) or lo)
        if lo < 1 or hi < lo or hi > total:
            raise SkeletonError("页码“%s”超出范围，这份 PPT 共 %d 页" % (part, total))
        for n in range(lo, hi + 1):
            if n not in out:
                out.append(n)
    if not out:
        raise SkeletonError("--slides 不能为空")
    return out


# --------------------------------------------------------------------------- 形状遍历

def unwrap_alt_content(tree):
    """把形状层级上的 mc:AlternateContent 展开成 Choice 里的内容（没有 Choice 就取 Fallback）。"""
    for ac in list(tree.iter(MC + "AlternateContent")):
        parent = ac.getparent()
        if parent is None or parent.tag not in (P + "spTree", P + "grpSp"):
            continue
        branch = ac.find(MC + "Choice")
        if branch is None:
            branch = ac.find(MC + "Fallback")
        children = [c for c in branch] if branch is not None else []
        idx = parent.index(ac)
        parent.remove(ac)
        for k, child in enumerate(children):
            parent.insert(idx + k, child)


def sp_tree(root):
    return root.find(P + "cSld/" + P + "spTree")


def xfrm_el(el):
    tag = el.tag
    if tag == P + "graphicFrame":
        return el.find(P + "xfrm")
    if tag == P + "grpSp":
        return el.find(P + "grpSpPr/" + A + "xfrm")
    if tag == P + "contentPart":
        return el.find(A + "xfrm")
    return el.find(P + "spPr/" + A + "xfrm")


def xfrm_values(x):
    if x is None:
        return None
    off, ext = x.find(A + "off"), x.find(A + "ext")
    if off is None or ext is None:
        return None
    vals = (_int(off.get("x")), _int(off.get("y")), _int(ext.get("cx")), _int(ext.get("cy")))
    return None if None in vals else vals


def _child_mapping(grp, mapping):
    x = grp.find(P + "grpSpPr/" + A + "xfrm")
    if x is None:
        return mapping
    off, ext = x.find(A + "off"), x.find(A + "ext")
    choff, chext = x.find(A + "chOff"), x.find(A + "chExt")
    vals = [_int(e.get(k)) if e is not None else None
            for e, k in ((off, "x"), (off, "y"), (ext, "cx"), (ext, "cy"),
                         (choff, "x"), (choff, "y"), (chext, "cx"), (chext, "cy"))]
    if None in vals:
        return mapping
    ox, oy, cx, cy, cox, coy, ccx, ccy = vals
    (ax, kx), (ay, ky) = mapping
    sx = kx * cx / ccx if ccx else kx
    sy = ky * cy / ccy if ccy else ky
    return ((ax + kx * ox - cox * sx, sx), (ay + ky * oy - coy * sy, sy))


def walk_shapes(tree, path=(), mapping=((0.0, 1.0), (0.0, 1.0))):
    """按文档顺序产出 (元素, 路径, 幻灯片坐标下的 box)，进入组合。"""
    idx = 0
    for el in tree:
        if el.tag not in SHAPE_TAGS:
            continue
        idx += 1
        p = path + (idx,)
        vals = xfrm_values(xfrm_el(el))
        box = None
        if vals is not None:
            (ax, kx), (ay, ky) = mapping
            box = (ax + kx * vals[0], ay + ky * vals[1], kx * vals[2], ky * vals[3])
        yield el, p, box
        if el.tag == P + "grpSp":
            for item in walk_shapes(el, p, _child_mapping(el, mapping)):
                yield item


def nv_el(el):
    for c in el:
        if c.tag in NV_TAGS:
            return c
    return None


def cnvpr(el):
    nv = nv_el(el)
    return nv.find(P + "cNvPr") if nv is not None else None


def ph_of(el):
    nv = nv_el(el)
    if nv is None:
        return None
    nvpr = nv.find(P + "nvPr")
    return nvpr.find(P + "ph") if nvpr is not None else None


# --------------------------------------------------------------------------- 文字与字号

def para_text(p):
    out = []
    for c in p:
        if c.tag in (A + "r", A + "fld"):
            t = c.find(A + "t")
            out.append((t.text or "") if t is not None else "")
        elif c.tag == A + "br":
            out.append("\n")
    return "".join(out)


RPR_NOISE = {"lang", "altLang", "dirty", "err", "smtClean", "noProof", "bmk"}
EDGE_SPACE = " \t\u3000\u00a0"


def _run_text(r):
    t = r.find(A + "t")
    return (t.text or "") if t is not None else ""


def fmt_key(run):
    """两个 run 的“肉眼格式”是否相同：忽略语言、拼写标记、超链接。"""
    rpr = run.find(A + "rPr")
    if rpr is None:
        return ""
    parts = [sorted((k, v) for k, v in rpr.attrib.items() if k not in RPR_NOISE)]
    for ch in rpr:
        if isinstance(ch.tag, str) and ch.tag not in HLINK_TAGS and ch.tag != A + "extLst":
            parts.append(etree.tostring(ch, method="c14n"))
    return repr(parts)


def para_groups(p):
    """把段落按格式切成 run 组：每组 = 连续且格式相同的 run（含其间的换行和空白 run）。

    组外（组与组之间、段首段尾）的换行原样保留。每组：items（要整体替换的元素）、text（去掉首尾空格）、
    lead / trail（首尾空格，替换时补回）、first（首个有字的 run，格式来源）。
    """
    kids = [c for c in p if c.tag in (A + "r", A + "fld", A + "br")]
    groups, pending = [], []
    for r in (c for c in kids if c.tag != A + "br"):
        if not _run_text(r).strip():
            (groups[-1]["runs"] if groups else pending).append(r)
            continue
        key = fmt_key(r)
        if groups and groups[-1]["key"] == key:
            groups[-1]["runs"].append(r)
        else:
            groups.append({"key": key, "runs": pending + [r], "first": r})
            pending = []
    for g in groups:
        i0, i1 = kids.index(g["runs"][0]), kids.index(g["runs"][-1])
        g["items"] = kids[i0:i1 + 1]
        raw = "".join("\n" if c.tag == A + "br" else _run_text(c) for c in g["items"])
        g["text"] = raw.strip(EDGE_SPACE)
        g["lead"] = raw[:len(raw) - len(raw.lstrip(EDGE_SPACE))]
        g["trail"] = raw[len(raw.rstrip(EDGE_SPACE)):]
    return groups


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _lvl_sz(style, lvl):
    if style is None:
        return None
    pp = style.find(A + "lvl%dpPr" % lvl)
    d = pp.find(A + "defRPr") if pp is not None else None
    sz = _num(d.get("sz")) if d is not None else None
    return sz / 100 if sz else None


def _find_ph_sp(root, ph, by_type_only=False):
    """在版式/母版里找与 ph 对应的占位符形状。"""
    idx, typ = ph.get("idx"), ph.get("type")
    cands = [e for e in root.iter(P + "ph")]
    pick = None
    if not by_type_only and idx is not None:
        pick = next((e for e in cands if e.get("idx") == idx), None)
    if pick is None:
        want = typ or "obj"
        if by_type_only:
            want = "title" if want in ("title", "ctrTitle") else "body"
        pick = next((e for e in cands if (e.get("type") or "obj") == want), None)
    if pick is None:
        return None
    return pick.getparent().getparent().getparent()


def resolve_font_pt(prs, slide, shape, para, run):
    """尽力解析一个 run 的字号（pt）：run → 形状 lstStyle → 版式/母版占位符 → 母版样式 → 默认。"""
    rpr = run.find(A + "rPr")
    sz = _num(rpr.get("sz")) if rpr is not None else None
    if sz:
        return sz / 100
    ppr = para.find(A + "pPr")
    lvl = (_int(ppr.get("lvl")) or 0) + 1 if ppr is not None else 1
    scale = 1.0
    chain = []
    tb = shape.find(P + "txBody")
    if tb is not None:
        chain.append(tb.find(A + "lstStyle"))
        fit = tb.find(A + "bodyPr/" + A + "normAutofit")
        fs = _num(fit.get("fontScale")) if fit is not None else None
        if fs:
            scale = fs / 100000
    ph = ph_of(shape)
    if ph is not None:
        try:
            layout = slide.slide_layout
            lsp = _find_ph_sp(layout._element, ph)
            if lsp is not None:
                chain.append(lsp.find(P + "txBody/" + A + "lstStyle"))
            master = layout.slide_master._element
            msp = _find_ph_sp(master, ph, by_type_only=True)
            if msp is not None:
                chain.append(msp.find(P + "txBody/" + A + "lstStyle"))
            typ = ph.get("type") or "obj"
            name = "titleStyle" if typ in ("title", "ctrTitle") else (
                "bodyStyle" if typ in ("body", "subTitle", "obj") else "otherStyle")
            chain.append(master.find(P + "txStyles/" + P + name))
        except Exception:  # 版式缺失之类的怪文件，退回默认字号
            pass
    chain.append(prs.part._element.find(P + "defaultTextStyle"))
    for style in chain:
        sz = _lvl_sz(style, lvl)
        if sz:
            return round(sz * scale, 1)
    return round(18.0 * scale, 1)


def inherited_box(slide, ph):
    """没有自带位置的占位符：从版式、母版继承位置。"""
    layout = slide.slide_layout
    for root, by_type in ((layout._element, False), (layout.slide_master._element, True)):
        sp = _find_ph_sp(root, ph, by_type_only=by_type)
        vals = xfrm_values(xfrm_el(sp)) if sp is not None else None
        if vals:
            return vals
    return None


def char_count(text):
    return len(re.sub(r"\s+", "", text))


def char_budget(n):
    if n <= 4:
        return [max(1, n - 1), n + 1]
    lo = max(1, int(n * 0.8 + 0.5))
    hi = max(lo, int(n * 1.2 + 0.5))
    return [lo, hi]


def _fmt_pt(v):
    return int(v) if float(v) == int(v) else v


# --------------------------------------------------------------------------- 提取

def prepare_slide(slide):
    tree = sp_tree(slide._element)
    if tree is not None:
        unwrap_alt_content(tree)


def collect_entries(prs, slide, slide_no):
    """一页里每个非空文字段落一条记录；ids 只依赖文档顺序，同一份文件重复提取结果一致。"""
    tree = sp_tree(slide._element)
    entries = []
    if tree is None:
        return entries
    width, height = int(prs.slide_width), int(prs.slide_height)
    seq = 0

    def add(shape, para, box, kind, location):
        nonlocal seq
        if box is None and ph_of(shape) is not None:
            try:
                box = inherited_box(slide, ph_of(shape))
            except Exception:
                box = None
        groups = para_groups(para)
        for gi, g in enumerate(groups, 1):
            seq += 1
            n = char_count(g["text"])
            entries.append({
                "id": "s%02d-%03d" % (slide_no, seq),
                "slide": slide_no,
                "location": location + (" · 片段%d/%d" % (gi, len(groups)) if len(groups) > 1 else ""),
                "kind": kind,
                "text": g["text"],
                "font_pt": _fmt_pt(resolve_font_pt(prs, slide, shape, para, g["first"])),
                "chars": n,
                "budget": char_budget(n),
                "box_pct": [round(box[0] / width * 100, 1), round(box[1] / height * 100, 1),
                            round(box[2] / width * 100, 1), round(box[3] / height * 100, 1)] if box else None,
                "_para": para,
                "_group": g,
            })

    for el, path, box in list(walk_shapes(tree)):
        ph = ph_of(el)
        if ph is not None and ph.get("type") in FOOTER_PH:
            continue
        spath = "/".join(str(i) for i in path)
        if el.tag == P + "sp":
            tb = el.find(P + "txBody")
            if tb is None:
                continue
            paras = [p for p in tb.findall(A + "p") if para_text(p).strip()]
            if ph is not None:
                kind = "占位符:%s" % (ph.get("type") or "obj")
            elif (el.find(P + "nvSpPr/" + P + "cNvSpPr") is not None
                  and el.find(P + "nvSpPr/" + P + "cNvSpPr").get("txBox") == "1"):
                kind = "文本框"
            else:
                kind = "形状"
            for k, p in enumerate(paras, 1):
                loc = "形状 %s" % spath + (" · 段%d" % k if len(paras) > 1 else "")
                add(el, p, box, kind, loc)
        elif el.tag == P + "graphicFrame":
            tbl = el.find(A + "graphic/" + A + "graphicData/" + A + "tbl")
            if tbl is None:
                continue
            cols = [_int(g.get("w")) or 0 for g in tbl.findall(A + "tblGrid/" + A + "gridCol")]
            rows = tbl.findall(A + "tr")
            heights = [_int(tr.get("h")) or 0 for tr in rows]
            for ri, tr in enumerate(rows, 1):
                for ci, tc in enumerate(tr.findall(A + "tc"), 1):
                    tb = tc.find(A + "txBody")
                    if tb is None:
                        continue
                    cell_box = box
                    if box and ci <= len(cols):
                        cell_box = (box[0] + sum(cols[:ci - 1]), box[1] + sum(heights[:ri - 1]),
                                    cols[ci - 1], heights[ri - 1])
                    paras = [p for p in tb.findall(A + "p") if para_text(p).strip()]
                    for k, p in enumerate(paras, 1):
                        loc = "表格 %s · 行%d列%d" % (spath, ri, ci) + (" · 段%d" % k if len(paras) > 1 else "")
                        add(tc, p, cell_box, "表格单元格", loc)
    return entries


MANIFEST_HEAD = (
    "# 骨架稿说明，可随骨架 pptx 一起带出。\n"
    "# purpose：用一句通用的话描述该页的用途，不得包含项目、客户、产品信息。\n"
    "# 值里有半角冒号时请加双引号。\n"
    "slides:\n"
)


def cmd_extract(args):
    work = args.output
    prs = Presentation(args.input)
    total = len(prs.slides)
    selected = parse_slides(args.slides, total)
    rep_path = os.path.join(work, "replacements.json")
    man_path = os.path.join(work, "manifest.yaml")
    if not args.force and (os.path.exists(rep_path) or os.path.exists(man_path)):
        raise SkeletonError("%s 里已有 replacements.json / manifest.yaml，为避免覆盖已填内容已停止；"
                            "确需重来请加 --force 或换一个目录" % work)
    os.makedirs(work, exist_ok=True)
    items = []
    for n in selected:
        slide = prs.slides[n - 1]
        prepare_slide(slide)
        items.extend(collect_entries(prs, slide, n))
    for e in items:
        e.pop("_para")
        e.pop("_group")
    with open(os.path.join(work, "texts.json"), "w", encoding="utf-8") as f:
        head = {"_notice": "本文件含原文，仅限本机使用，严禁带出公司。", "slides": selected,
                "slide_size_emu": [int(prs.slide_width), int(prs.slide_height)]}
        f.write(json.dumps(head, ensure_ascii=False)[:-1] + ',\n "items": [\n')
        f.write(",\n".join("  " + json.dumps(e, ensure_ascii=False) for e in items))
        f.write("\n]}\n")
    write_json(rep_path, {e["id"]: "" for e in items})
    with open(man_path, "w", encoding="utf-8") as f:
        f.write(MANIFEST_HEAD)
        for i in range(1, len(selected) + 1):
            f.write('  - index: %d\n    purpose: ""\n' % i)
    print("已提取 %d 页、%d 段文字 -> %s" % (len(selected), len(items), work))
    print("  texts.json        原文（只读，严禁带出）")
    print("  replacements.json 按 id 填写引导文字")
    print("  manifest.yaml     填写每页的通用用途 purpose（第 i 项对应骨架第 i 页：%s）"
          % ", ".join(str(n) for n in selected))


# --------------------------------------------------------------------------- 应用：替换文字

def rewrite_group(p, group, text):
    """用 text 替换一组 run：沿用组内首个有字的 run 的格式；\\n 变成换行；首尾空格按原样补回。"""
    first = group["first"]
    rpr = None
    if first.find(A + "rPr") is not None:
        rpr = copy.deepcopy(first.find(A + "rPr"))
    elif p.find(A + "endParaRPr") is not None:
        rpr = copy.deepcopy(p.find(A + "endParaRPr"))
        rpr.tag = A + "rPr"
    at = list(p).index(group["items"][0])
    for c in group["items"]:
        p.remove(c)
    new = []
    for i, seg in enumerate((group["lead"] + text + group["trail"]).split("\n")):
        if i:
            br = etree.Element(A + "br")
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            new.append(br)
        if seg:
            r = etree.Element(A + "r")
            if rpr is not None:
                r.append(copy.deepcopy(rpr))
            etree.SubElement(r, A + "t").text = seg
            new.append(r)
    for k, e in enumerate(new):
        p.insert(at + k, e)


def clear_paragraphs(txbody):
    """清空文字，只留第一段的 pPr/endParaRPr（用于母版、版式里的文字）。"""
    paras = txbody.findall(A + "p")
    for p in paras[1:]:
        txbody.remove(p)
    if paras:
        for c in list(paras[0]):
            if c.tag not in (A + "pPr", A + "endParaRPr"):
                paras[0].remove(c)
    else:
        etree.SubElement(txbody, A + "p")


def tidy_paragraphs(root):
    """段落里混着数学公式之类的怪东西时，只留 pPr / r / br / fld / endParaRPr。"""
    for p in root.iter(A + "p"):
        for c in list(p):
            if c.tag not in (A + "pPr", A + "r", A + "br", A + "fld", A + "endParaRPr"):
                p.remove(c)


def length_warnings(items, reps):
    warns = []
    for e in items:
        orig, new = e["text"], reps[e["id"]]
        if not any(ch.isalpha() for ch in orig):
            continue
        n, m = e["chars"], char_count(new)
        if (n < 6 and m > n + 3) or (n >= 6 and (m / n > 1.3 or m / n < 0.6)):
            warns.append("%s 原文 %d 字，引导文字 %d 字（建议 %d-%d 字）" % (e["id"], n, m, *e["budget"]))
    return warns


# --------------------------------------------------------------------------- 应用：剥离

def build_rect(src, label=None, label_ok=True, area_ratio=0.0):
    """生成与 src 同位置同大小的浅灰矩形；做不到（没有位置信息）返回 None。"""
    c = cnvpr(src)
    sid = c.get("id") if c is not None else "1"
    xf = xfrm_el(src)
    new_x = None
    if xf is not None:
        if src.tag == P + "graphicFrame":
            new_x = etree.Element(A + "xfrm")
            for ch in xf:
                new_x.append(copy.deepcopy(ch))
        else:
            new_x = copy.deepcopy(xf)
    ph = ph_of(src)
    if new_x is None and ph is None:
        return None
    vals = xfrm_values(xf)
    para = EMPTY_PARA
    if label and label_ok and area_ratio < 0.25:
        small = min(vals[2], vals[3]) if vals else 10 ** 9
        sz = 1200 if small >= 500000 else (800 if small >= 250000 else 0)
        if sz:
            para = LABEL_PARA.format(sz=sz, label=label)
    sp = etree.fromstring(RECT_TPL.format(id=sid, para=para))
    sppr = sp.find(P + "spPr")
    if new_x is not None:
        sppr.insert(0, new_x)
    if ph is not None and new_x is None:
        sp.find(P + "nvSpPr/" + P + "nvPr").append(copy.deepcopy(ph))
    if src.tag == P + "pic":
        old = src.find(P + "spPr")
        if old is not None:
            geom = old.find(A + "prstGeom")
            if geom is not None:
                sppr.replace(sppr.find(A + "prstGeom"), copy.deepcopy(geom))
            ln = old.find(A + "ln")
            if ln is not None:
                sppr.replace(sppr.find(A + "ln"), copy.deepcopy(ln))
            fx = old.find(A + "effectLst")
            if fx is not None:
                sppr.append(copy.deepcopy(fx))
    return sp


def graphic_label(gf):
    gd = gf.find(A + "graphic/" + A + "graphicData")
    uri = (gd.get("uri") if gd is not None else "") or ""
    if uri.endswith("/drawingml/2006/table"):
        return None
    low = uri.lower()
    if "chart" in low:
        return "图表"
    if "diagram" in low:
        return "图示"
    return "对象"


def pic_label(pic):
    for e in pic.iter():
        if isinstance(e.tag, str) and local(e.tag) in MEDIA_LOCALS:
            return "对象"
    return "图片"


def replace_visuals(tree, kind, slide_area):
    """图片/图表/SmartArt/OLE/媒体 -> 灰色矩形；母版、版式里的所有图形框都换掉。"""
    for el, _path, box in list(walk_shapes(tree)):
        ph = ph_of(el)
        if ph is not None and ph.get("type") in FOOTER_PH:
            el.getparent().remove(el)
            continue
        label = None
        if el.tag == P + "pic":
            label = pic_label(el)
        elif el.tag == P + "contentPart":
            label = "对象"
        elif el.tag == P + "graphicFrame":
            label = graphic_label(el)
            if label is None and kind != "slide":
                label = "对象"
        else:
            if kind != "slide" and el.tag == P + "sp":
                tb = el.find(P + "txBody")
                if tb is not None:
                    clear_paragraphs(tb)
            continue
        if label is None:
            continue
        ratio = (box[2] * box[3] / slide_area) if (kind != "slide" and box and slide_area) else 0.0
        rect = build_rect(el, label, True, ratio)
        parent = el.getparent()
        if rect is None:
            parent.remove(el)
        else:
            parent.replace(el, rect)


def common_scrub(root):
    """超链接、标签、图片填充、形状名与替代文字。"""
    for e in list(root.iter()):
        if not isinstance(e.tag, str):
            continue
        t = e.tag
        if t in HLINK_TAGS or t in (P + "custDataLst", A + "buBlip"):
            if e.getparent() is not None:
                e.getparent().remove(e)
        elif t == A + "blipFill":
            for ch in list(e):
                e.remove(ch)
            e.attrib.clear()
            e.tag = A + "solidFill"
            etree.SubElement(e, A + "srgbClr").set("val", GREY_FILL)
        elif t == P + "cNvPr":
            e.set("name", "Shape %s" % e.get("id", "0"))
            for k in ("descr", "title"):
                if k in e.attrib:
                    del e.attrib[k]


def strip_unreferenced(root, dropped):
    """删除仍引用已丢弃关系的元素，避免留下悬空 rId。"""
    if not dropped:
        return
    for e in list(root.iter()):
        if not isinstance(e.tag, str) or e.getparent() is None:
            continue
        if any(k.startswith(R) and v in dropped for k, v in e.attrib.items()):
            e.getparent().remove(e)


def prune_rels(part, keep_types):
    """只保留 keep_types 里的关系；其余（含外部链接）全部丢弃，孤儿部件因此不会被保存。"""
    rels = part.rels
    dropped = {rid for rid, rel in list(rels.items()) if rel.is_external or rel.reltype not in keep_types}
    root = getattr(part, "_element", None)
    if root is not None:
        strip_unreferenced(root, dropped)
    for rid in dropped:
        rels.pop(rid)


def set_part_xml(part, xml_bytes):
    if hasattr(part, "_element"):
        part._element = parse_xml(xml_bytes)
    else:
        part._blob = xml_bytes


def rel_targets(part, reltype):
    return [rel.target_part for rel in part.rels.values() if not rel.is_external and rel.reltype == reltype]


def scrub_theme(part):
    if hasattr(part, "_element"):
        root = part._element
    else:
        root = etree.fromstring(part.blob)
    root.attrib.pop("name", None)
    for e in list(root.iter()):
        if not isinstance(e.tag, str):
            continue
        ln = local(e.tag)
        if ln in ("clrScheme", "fontScheme", "fmtScheme") or e.tag == A + "custClr":
            e.attrib.pop("name", None)
        elif ln == "blipFill":
            for ch in list(e):
                e.remove(ch)
            e.attrib.clear()
            e.tag = A + "solidFill"
            etree.SubElement(e, A + "schemeClr").set("val", "phClr")
    for e in list(root):
        if e.tag == A + "extLst":
            root.remove(e)
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    prune_rels(part, set())


def scrub_table_styles(part):
    root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    for i, e in enumerate(root.iter(A + "tblStyle"), 1):
        e.set("styleName", "Table Style %d" % i)
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    prune_rels(part, set())


def scrub_slide_part(slide, slide_area):
    root = slide._element
    tree = sp_tree(root)
    if tree is not None:
        unwrap_alt_content(tree)
        replace_visuals(tree, "slide", slide_area)
    tidy_paragraphs(root)
    common_scrub(root)
    root.attrib.pop("show", None)
    cs = root.find(P + "cSld")
    if cs is not None:
        cs.attrib.pop("name", None)
    for e in list(root):
        if e.tag in (P + "timing", P + "transition") or (
                e.tag == MC + "AlternateContent" and any(local(x.tag) == "transition" for x in e.iter())):
            root.remove(e)
    prune_rels(slide.part, {RT.SLIDE_LAYOUT})


def scrub_layout_or_master(part, kind, number, slide_area):
    root = part._element
    tree = sp_tree(root)
    if tree is not None:
        unwrap_alt_content(tree)
        replace_visuals(tree, kind, slide_area)
    common_scrub(root)
    cs = root.find(P + "cSld")
    if cs is not None:
        if kind == "layout":
            cs.set("name", "Layout %d" % number)
        else:
            cs.attrib.pop("name", None)
    root.attrib.pop("matchingName", None)
    hf = root.find(P + "hf")
    if hf is not None:
        root.remove(hf)
    prune_rels(part, {RT.SLIDE_MASTER} if kind == "layout" else {RT.SLIDE_LAYOUT, RT.THEME})


def strip_deck(prs, selected, slide_area):
    """保留 selected（按顺序），剥离其余一切；返回最终保留的幻灯片对象列表。"""
    pres_part = prs.part
    id_lst = prs.slides._sldIdLst
    all_ids = list(id_lst)
    keep_ids = [all_ids[n - 1] for n in selected]
    keep_slides = [prs.slides[n - 1] for n in selected]
    for sid in all_ids:
        id_lst.remove(sid)
        if sid not in keep_ids:
            pres_part.rels.pop(sid.get(R + "id"))
    for sid in keep_ids:
        id_lst.append(sid)
    pres_part.rename_slide_parts([sid.get(R + "id") for sid in keep_ids])  # slide1.xml, slide2.xml ...

    for slide in keep_slides:
        scrub_slide_part(slide, slide_area)

    # 去掉没用到的版式和母版
    used_layouts = {rel_targets(s.part, RT.SLIDE_LAYOUT)[0] for s in keep_slides}
    masters = rel_targets(pres_part, RT.SLIDE_MASTER)
    used_masters = {rel_targets(lp, RT.SLIDE_MASTER)[0] for lp in used_layouts}
    master_lst = pres_part._element.find(P + "sldMasterIdLst")
    for mid in list(master_lst) if master_lst is not None else []:
        rid = mid.get(R + "id")
        if pres_part.rels[rid].target_part not in used_masters:
            master_lst.remove(mid)
            pres_part.rels.pop(rid)
    layout_no = 0
    themes = set()
    for m in [m for m in masters if m in used_masters]:
        lst = m._element.find(P + "sldLayoutIdLst")
        for lid in list(lst) if lst is not None else []:
            rid = lid.get(R + "id")
            if m.rels[rid].target_part not in used_layouts:
                lst.remove(lid)
                m.rels.pop(rid)
        for lp in rel_targets(m, RT.SLIDE_LAYOUT):
            layout_no += 1
            scrub_layout_or_master(lp, "layout", layout_no, slide_area)
        themes.update(rel_targets(m, RT.THEME))
        scrub_layout_or_master(m, "master", 0, slide_area)

    # 演示文稿级别
    pres = pres_part._element
    for tag in PRESENTATION_DROP:
        for e in pres.findall(P + tag):
            pres.remove(e)
    themes.update(rel_targets(pres_part, RT.THEME))
    for th in themes:
        scrub_theme(th)
    for pp in rel_targets(pres_part, RT.PRES_PROPS):
        set_part_xml(pp, PRES_PROPS_XML)
        prune_rels(pp, set())
    for vp in rel_targets(pres_part, RT.VIEW_PROPS):
        set_part_xml(vp, VIEW_PROPS_XML)
        prune_rels(vp, set())
    for ts in rel_targets(pres_part, RT.TABLE_STYLES):
        scrub_table_styles(ts)
    prune_rels(pres_part, {RT.SLIDE, RT.SLIDE_MASTER, RT.THEME, RT.PRES_PROPS, RT.VIEW_PROPS, RT.TABLE_STYLES})

    # 文档属性
    pkg = pres_part.package
    pkg_rels = getattr(pkg, "_rels", None) or pkg.rels  # 不同版本 python-pptx 的属性名
    cp = prs.core_properties
    el = cp._element
    for ch in list(el):
        el.remove(ch)
    cp.created = FIXED_DATE
    cp.modified = FIXED_DATE
    cp.revision = 1
    for rel in list(pkg_rels.values()):
        if rel.reltype == RT.EXTENDED_PROPERTIES and not rel.is_external:
            rel.target_part._blob = APP_XML
    keep = {RT.OFFICE_DOCUMENT, RT.CORE_PROPERTIES, RT.EXTENDED_PROPERTIES}
    for rid in [rid for rid, rel in list(pkg_rels.items()) if rel.is_external or rel.reltype not in keep]:
        pkg_rels.pop(rid)
    return keep_slides


def cmd_apply(args):
    work = args.work_dir
    out_abs = os.path.abspath(args.output)
    work_abs = os.path.abspath(work)
    try:
        inside = os.path.commonpath([out_abs, work_abs]) == work_abs
    except ValueError:  # Windows 不同盘符
        inside = False
    if inside:
        raise SkeletonError("输出文件不能放在 WORK_DIR 里（里面有含原文的 texts.json），请换一个目录")
    try:
        meta = read_json(os.path.join(work, "texts.json"))
        reps = read_json(os.path.join(work, "replacements.json"))
    except (OSError, ValueError) as exc:
        raise SkeletonError("读取 WORK_DIR 失败：%s（请先运行 extract）" % exc)
    selected, items = meta["slides"], meta["items"]
    prs = Presentation(args.input)
    if max(selected) > len(prs.slides):
        raise SkeletonError("IN.pptx 只有 %d 页，与 extract 时的文件不一致" % len(prs.slides))

    entries = []
    for n in selected:
        slide = prs.slides[n - 1]
        prepare_slide(slide)
        entries.extend(collect_entries(prs, slide, n))
    if [(e["id"], e["text"]) for e in entries] != [(e["id"], e["text"]) for e in items]:
        raise SkeletonError("IN.pptx 与 texts.json 对不上（文件被改过，或不是 extract 时用的那份），请重新 extract")

    unknown = [k for k in reps if k not in {e["id"] for e in items}]
    if unknown:
        raise SkeletonError("replacements.json 里有不存在的 id：%s" % "、".join(unknown[:10]))
    missing = [e["id"] for e in items if not isinstance(reps.get(e["id"]), str) or not reps[e["id"]].strip()]
    if missing:
        raise SkeletonError("还有 %d 条没有填写引导文字，不能生成骨架：%s%s"
                            % (len(missing), "、".join(missing[:15]), " ……" if len(missing) > 15 else ""))

    for e in entries:
        rewrite_group(e["_para"], e["_group"], reps[e["id"]].strip())
    area = int(prs.slide_width) * int(prs.slide_height)
    strip_deck(prs, selected, area)
    out_dir = os.path.dirname(out_abs)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    prs.save(args.output)

    print("已生成骨架：%s（%d 页，替换 %d 段文字）" % (args.output, len(selected), len(items)))
    for w in length_warnings(items, reps):
        print("提示（长度偏差较大，可能溢出或过空）：" + w)
    print("下一步：运行 check 审计；通过前不要把任何文件带出公司。")


# --------------------------------------------------------------------------- 审计

SENT = "\x00"   # 单元分隔
MASK = "\x01"   # 被 --allow 屏蔽的词


def norm(s):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s).casefold())


def weight(ch):
    o = ord(ch)
    cjk = 0x4E00 <= o <= 0x9FFF or 0x3400 <= o <= 0x4DBF or 0xF900 <= o <= 0xFAFF or 0x20000 <= o <= 0x2A6DF
    return 8 if cjk else 5


def mask_words(s, allow):
    for w in allow:
        s = s.replace(w, MASK)
    return s


def common_leaks(needle, hay):
    """needle 里（被屏蔽词切开的各段）与 hay 的最长公共片段，权重 >= 阈值才算。"""
    found = []
    for seg in needle.split(MASK):
        n, i = len(seg), 0
        while i < n:
            w, j = 0, i
            while j < n and w < LEAK_THRESHOLD:
                w += weight(seg[j])
                j += 1
            if w < LEAK_THRESHOLD:
                break
            if seg[i:j] in hay:
                k = j
                while k < n and seg[i:k + 1] in hay:
                    k += 1
                found.append(seg[i:k])
                i = k
            else:
                i += 1
    return found


NUMBER_RE = re.compile(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?")


def numbers_in(text):
    out = set()
    for tok in NUMBER_RE.findall(unicodedata.normalize("NFKC", text)):
        tok = tok.replace(",", "")
        if sum(c.isdigit() for c in tok) >= 3:
            out.add(tok)
    return out


def mask_raw(text, allow_raw):
    t = unicodedata.normalize("NFKC", text)
    for w in allow_raw:
        t = re.sub(re.escape(unicodedata.normalize("NFKC", w)), MASK, t, flags=re.I)
    return t


def is_trivial(n):
    s = n.replace(MASK, "")
    if len(s) <= 1:
        return True
    return not any(c.isalpha() for c in s) and sum(c.isdigit() for c in s) < 3


def _xml_parser():
    return etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)


def collect_zip(path):
    """读 pptx 的每个部件；返回 content（正文段落等）、attrs（名称/替代文字）、strings（一切字符串）。"""
    content, attrs, strings, names = [], [], [], []
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.endswith("/"):
                continue
            names.append(name)
            strings.append((name, name))
            low = name.lower()
            if not (low.endswith(".xml") or low.endswith(".rels")):
                continue
            data = zf.read(name)
            try:
                root = etree.fromstring(data, _xml_parser())
            except etree.XMLSyntaxError:
                strings.append((name, data.decode("utf-8", "ignore")))
                continue
            for e in root.iter():
                if not isinstance(e.tag, str):
                    continue
                if e.text and e.text.strip():
                    strings.append((name, e.text))
                for k, v in e.attrib.items():
                    if v.strip():
                        strings.append((name, v))
                        if local(k) in TEXT_ATTRS and not GENERIC_NAME.match(v):
                            attrs.append((name, v))
            if name == "docProps/core.xml":
                content.extend((name, e.text) for e in root
                               if isinstance(e.tag, str) and local(e.tag) not in CORE_IGNORE and (e.text or "").strip())
            elif name == "docProps/app.xml":
                content.extend((name, e.text) for e in root.iter()
                               if isinstance(e.tag, str) and local(e.tag) in APP_TEXT_TAGS and (e.text or "").strip())
            elif name == "docProps/custom.xml":
                content.extend((name, e.text) for e in root.iter() if isinstance(e.tag, str) and (e.text or "").strip())
            elif name.startswith("ppt/") and low.endswith(".xml"):
                for p in root.iter(A + "p"):
                    txt = "".join((x.text or "") if local(x.tag) == "t" else ("\n" if x.tag == A + "br" else "")
                                  for x in p.iter() if isinstance(x.tag, str))
                    if txt.strip():
                        content.append((name, txt))
                for e in root.iter():
                    ln = local(e.tag)
                    if not (e.text or "").strip():
                        continue
                    if ln in ("v", "text"):
                        content.append((name, e.text))
                    elif ln == "t" and not any(a.tag == A + "p" for a in e.iterancestors()):
                        content.append((name, e.text))
    return {"content": content, "attrs": attrs, "strings": strings, "names": names}


def manifest_values(path):
    """manifest.yaml 里所有值（不含注释和键名）。"""
    vals, purposes = [], []
    if not os.path.exists(path):
        return vals, purposes
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            s = line.rstrip("\n")
            if not s.strip() or s.lstrip().startswith("#"):
                continue
            m = re.match(r"^\s*(?:-\s*)?([A-Za-z_][\w-]*)\s*:\s*(.*?)\s*$", s)
            key, val = (m.group(1), m.group(2)) if m else (None, s.strip())
            if len(val) >= 2 and val[0] == val[-1] and val[0] in "\"'":
                val = val[1:-1]
            if key == "purpose":
                purposes.append(val)
            if val and key not in ("index", "slides"):
                vals.append(val)
    return vals, purposes


def load_words(path):
    if not path:
        return []
    with open(path, encoding="utf-8-sig") as f:
        return [ln.strip() for ln in f if ln.strip() and not ln.lstrip().startswith("#")]


def part_label(name):
    m = re.match(r"ppt/slides/slide(\d+)\.xml$", name)
    return "OUT 第 %s 页" % m.group(1) if m else "OUT:" + name


LOCAL_LABELS = ("manifest.yaml", "replacements.json", "文件名")


def short(s, n=36):
    return s if len(s) <= n else s[:n] + "…"


def cmd_check(args):
    allow_raw = [w for w in (args.allow or []) if w.strip()]
    allow = sorted({norm(w) for w in allow_raw if norm(w)}, key=len, reverse=True)
    words = load_words(args.words)
    work = args.work_dir
    rep_path = os.path.join(work, "replacements.json")
    if not os.path.exists(rep_path):
        raise SkeletonError("找不到 %s" % rep_path)
    reps = [v for v in read_json(rep_path).values() if isinstance(v, str) and v.strip()]
    man_vals, purposes = manifest_values(os.path.join(work, "manifest.yaml"))

    orig = collect_zip(args.orig)
    out = collect_zip(args.out)
    findings = []  # (类别, 说明)

    # ---- (a) 与原文重合的片段
    hay_units = ([(part_label(p), t) for p, t in out["content"] + out["attrs"]]
                 + [("manifest.yaml", t) for t in man_vals] + [("replacements.json", t) for t in reps])
    hay_norm = [(loc, mask_words(norm(t), allow)) for loc, t in hay_units]
    hay = SENT.join(t for _, t in hay_norm)
    needles = {mask_words(norm(t), allow) for _, t in orig["content"] + orig["attrs"]}
    seen = set()
    for nd in sorted(needles):
        for frag in common_leaks(nd, hay):
            if frag in seen:
                continue
            seen.add(frag)
            locs = sorted({loc for loc, t in hay_norm if frag in t})
            findings.append(("a", "与原文重合：「%s」出现在 %s" % (short(frag), "、".join(locs))))

    # ---- (b) 原文中的数字（>=3 位）
    orig_nums = set()
    for _, t in orig["content"]:
        orig_nums |= numbers_in(mask_raw(t, allow_raw))
    num_units = ([(part_label(p), t) for p, t in out["content"]]
                 + [("manifest.yaml", t) for t in man_vals] + [("replacements.json", t) for t in reps])
    num_locs = {}
    for loc, t in num_units:
        for n in numbers_in(mask_raw(t, allow_raw)) & orig_nums:
            num_locs.setdefault(n, set()).add(loc)
    for n in sorted(num_locs):
        findings.append(("b", "沿用了原文数字「%s」：%s（请改成 N / XX）" % (n, "、".join(sorted(num_locs[n])))))

    # ---- (c) 敏感词
    all_units = (list(out["strings"]) + list(out["content"]) + [("manifest.yaml", t) for t in man_vals]
                 + [("replacements.json", t) for t in reps] + [("文件名", os.path.basename(args.out))])
    all_norm = [(loc, norm(t)) for loc, t in all_units]
    for w in words:
        nw = norm(w)
        if not nw:
            continue
        locs = sorted({loc if loc in LOCAL_LABELS else part_label(loc) for loc, t in all_norm if nw in t})
        if locs:
            findings.append(("c", "出现敏感词「%s」：%s" % (w, "、".join(locs[:6]) + (" 等" if len(locs) > 6 else ""))))

    # ---- (d) 不该留下的部件
    for name in sorted(out["names"]):
        if ALLOWED_PARTS.match(name):
            continue
        cat = next((c for pre, c in PART_CATEGORIES if name.startswith(pre)), "非预期部件")
        findings.append(("d", "残留部件 %s（%s）" % (name, cat)))

    # ---- (e) 文档属性
    with zipfile.ZipFile(args.out) as zf:
        for name in ("docProps/core.xml", "docProps/app.xml"):
            if name not in zf.namelist():
                continue
            root = etree.fromstring(zf.read(name), _xml_parser())
            for e in root.iter():
                ln = local(e.tag)
                if not ln or not (e.text or "").strip() or e is root:
                    continue
                if name.endswith("core.xml") and ln not in CORE_IGNORE:
                    findings.append(("e", "docProps/core.xml 的 %s 非空：「%s」" % (ln, short(e.text.strip()))))
                if name.endswith("app.xml") and ln in APP_TEXT_TAGS:
                    findings.append(("e", "docProps/app.xml 的 %s 非空：「%s」" % (ln, short(e.text.strip()))))

        # ---- (f) 超链接
        for name in zf.namelist():
            low = name.lower()
            if low.endswith(".rels"):
                root = etree.fromstring(zf.read(name), _xml_parser())
                for rel in root:
                    typ, mode = rel.get("Type", ""), rel.get("TargetMode", "")
                    if typ.endswith("/hyperlink") or mode == "External":
                        findings.append(("f", "%s 中有外部/超链接关系：%s" % (name, short(rel.get("Target", "")))))
            elif low.endswith(".xml") and name.startswith("ppt/"):
                root = etree.fromstring(zf.read(name), _xml_parser())
                if any(e.tag in HLINK_TAGS for e in root.iter()):
                    findings.append(("f", "%s 中仍有超链接/点击动作" % name))

    # ---- (g) 与原文逐字相同的文字
    orig_set = {mask_words(norm(t), allow) for p, t in orig["content"] if p.startswith("ppt/")}
    same = set()
    for p, t in out["content"]:
        if not p.startswith("ppt/"):
            continue
        n = mask_words(norm(t), allow)
        if (n.replace(MASK, "") and n not in PLACEHOLDER_LABELS and not is_trivial(n)
                and n in orig_set and n not in same):
            same.add(n)
            findings.append(("g", "文字与原文完全相同：「%s」（%s）" % (short(t.strip()), part_label(p))))

    # ---- (h) manifest 完整性
    n_slides = len([n for n in out["names"] if re.match(r"ppt/slides/slide\d+\.xml$", n)])
    if len(purposes) != n_slides:
        findings.append(("h", "manifest.yaml 里有 %d 个 purpose，骨架共 %d 页" % (len(purposes), n_slides)))
    for i, p in enumerate(purposes, 1):
        if not p.strip():
            findings.append(("h", "manifest.yaml 第 %d 页的 purpose 未填写" % i))

    # ---- 报告
    print("== 脱敏检查 ==")
    print("骨架：%s（%d 页）" % (args.out, n_slides))
    if allow_raw:
        print("已允许的词（--allow，请确认它们不含项目信息）：" + "、".join(allow_raw))
    if not words:
        print("提示：未提供敏感词文件（--words），只做了原文比对和结构检查。")
    if not findings:
        print("结果：通过，0 项问题。仍需人工逐页检查后才能带出。")
        return 0
    labels = {"a": "原文片段", "b": "数字", "c": "敏感词", "d": "残留部件", "e": "文档属性",
              "f": "超链接", "g": "文字未改", "h": "manifest"}
    for code in "abcdefgh":
        group = [m for c, m in findings if c == code]
        if group:
            print("\n[%s] %s（%d 项）" % (code, labels[code], len(group)))
            for m in group:
                print("  - " + m)
    print("\n结果：未通过，共 %d 项问题。请修改 replacements.json / manifest.yaml 后重新 apply 并 check。" % len(findings))
    return 1


# --------------------------------------------------------------------------- 渲染

def find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.exists(p):
            return p
    return None


def pdf_to_png(pdf, out_dir, stem):
    """依次尝试 pypdfium2、pdftoppm、PyMuPDF；都没有返回空列表。"""
    try:
        import pypdfium2 as pdfium  # type: ignore
        doc = pdfium.PdfDocument(pdf)
        paths = []
        for i in range(len(doc)):
            page = doc[i]
            bmp = page.render(scale=1280.0 / page.get_width())
            path = os.path.join(out_dir, "%s-%02d.png" % (stem, i + 1))
            bmp.to_pil().save(path)
            paths.append(path)
        return paths
    except ImportError:
        pass
    ppm = shutil.which("pdftoppm")
    if ppm:
        prefix = os.path.join(out_dir, stem)
        subprocess.run([ppm, "-png", "-scale-to-x", "1280", "-scale-to-y", "-1", pdf, prefix], check=True)
        return sorted(os.path.join(out_dir, f) for f in os.listdir(out_dir)
                      if f.startswith(stem) and f.endswith(".png"))
    try:
        import fitz  # type: ignore
        doc = fitz.open(pdf)
        paths = []
        for i, page in enumerate(doc):
            path = os.path.join(out_dir, "%s-%02d.png" % (stem, i + 1))
            page.get_pixmap(matrix=fitz.Matrix(1280.0 / page.rect.width, 1280.0 / page.rect.width)).save(path)
            paths.append(path)
        return paths
    except ImportError:
        return []


def cmd_render(args):
    soffice = find_soffice()
    if not soffice:
        print("未找到 soffice（LibreOffice），跳过渲染；请直接用 PowerPoint/WPS 打开 %s 查看。" % args.input)
        return 0
    os.makedirs(args.output, exist_ok=True)
    stem = os.path.splitext(os.path.basename(args.input))[0]
    with tempfile.TemporaryDirectory() as tmp:
        profile = "file:///" + os.path.join(tmp, "profile").replace("\\", "/").lstrip("/")
        try:
            subprocess.run([soffice, "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf",
                            "--outdir", tmp, os.path.abspath(args.input)],
                           check=True, timeout=300, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except (subprocess.SubprocessError, OSError) as exc:
            print("soffice 转换失败，跳过渲染：%s" % exc)
            return 0
        pdf = os.path.join(tmp, stem + ".pdf")
        if not os.path.exists(pdf):
            print("soffice 没有产出 PDF，跳过渲染。")
            return 0
        pngs = pdf_to_png(pdf, args.output, stem)
        if not pngs:
            dst = os.path.join(args.output, stem + ".pdf")
            shutil.copy(pdf, dst)
            print("没有可用的 PDF 转图工具（pypdfium2 / pdftoppm / PyMuPDF），已输出 PDF：%s" % dst)
            return 0
    print("已渲染 %d 页：" % len(pngs))
    for p in pngs:
        print("  " + p)
    return 0


# --------------------------------------------------------------------------- 入口

def build_parser():
    ap = argparse.ArgumentParser(
        prog="skeleton.py",
        description="把 PPT 页脱敏成只保留版式的骨架 pptx（文字换成引导性描述）。",
        epilog="典型流程：extract -> 填 replacements.json / manifest.yaml -> apply -> check -> render（可选）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("extract", help="提取所选页的文字，生成 texts.json / replacements.json / manifest.yaml")
    p.add_argument("input", metavar="IN.pptx")
    p.add_argument("--slides", required=True, help="页码，如 1,3-5（按给定顺序）；all 表示全部")
    p.add_argument("-o", "--output", required=True, metavar="WORK_DIR")
    p.add_argument("--force", action="store_true", help="覆盖 WORK_DIR 里已有的 replacements.json / manifest.yaml")
    p.set_defaults(func=cmd_extract)

    p = sub.add_parser("apply", help="写入引导文字并剥离敏感内容，生成骨架 pptx")
    p.add_argument("input", metavar="IN.pptx")
    p.add_argument("work_dir", metavar="WORK_DIR")
    p.add_argument("-o", "--output", required=True, metavar="OUT.pptx")
    p.set_defaults(func=cmd_apply)

    p = sub.add_parser("check", help="审计 OUT.pptx 是否仍有泄露，有问题时退出码为 1")
    p.add_argument("orig", metavar="ORIG.pptx")
    p.add_argument("out", metavar="OUT.pptx")
    p.add_argument("work_dir", metavar="WORK_DIR")
    p.add_argument("--words", metavar="SENSITIVE.txt", help="敏感词文件，一行一个（# 开头为注释）")
    p.add_argument("--allow", action="append", default=[], metavar="WORD",
                   help="用户确认无害的通用词，可重复；会在报告里列出")
    p.set_defaults(func=cmd_check)

    p = sub.add_parser("render", help="用 soffice 把 pptx 转成 PNG（没有 soffice 就跳过）")
    p.add_argument("input", metavar="OUT.pptx")
    p.add_argument("-o", "--output", required=True, metavar="DIR")
    p.set_defaults(func=cmd_render)
    return ap


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    args = build_parser().parse_args(argv)
    try:
        rc = args.func(args)
    except SkeletonError as exc:
        sys.stderr.write("错误：%s\n" % exc)
        return 1
    except (OSError, zipfile.BadZipFile, KeyError) as exc:
        sys.stderr.write("错误：读写文件失败（%s: %s）\n" % (type(exc).__name__, exc))
        return 1
    return rc or 0


if __name__ == "__main__":
    sys.exit(main())
